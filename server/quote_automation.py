"""Deterministic material-quote workflow.

This module owns the new per-bid quote request, message matching, review, price
application, history, follow-up, and simulation behavior. It intentionally does
not import or call the OpenAI client.
"""

from __future__ import annotations

import csv
import email
import hashlib
import hmac
import io
import itertools
import json
import math
import os
import re
import shutil
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta, timezone
from email import policy
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

import pdfplumber

from models import (
    _get_conn,
    delete_quote_email_draft,
    get_quote_email_draft,
    get_price_history,
    list_jobs,
    list_quote_email_drafts,
    list_quote_requests,
    load_job,
    log_activity,
    record_imported_file,
    record_job_artifact,
    record_material_price_decision,
    save_quotes,
    save_quote_email_draft,
    save_vendor_prices_from_quotes,
)
from quote_evidence import normalize_quote_unit


ARTIFACT_ROOT = Path(os.environ.get("ARTIFACT_ROOT", "/data/artifacts")).resolve()
DENVER = ZoneInfo("America/Denver")
OPEN_REQUEST_STATUSES = {
    "approved",
    "sending",
    "send_uncertain",
    "sent",
    "waiting",
    "overdue",
    "received_partial",
    "needs_review",
    "send_failed",
}
MATCHABLE_REQUEST_STATUSES = OPEN_REQUEST_STATUSES | {"stale"}
ASSIGNMENT_LEASE_SECONDS = 120
SIMULATION_SCENARIOS = {
    "all",
    "normal_reply",
    "changed_subject",
    "standalone_email",
    "ambiguous_bid",
    "partial_quote",
    "changed_price",
    "wrong_unit",
    "duplicate_reply",
    "broken_attachment",
    "no_response",
    "materials_changed",
    "send_failure",
}


class AssignmentRejectedError(ValueError):
    """The assignment failed after its review state was safely updated."""


class AssignmentLeaseLostError(ValueError):
    """A newer worker owns this email assignment."""
MAX_ATTACHMENT_BYTES = 25 * 1024 * 1024
MAX_SPREADSHEET_ROWS = 50_000
GENERIC_EMAIL_DOMAINS = {
    "gmail.com",
    "googlemail.com",
    "outlook.com",
    "hotmail.com",
    "live.com",
    "yahoo.com",
    "icloud.com",
}
HEADER_ALIASES = {
    "product": "description",
    "product_name": "description",
    "item": "description",
    "item_code": "item_code",
    "item_no": "item_code",
    "item_number": "item_code",
    "sku": "item_code",
    "code": "item_code",
    "description": "description",
    "vendor": "vendor",
    "manufacturer": "vendor",
    "unit_price": "unit_price",
    "unit_price_usd": "unit_price",
    "unit_cost": "unit_price",
    "price_each": "unit_price",
    "price_per_unit": "unit_price",
    "price": "ambiguous_price",
    "net_price": "ambiguous_price",
    "cost": "total_cost",
    "total": "total_cost",
    "total_cost": "total_cost",
    "extended_cost": "total_cost",
    "line_total": "total_cost",
    "amount": "total_cost",
    "unit": "unit",
    "uom": "unit",
    "freight": "freight",
    "lead_time": "lead_time",
    "notes": "notes",
}


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def iso_now() -> str:
    return utc_now().isoformat()


def _exact_unit_text(value: Any) -> str:
    return " ".join(str(value or "").strip().upper().split())


def _exact_unit_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", _exact_unit_text(value))


def delivery_failure_status(
    *,
    microsoft_may_have_accepted: bool = False,
) -> str:
    return "send_uncertain" if microsoft_may_have_accepted else "send_failed"


def _json_loads(value: Any, default: Any) -> Any:
    if isinstance(value, (dict, list)):
        return value
    try:
        parsed = json.loads(value or "")
    except (TypeError, ValueError, json.JSONDecodeError):
        return default
    return parsed if isinstance(parsed, type(default)) else default


def normalize_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def normalized_tokens(value: Any) -> tuple[str, ...]:
    return tuple(re.findall(r"[a-z0-9]+", str(value or "").lower()))


def contains_exact_fact(searchable: Any, fact: Any) -> bool:
    search_tokens = normalized_tokens(searchable)
    fact_tokens = normalized_tokens(fact)
    if not fact_tokens or len(fact_tokens) > len(search_tokens):
        return False
    width = len(fact_tokens)
    return any(
        search_tokens[index : index + width] == fact_tokens
        for index in range(len(search_tokens) - width + 1)
    )


def normalize_email(value: Any) -> str:
    text = str(value or "").strip().lower()
    match = re.search(r"<([^>]+)>", text)
    return (match.group(1) if match else text).strip()


def is_valid_email(value: Any) -> bool:
    address = normalize_email(value)
    return bool(
        re.fullmatch(
            r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@"
            r"[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?"
            r"(?:\.[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?)+",
            address,
        )
    )


def email_domain(value: Any) -> str:
    address = normalize_email(value)
    return address.rsplit("@", 1)[-1] if "@" in address else ""


def normalize_item_code(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _safe_name(value: str) -> str:
    clean = os.path.basename(value or "artifact")
    return re.sub(r"[^A-Za-z0-9._-]+", "_", clean) or "artifact"


def _artifact_relative(path: Path) -> str:
    return str(path.resolve().relative_to(ARTIFACT_ROOT))


def _write_artifact(
    data: bytes,
    *,
    filename: str,
    job_id: int | None = None,
    artifact_kind: str = "vendor_quote",
    conn=None,
    path_token: str = "",
    record_artifact: bool = True,
) -> dict:
    digest = _sha256_bytes(data)
    root = ARTIFACT_ROOT / (str(job_id) if job_id else "shared") / "uploads"
    root.mkdir(parents=True, exist_ok=True)
    token = f"{_safe_name(path_token)[:40]}_" if path_token else ""
    path = root / f"{digest[:12]}_{token}{_safe_name(filename)}"
    created = not path.exists()
    if created:
        path.write_bytes(data)
    result = {
        "file_name": filename,
        "file_hash": digest,
        "file_size": len(data),
        "artifact_path": _artifact_relative(path),
        "absolute_path": str(path),
        "artifact_kind": artifact_kind,
        "created": created,
    }
    if job_id and record_artifact:
        try:
            record_job_artifact(
                job_id,
                artifact_kind,
                result["artifact_path"],
                digest,
                len(data),
                conn=conn,
            )
        except Exception:
            if created:
                try:
                    path.unlink()
                except OSError:
                    pass
            raise
    return result


def _manifest_artifact_source(manifest: dict) -> dict | None:
    relative = str(manifest.get("artifact_path") or "")
    try:
        source = (ARTIFACT_ROOT / relative).resolve()
        source.relative_to(ARTIFACT_ROOT)
    except (ValueError, OSError):
        return None
    if not source.is_file():
        return None
    return {
        "file_name": str(manifest.get("file_name") or source.name),
        "file_hash": str(manifest.get("file_hash") or _sha256_bytes(source.read_bytes())),
        "file_size": int(manifest.get("file_size") or source.stat().st_size),
        "artifact_path": relative,
        "absolute_path": str(source),
        "artifact_kind": str(manifest.get("artifact_kind") or "vendor_quote"),
    }


def _copy_artifact_to_job(
    manifest: dict,
    job_id: int,
    *,
    conn=None,
    path_token: str = "",
    record_artifact: bool = True,
) -> dict | None:
    source_manifest = _manifest_artifact_source(manifest)
    if not source_manifest:
        return None
    source = Path(source_manifest["absolute_path"])
    return _write_artifact(
        source.read_bytes(),
        filename=source_manifest["file_name"],
        job_id=job_id,
        artifact_kind="vendor_quote",
        conn=conn,
        path_token=path_token,
        record_artifact=record_artifact,
    )


def _stage_assignment_artifacts(
    parsed_sources: Iterable[dict],
    *,
    job_id: int,
    assignment_token: str,
) -> tuple[list[dict], list[str]]:
    staged = []
    created_paths = []
    try:
        for manifest in parsed_sources:
            copied = _copy_artifact_to_job(
                manifest,
                job_id,
                path_token=f"assign-{assignment_token}",
                record_artifact=False,
            )
            if not copied:
                raise RuntimeError(
                    "Could not preserve quote proof: "
                    f"{manifest.get('file_name') or 'file'}."
                )
            staged.append(copied)
            if copied.get("created"):
                created_paths.append(str(copied["artifact_path"]))
        return staged, created_paths
    except Exception:
        _remove_uncommitted_artifact_files(created_paths)
        raise


def _remove_uncommitted_artifact_files(artifact_paths: Iterable[str]) -> None:
    paths = {str(value or "") for value in artifact_paths if str(value or "")}
    if not paths:
        return
    conn = _get_conn()
    try:
        referenced = {
            str(row["artifact_path"])
            for row in conn.execute(
                f"""SELECT artifact_path FROM job_artifacts
                    WHERE artifact_path IN ({",".join("?" for _ in paths)})""",
                tuple(paths),
            ).fetchall()
        }
    finally:
        conn.close()
    for relative in paths - referenced:
        try:
            path = (ARTIFACT_ROOT / relative).resolve()
            path.relative_to(ARTIFACT_ROOT)
            if path.is_file():
                path.unlink()
        except (OSError, ValueError):
            continue


def cleanup_orphan_assignment_artifacts() -> dict:
    pattern = re.compile(r"^[0-9a-f]{12}_assign-[0-9a-f]{32}_")
    candidates = [
        path
        for path in ARTIFACT_ROOT.glob("*/uploads/*")
        if path.is_file() and pattern.match(path.name)
    ]
    if not candidates:
        return {"checked": 0, "removed": 0}
    conn = _get_conn()
    removed = 0
    try:
        for path in candidates:
            try:
                relative = _artifact_relative(path)
            except (OSError, ValueError):
                continue
            referenced = conn.execute(
                """SELECT 1 FROM job_artifacts WHERE artifact_path=?
                   UNION ALL
                   SELECT 1 FROM imported_files WHERE artifact_path=?
                   LIMIT 1""",
                (relative, relative),
            ).fetchone()
            if referenced:
                continue
            try:
                path.unlink()
                removed += 1
            except OSError:
                continue
    finally:
        conn.close()
    return {"checked": len(candidates), "removed": removed}


def _material_snapshot(materials: Iterable[dict]) -> list[dict]:
    result = []
    for material in materials:
        result.append(
            {
                "id": material.get("id"),
                "item_code": str(material.get("item_code") or "").strip(),
                "description": str(material.get("description") or "").strip(),
                "unit": _exact_unit_text(material.get("unit")),
                "quantity": round(
                    float(material.get("order_qty") or material.get("installed_qty") or 0),
                    4,
                ),
                "vendor": str(material.get("vendor") or "").strip(),
            }
        )
    return result


def snapshot_hash(materials: Iterable[dict]) -> str:
    payload = json.dumps(
        _material_snapshot(materials),
        sort_keys=True,
        separators=(",", ":"),
    )
    return _sha256_bytes(payload.encode("utf-8"))


def _job_source_fingerprint(job: dict) -> str:
    materials = [
        {
            key: value
            for key, value in material.items()
            if key != "vendor"
        }
        for material in _material_snapshot(job.get("materials") or [])
    ]
    values = {
        "project_name": job.get("project_name") or "",
        "gc_name": job.get("gc_name") or "",
        "address": job.get("address") or "",
        "city": job.get("city") or "",
        "state": job.get("state") or "",
        "zip": job.get("zip") or "",
        "materials": materials,
    }
    payload = json.dumps(values, sort_keys=True, separators=(",", ":"))
    return _sha256_bytes(payload.encode("utf-8"))


def _quote_subject(job: dict) -> str:
    project = str(job.get("project_name") or "Project").strip()
    return f"Quote Request - {project} - Flooring Materials"


def _quote_body(job: dict, materials: list[dict]) -> str:
    lines = [f"Project: {job.get('project_name') or ''}"]
    if job.get("gc_name"):
        lines.append(f"General Contractor: {job['gc_name']}")
    location = ", ".join(
        str(value).strip()
        for value in (
            job.get("address"),
            job.get("city"),
            job.get("state"),
            job.get("zip"),
        )
        if str(value or "").strip()
    )
    if location:
        lines.append(f"Location: {location}")
    lines.extend(
        [
            "",
            "Please provide unit pricing, freight, availability, and lead time for:",
            "",
        ]
    )
    for material in materials:
        code = str(material.get("item_code") or "").strip()
        description = str(material.get("description") or "").strip()
        label = " - ".join(value for value in (code, description) if value)
        quantity = float(material.get("order_qty") or material.get("installed_qty") or 0)
        unit = normalize_quote_unit(material.get("unit"))
        lines.append(f"- {label}: {quantity:g} {unit}".rstrip())
    lines.extend(["", "Thank you."])
    return "\n".join(lines)


def _vendor_rows() -> list[dict]:
    conn = _get_conn()
    try:
        return [
            dict(row)
            for row in conn.execute(
                "SELECT id, name, contact_name, contact_email FROM vendors ORDER BY name"
            ).fetchall()
        ]
    finally:
        conn.close()


def _vendor_for_material(name: str, vendors: list[dict]) -> dict | None:
    normalized = normalize_text(name)
    if not normalized:
        return None
    exact = [vendor for vendor in vendors if normalize_text(vendor.get("name")) == normalized]
    return exact[0] if len(exact) == 1 else None


def save_vendor_mapping(vendor_name: str, vendor_email: str) -> int:
    name = str(vendor_name or "").strip()
    email_address = normalize_email(vendor_email)
    if not name:
        raise ValueError("A vendor name is required.")
    if not is_valid_email(email_address):
        raise ValueError("Enter a valid vendor email address.")
    conn = _get_conn()
    try:
        now = iso_now()
        row = conn.execute(
            "SELECT id FROM vendors WHERE lower(name)=lower(?)",
            (name,),
        ).fetchone()
        if row:
            conn.execute(
                """UPDATE vendors
                   SET contact_email=?, updated_at=?
                   WHERE id=?""",
                (email_address, now, row["id"]),
            )
            vendor_id = int(row["id"])
        else:
            cursor = conn.execute(
                """INSERT INTO vendors
                   (name, contact_email, created_at, updated_at)
                   VALUES (?, ?, ?, ?)""",
                (name, email_address, now, now),
            )
            vendor_id = int(cursor.lastrowid)
        conn.commit()
        return vendor_id
    finally:
        conn.close()


def _request_material_ids(request: dict) -> set[int]:
    raw_items = (
        request.get("material_snapshot")
        or request.get("material_ids")
        or request.get("material_snapshot_json")
        or []
    )
    if isinstance(raw_items, str):
        raw_items = _json_loads(raw_items, [])
    result = set()
    for item in raw_items if isinstance(raw_items, list) else []:
        value = item.get("id") if isinstance(item, dict) else item
        try:
            result.add(int(value))
        except (TypeError, ValueError):
            continue
    return result


def _material_price_history(material: dict, job_id: int) -> dict:
    history = get_price_history(
        item_code=material.get("item_code"),
        exclude_job_id=job_id,
        verified_only=True,
    )
    unit = normalize_quote_unit(material.get("unit"))
    verified = [
        record
        for record in history.get("records") or []
        if str(record.get("source_hash") or "").strip()
        and normalize_quote_unit(record.get("unit")) == unit
        and float(record.get("unit_price") or 0) > 0
    ]
    prediction = None
    if len(verified) >= 3:
        prediction = round(
            sum(float(record["unit_price"]) for record in verified) / len(verified),
            2,
        )
    latest = verified[0] if verified else None
    previous = verified[1] if len(verified) > 1 else None
    percentage_change = None
    if latest and previous and float(previous.get("unit_price") or 0) > 0:
        percentage_change = round(
            (
                float(latest.get("unit_price") or 0)
                - float(previous.get("unit_price") or 0)
            )
            / float(previous["unit_price"])
            * 100,
            1,
        )
    latest_at = _parse_iso(
        (latest or {}).get("verified_at")
        or (latest or {}).get("quote_date")
        or (latest or {}).get("created_at")
    )
    age_days = (
        max(0, (utc_now() - latest_at).days)
        if latest_at
        else None
    )
    verified_prices = [float(record["unit_price"]) for record in verified]
    return {
        **history,
        "records": verified,
        "min": min(verified_prices) if verified_prices else None,
        "max": max(verified_prices) if verified_prices else None,
        "avg": (
            round(sum(verified_prices) / len(verified_prices), 2)
            if verified_prices
            else None
        ),
        "latest": latest,
        "verified_same_unit_count": len(verified),
        "predicted_price": prediction,
        "prediction_is_suggestion": prediction is not None,
        "age_days": age_days,
        "percentage_change": percentage_change,
        "similar_bid_count": len(
            {record.get("job_id") for record in verified if record.get("job_id")}
        ),
    }


def build_quote_plan(job_id: int) -> dict:
    job = load_job(job_id)
    if not job:
        raise ValueError("Job not found")
    vendors = _vendor_rows()
    open_requests = [
        request
        for request in list_quote_requests(job_id)
        if str(request.get("status") or "").lower() in OPEN_REQUEST_STATUSES
    ]
    requested_ids = {
        material_id
        for request in open_requests
        for material_id in _request_material_ids(request)
    }
    groups: dict[str, dict] = {}
    price_history_by_material: dict[str, dict] = {}
    for material in job.get("materials") or []:
        if material.get("id") is not None:
            price_history_by_material[str(material["id"])] = _material_price_history(
                material,
                job_id,
            )
        if float(material.get("unit_price") or 0) > 0:
            continue
        vendor_name = str(material.get("vendor") or "").strip()
        key = normalize_text(vendor_name) or "unassigned"
        vendor = _vendor_for_material(vendor_name, vendors)
        group = groups.setdefault(
            key,
            {
                "vendor_id": vendor.get("id") if vendor else None,
                "vendor_name": vendor.get("name") if vendor else (vendor_name or "Unassigned"),
                "vendor_email": str(vendor.get("contact_email") or "").strip() if vendor else "",
                "contact_name": str(vendor.get("contact_name") or "").strip() if vendor else "",
                "materials": [],
                "issues": [],
            },
        )
        row = {
            "id": material.get("id"),
            "item_code": material.get("item_code") or "",
            "description": material.get("description") or "",
            "unit": _exact_unit_text(material.get("unit")),
            "quantity": float(material.get("order_qty") or material.get("installed_qty") or 0),
            "already_requested": material.get("id") in requested_ids,
            "price_history": price_history_by_material.get(
                str(material.get("id")),
                {},
            ),
        }
        group["materials"].append(row)

    planned_groups = []
    for group in sorted(groups.values(), key=lambda value: value["vendor_name"].lower()):
        unsent = [material for material in group["materials"] if not material["already_requested"]]
        if group["vendor_name"] == "Unassigned":
            group["issues"].append("Assign a vendor before sending.")
        if not group["vendor_email"]:
            group["issues"].append("Add a vendor email before sending.")
        if not unsent:
            group["issues"].append("Every unpriced material in this group already has an open request.")
        group["subject"] = _quote_subject(job)
        source_materials = [
            material
            for material in job.get("materials") or []
            if material.get("id") in {item["id"] for item in unsent}
        ]
        group["body"] = _quote_body(job, source_materials)
        group["request_fingerprint"] = snapshot_hash(source_materials)
        group["can_send"] = not group["issues"] and bool(unsent)
        group["materials_to_send"] = unsent
        planned_groups.append(group)

    return {
        "job_id": job_id,
        "source_fingerprint": _job_source_fingerprint(job),
        "groups": planned_groups,
        "unpriced_count": sum(len(group["materials"]) for group in planned_groups),
        "ready_group_count": sum(1 for group in planned_groups if group["can_send"]),
        "price_history": price_history_by_material,
        "matching_engine": {"name": "deterministic-v1", "ai_calls": 0},
    }


def _draft_material_ids(group: dict) -> list[int]:
    raw_items = (
        group.get("material_ids")
        or group.get("materials_to_send")
        or group.get("materials")
        or []
    )
    result = []
    for item in raw_items if isinstance(raw_items, list) else []:
        value = item.get("id") if isinstance(item, dict) else item
        try:
            material_id = int(value)
        except (TypeError, ValueError):
            continue
        if material_id not in result:
            result.append(material_id)
    return result


def _draft_group_issues(
    group: dict,
    *,
    materials: list[dict],
    current_materials: dict[int, dict],
    open_material_ids: set[int],
    stale: bool,
) -> list[str]:
    issues = []
    if stale:
        issues.append("Bid materials changed after this draft was saved.")
    vendor_name = str(group.get("vendor_name") or "").strip()
    if not vendor_name or normalize_text(vendor_name) == "unassigned":
        issues.append("Choose a vendor.")
    if not is_valid_email(group.get("vendor_email")):
        issues.append("Add a valid vendor email.")
    if not str(group.get("subject") or "").strip():
        issues.append("Add an email subject.")
    if not str(group.get("body") or "").strip():
        issues.append("Add the email message.")
    if not materials:
        issues.append("Choose at least one material.")
    missing_ids = [
        material_id
        for material_id in _draft_material_ids(group)
        if material_id not in current_materials
    ]
    if missing_ids:
        issues.append("A material is no longer part of this bid.")
    if any(float(material.get("unit_price") or 0) > 0 for material in materials):
        issues.append("A material already has a price.")
    if any(int(material["id"]) in open_material_ids for material in materials):
        issues.append("A material already has an open quote request.")
    if materials and not hmac.compare_digest(
        snapshot_hash(materials),
        str(group.get("request_fingerprint") or ""),
    ):
        issues.append("The material group changed after this draft was saved.")
    return list(dict.fromkeys(issues))


def material_quote_email_draft(job_id: int) -> dict | None:
    """Return one draft with current material and send-safety state."""
    draft = get_quote_email_draft(job_id)
    if not draft:
        return None
    job = load_job(job_id)
    if not job:
        return None
    current_fingerprint = _job_source_fingerprint(job)
    stale = not hmac.compare_digest(
        current_fingerprint,
        str(draft.get("source_fingerprint") or ""),
    )
    current_materials = {
        int(material["id"]): material
        for material in job.get("materials") or []
        if material.get("id") is not None
    }
    open_material_ids = {
        material_id
        for request in list_quote_requests(job_id)
        if str(request.get("status") or "").lower() in OPEN_REQUEST_STATUSES
        for material_id in _request_material_ids(request)
    }
    groups = []
    for stored_group in draft.get("groups") or []:
        group = dict(stored_group)
        material_ids = _draft_material_ids(group)
        materials = [
            current_materials[material_id]
            for material_id in material_ids
            if material_id in current_materials
        ]
        issues = _draft_group_issues(
            group,
            materials=materials,
            current_materials=current_materials,
            open_material_ids=open_material_ids,
            stale=stale,
        )
        groups.append(
            {
                **group,
                "material_ids": material_ids,
                "materials": [
                    {
                        "id": material.get("id"),
                        "item_code": material.get("item_code") or "",
                        "description": material.get("description") or "",
                        "unit": _exact_unit_text(material.get("unit")),
                        "quantity": float(
                            material.get("order_qty")
                            or material.get("installed_qty")
                            or 0
                        ),
                        "unit_price": float(material.get("unit_price") or 0),
                    }
                    for material in materials
                ],
                "issues": issues,
                "can_send": not issues,
            }
        )
    return {
        **draft,
        "project_name": job.get("project_name") or "",
        "gc_name": job.get("gc_name") or "",
        "slug": job.get("slug") or "",
        "current_source_fingerprint": current_fingerprint,
        "stale": stale,
        "groups": groups,
        "group_count": len(groups),
        "ready_group_count": sum(1 for group in groups if group["can_send"]),
        "needs_setup_count": sum(1 for group in groups if not group["can_send"]),
    }


def save_material_quote_email_draft(
    job_id: int,
    *,
    source_fingerprint: str,
    groups: list[dict],
    prepared_by: str = "",
) -> dict:
    """Validate and save bid-owned vendor email groups without sending."""
    job = load_job(job_id)
    if not job:
        raise ValueError("Job not found")
    current_source_fingerprint = _job_source_fingerprint(job)
    if not source_fingerprint or not hmac.compare_digest(
        current_source_fingerprint,
        str(source_fingerprint),
    ):
        raise ValueError(
            "This quote plan is out of date. Refresh it before saving."
        )
    if not isinstance(groups, list) or not groups:
        raise ValueError("Add at least one vendor group.")
    current_materials = {
        int(material["id"]): material
        for material in job.get("materials") or []
        if material.get("id") is not None
    }
    open_material_ids = {
        material_id
        for request in list_quote_requests(job_id)
        if str(request.get("status") or "").lower() in OPEN_REQUEST_STATUSES
        for material_id in _request_material_ids(request)
    }
    existing = get_quote_email_draft(job_id) or {}
    existing_group_ids = {
        tuple(sorted(_draft_material_ids(group))): str(group.get("group_id") or "")
        for group in existing.get("groups") or []
        if _draft_material_ids(group)
    }
    used_material_ids: set[int] = set()
    normalized_groups = []
    for raw_group in groups[:100]:
        if not isinstance(raw_group, dict):
            continue
        material_ids = _draft_material_ids(raw_group)
        if not material_ids:
            continue
        duplicate_ids = used_material_ids.intersection(material_ids)
        if duplicate_ids:
            raise ValueError(
                f"Material {min(duplicate_ids)} appears in more than one vendor group."
            )
        missing_ids = set(material_ids) - set(current_materials)
        if missing_ids:
            raise ValueError(
                f"Material {min(missing_ids)} is no longer part of this bid."
            )
        unavailable_ids = {
            material_id
            for material_id in material_ids
            if (
                float(current_materials[material_id].get("unit_price") or 0) > 0
                or material_id in open_material_ids
            )
        }
        if unavailable_ids:
            raise ValueError(
                f"Material {min(unavailable_ids)} is already priced or requested."
            )
        selected = [
            material
            for material in job.get("materials") or []
            if int(material.get("id") or 0) in set(material_ids)
        ]
        current_group_id = existing_group_ids.get(tuple(sorted(material_ids)))
        group_id = current_group_id or uuid.uuid4().hex
        vendor_name = str(raw_group.get("vendor_name") or "").strip()
        vendor = _vendor_for_material(vendor_name, _vendor_rows()) if vendor_name else None
        normalized_groups.append(
            {
                "group_id": group_id,
                "vendor_id": raw_group.get("vendor_id") or (
                    vendor.get("id") if vendor else None
                ),
                "vendor_name": vendor_name or "Unassigned",
                "vendor_email": normalize_email(
                    raw_group.get("vendor_email")
                    or (vendor.get("contact_email") if vendor else "")
                ),
                "contact_name": str(
                    raw_group.get("contact_name")
                    or (vendor.get("contact_name") if vendor else "")
                    or ""
                ).strip(),
                "material_ids": material_ids,
                "subject": str(raw_group.get("subject") or "").strip()
                or _quote_subject(job),
                "body": str(raw_group.get("body") or "").strip()
                or _quote_body(job, selected),
                "request_fingerprint": snapshot_hash(selected),
            }
        )
        used_material_ids.update(material_ids)
    if not normalized_groups:
        raise ValueError("Add at least one vendor group with materials.")
    save_quote_email_draft(
        job_id,
        source_fingerprint=current_source_fingerprint,
        groups=normalized_groups,
        prepared_by=prepared_by,
    )
    return material_quote_email_draft(job_id)


def update_material_quote_email_draft_group(
    job_id: int,
    group_id: str,
    changes: dict,
) -> dict:
    """Edit recipient and email copy while keeping the material snapshot locked."""
    draft = get_quote_email_draft(job_id)
    if not draft:
        raise ValueError("No saved email draft exists for this bid.")
    allowed = {"vendor_email", "subject", "body"}
    updated = False
    groups = []
    for raw_group in draft.get("groups") or []:
        group = dict(raw_group)
        if str(group.get("group_id") or "") == str(group_id):
            for field in allowed:
                if field in changes:
                    value = str(changes.get(field) or "")
                    group[field] = (
                        normalize_email(value)
                        if field == "vendor_email"
                        else value.strip()
                    )
            updated = True
        groups.append(group)
    if not updated:
        raise ValueError("Email draft group not found.")
    save_quote_email_draft(
        job_id,
        source_fingerprint=str(draft.get("source_fingerprint") or ""),
        groups=groups,
        prepared_by=str(draft.get("prepared_by") or ""),
    )
    return material_quote_email_draft(job_id)


def remove_material_quote_email_draft(job_id: int) -> bool:
    return delete_quote_email_draft(job_id)


def _mailbox_quote_requests(mailbox_email: str | None) -> list[dict]:
    if not mailbox_email:
        return []
    conn = _get_conn()
    try:
        rows = conn.execute(
            """SELECT qr.*, j.project_name, j.gc_name, j.slug
               FROM quote_requests qr
               JOIN jobs j ON j.id=qr.job_id
               WHERE lower(qr.mailbox_email)=?
               ORDER BY qr.created_at DESC, qr.id DESC""",
            (normalize_email(mailbox_email),),
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()


def _all_quote_requests_for_bid_summary() -> list[dict]:
    """Return request rows for status counts without exposing mailbox content."""
    conn = _get_conn()
    try:
        rows = conn.execute(
            """SELECT qr.*, j.project_name, j.gc_name, j.slug
               FROM quote_requests qr
               JOIN jobs j ON j.id=qr.job_id
               ORDER BY qr.created_at DESC, qr.id DESC"""
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        conn.close()


def _email_center_request(request: dict, job: dict) -> dict:
    item = dict(request)
    item["material_snapshot"] = _json_loads(
        item.get("material_snapshot_json"),
        [],
    )
    item["material_ids"] = sorted(_request_material_ids(item))
    item["materials"] = _request_materials(int(item["id"]))
    item["followups"] = _followups_for_request(int(item["id"]))
    item["messages"] = _messages_for_request(int(item["id"]))
    workflow_state = _request_workflow_state(item, job)
    item["stale"] = workflow_state["stale"]
    item["status"] = workflow_state["status"]
    if item["stale"] and str(request.get("status") or "").lower() not in {
        "complete",
        "cancelled",
        "stale",
        "received",
    }:
        mark_request_stale(int(item["id"]))
    item.pop("send_token", None)
    return item


def material_quotes_email_center(
    *,
    mailbox_email: str | None,
    outlook: dict | None = None,
) -> dict:
    """Return the global mailbox view while keeping drafts bid-owned."""
    drafts = []
    for stored in list_quote_email_drafts():
        draft = material_quote_email_draft(int(stored["job_id"]))
        if draft:
            drafts.append(draft)
    requests = []
    if mailbox_email:
        jobs: dict[int, dict] = {}
        for request in _mailbox_quote_requests(mailbox_email):
            job_id = int(request["job_id"])
            job = jobs.get(job_id)
            if job is None:
                job = load_job(job_id) or {}
                jobs[job_id] = job
            requests.append(_email_center_request(request, job))
        review = quote_review(mailbox_email=mailbox_email)
    else:
        review = {"messages": [], "price_matches": [], "mailbox_locked": True}

    ready_count = sum(draft.get("ready_group_count", 0) for draft in drafts)
    needs_you_request_count = sum(
        1
        for request in requests
        if request.get("status")
        in {"send_failed", "send_uncertain", "stale", "needs_review"}
    )
    waiting_count = sum(
        1
        for request in requests
        if request.get("status")
        in {"approved", "sending", "sent", "waiting", "received_partial"}
    )
    overdue_count = sum(
        1 for request in requests if request.get("status") == "overdue"
    )
    complete_count = sum(
        1
        for request in requests
        if request.get("status") in {"complete", "received", "cancelled"}
    )
    needs_you_count = (
        sum(draft.get("needs_setup_count", 0) for draft in drafts)
        + needs_you_request_count
        + len(review.get("messages") or [])
        + len(review.get("price_matches") or [])
    )
    job_rows = [
        {
            "id": job.get("id"),
            "slug": job.get("slug") or "",
            "project_name": job.get("project_name") or "Untitled Bid",
            "gc_name": job.get("gc_name") or "",
            "city": job.get("city") or "",
            "state": job.get("state") or "",
        }
        for job in list_jobs()
    ]
    return {
        "outlook": outlook or {"configured": False, "connected": False},
        "mailbox_locked": not bool(mailbox_email),
        "drafts": drafts,
        "requests": requests,
        "review": review,
        "jobs": job_rows,
        "summary": {
            "ready_to_send": ready_count,
            "waiting": waiting_count,
            "overdue": overdue_count,
            "needs_you": needs_you_count,
            "complete": complete_count,
        },
        "matching_engine": {"name": "deterministic-v1", "ai_calls": 0},
    }


def material_quote_bid_summaries(mailbox_email: str | None = None) -> dict:
    """Return compact bid rows sorted by the estimator's next action."""
    request_rows = (
        _mailbox_quote_requests(mailbox_email)
        if mailbox_email
        else _all_quote_requests_for_bid_summary()
    )
    requests_by_job: dict[int, list[dict]] = {}
    for request in request_rows:
        requests_by_job.setdefault(int(request["job_id"]), []).append(request)
    review = (
        quote_review(mailbox_email=mailbox_email)
        if mailbox_email
        else {"messages": [], "price_matches": []}
    )
    matching_by_job: dict[int, int] = {}
    for message in review.get("messages") or []:
        possible_job_ids = {
            int(message["matched_job_id"])
            for _ in (0,)
            if message.get("matched_job_id") is not None
        }
        possible_job_ids.update(
            int(candidate["job_id"])
            for candidate in message.get("candidates") or []
            if candidate.get("job_id") is not None
        )
        for job_id in possible_job_ids:
            matching_by_job[job_id] = matching_by_job.get(job_id, 0) + 1
    price_review_by_job: dict[int, int] = {}
    for match in review.get("price_matches") or []:
        job_id = int(match["job_id"])
        price_review_by_job[job_id] = price_review_by_job.get(job_id, 0) + 1

    stage_labels = {
        "needs_review": "Needs Your Review",
        "overdue": "Overdue",
        "ready_to_send": "Ready to Send",
        "needs_setup": "Needs Setup",
        "waiting": "Waiting on Vendor",
        "complete": "Complete",
    }
    priority = {
        "needs_review": 0,
        "overdue": 1,
        "ready_to_send": 2,
        "needs_setup": 3,
        "waiting": 4,
        "complete": 5,
    }
    bids = []
    for summary in list_jobs():
        job_id = int(summary["id"])
        job = load_job(job_id) or {}
        materials = job.get("materials") or []
        unpriced = [
            material
            for material in materials
            if float(material.get("unit_price") or 0) <= 0
        ]
        requests = []
        requested_ids: set[int] = set()
        status_counts: dict[str, int] = {}
        for request in requests_by_job.get(job_id, []):
            state = _request_workflow_state(request, job)
            status = state["status"]
            status_counts[status] = status_counts.get(status, 0) + 1
            if status not in {"complete", "received", "cancelled"}:
                requested_ids.update(_request_material_ids(request))
            requests.append({**request, **state})
        unrequested_count = sum(
            1
            for material in unpriced
            if int(material.get("id") or 0) not in requested_ids
        )
        draft = material_quote_email_draft(job_id)
        ready_draft_count = draft.get("ready_group_count", 0) if draft else 0
        draft_setup_count = draft.get("needs_setup_count", 0) if draft else 0
        needs_matching = matching_by_job.get(job_id, 0)
        price_review = price_review_by_job.get(job_id, 0)
        request_needs_review = sum(
            status_counts.get(status, 0)
            for status in ("send_failed", "send_uncertain", "stale", "needs_review")
        )
        needs_review_count = (
            needs_matching
            + price_review
            + request_needs_review
            + (draft_setup_count if draft and draft.get("stale") else 0)
        )
        open_request_count = sum(
            count
            for status, count in status_counts.items()
            if status not in {"complete", "received", "cancelled"}
        )
        completed_request_count = sum(
            status_counts.get(status, 0)
            for status in ("complete", "received")
        )
        cancelled_request_count = status_counts.get("cancelled", 0)
        if needs_review_count:
            stage = "needs_review"
        elif status_counts.get("overdue", 0):
            stage = "overdue"
        elif ready_draft_count:
            stage = "ready_to_send"
        elif unrequested_count or not materials:
            stage = "needs_setup"
        elif open_request_count:
            stage = "waiting"
        else:
            stage = "complete"
        slug = job.get("slug") or str(job_id)
        if stage == "needs_setup":
            next_action = {
                "label": "Set Up Quote Email" if materials else "Open Bid",
                "url": (
                    f"/jobs/{slug}?step=quotes"
                    if materials
                    else f"/jobs/{slug}"
                ),
            }
        elif stage == "complete":
            next_action = {"label": "Open Bid", "url": f"/jobs/{slug}"}
        else:
            email_view = {
                "needs_review": "needs_you",
                "overdue": "overdue",
                "ready_to_send": "ready_to_send",
                "waiting": "waiting",
            }[stage]
            next_action = {
                "label": {
                    "needs_review": "Review Email",
                    "overdue": "Check Overdue Email",
                    "ready_to_send": "Review & Send",
                    "waiting": "Check Vendor Reply",
                }[stage],
                "url": f"/jobs/bids/quote-emails?job={job_id}&view={email_view}",
            }
        vendor_keys = {
            normalize_text(material.get("vendor")) or "unassigned"
            for material in unpriced
        }
        bids.append(
            {
                "job_id": job_id,
                "slug": job.get("slug") or "",
                "project_name": job.get("project_name") or "Untitled Bid",
                "gc_name": job.get("gc_name") or "",
                "salesperson": job.get("salesperson") or "",
                "city": job.get("city") or "",
                "state": job.get("state") or "",
                "material_count": len(materials),
                "unpriced_count": len(unpriced),
                "unrequested_count": unrequested_count,
                "vendor_group_count": len(vendor_keys),
                "draft_group_count": draft.get("group_count", 0) if draft else 0,
                "draft_ready_count": ready_draft_count,
                "draft_stale": bool(draft and draft.get("stale")),
                "request_count": len(requests),
                "active_request_count": open_request_count,
                "completed_request_count": completed_request_count,
                "cancelled_request_count": cancelled_request_count,
                "request_statuses": status_counts,
                "needs_matching_count": needs_matching,
                "price_review_count": price_review,
                "quote_stage": stage,
                "quote_stage_label": stage_labels[stage],
                "next_action": next_action,
                "updated_at": (
                    draft.get("updated_at")
                    if draft
                    else job.get("created_at")
                ),
            }
        )
    bids.sort(
        key=lambda bid: (
            priority[bid["quote_stage"]],
            str(bid.get("project_name") or "").lower(),
        )
    )
    summary_counts = {
        key: sum(1 for bid in bids if bid["quote_stage"] == key)
        for key in stage_labels
    }
    return {
        "bids": bids,
        "summary": summary_counts,
        "mailbox_locked": not bool(mailbox_email),
        "matching_engine": {"name": "deterministic-v1", "ai_calls": 0},
    }


def _parse_iso(value: Any) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed


def add_business_days(start: datetime, days: int) -> datetime:
    local = start.astimezone(DENVER)
    cursor = local.date()
    remaining = days
    while remaining > 0:
        cursor += timedelta(days=1)
        if cursor.weekday() < 5:
            remaining -= 1
    return datetime.combine(cursor, time(hour=9), tzinfo=DENVER).astimezone(timezone.utc)


def business_days_between(start: datetime, end: datetime) -> int:
    cursor = start.astimezone(DENVER).date()
    end_date = end.astimezone(DENVER).date()
    count = 0
    while cursor < end_date:
        cursor += timedelta(days=1)
        if cursor.weekday() < 5:
            count += 1
    return count


def schedule_followups(request_id: int, sent_at: datetime) -> None:
    now = iso_now()
    conn = _get_conn()
    try:
        for number, days in ((1, 3), (2, 6)):
            conn.execute(
                """INSERT OR IGNORE INTO quote_followup_events
                   (quote_request_id, followup_number, scheduled_for, status, created_at)
                   VALUES (?, ?, ?, 'scheduled', ?)""",
                (request_id, number, add_business_days(sent_at, days).isoformat(), now),
            )
        conn.commit()
    finally:
        conn.close()


def create_approved_request(
    job_id: int,
    *,
    vendor_name: str,
    vendor_email: str,
    vendor_id: int | None,
    material_ids: list[int],
    subject: str,
    body: str,
    reviewer_name: str,
    mailbox_email: str,
    expected_request_fingerprint: str,
    expected_source_fingerprint: str,
) -> dict:
    vendor_name = str(vendor_name or "").strip()
    if not vendor_name:
        raise ValueError("A vendor name is required.")
    if not is_valid_email(vendor_email):
        raise ValueError("Enter a valid vendor email address.")
    selected_ids = {int(value) for value in material_ids}
    if not selected_ids:
        raise ValueError("Select at least one unpriced material.")
    now = iso_now()
    conn = _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        job_row = conn.execute(
            "SELECT * FROM jobs WHERE id=?",
            (job_id,),
        ).fetchone()
        material_rows = conn.execute(
            "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
            (job_id,),
        ).fetchall()
        if not job_row:
            raise ValueError("Job not found")
        job = {
            **dict(job_row),
            "materials": [dict(row) for row in material_rows],
        }
        current_source_fingerprint = _job_source_fingerprint(job)
        if (
            not expected_source_fingerprint
            or not hmac.compare_digest(
                current_source_fingerprint,
                str(expected_source_fingerprint),
            )
        ):
            raise ValueError(
                "This quote plan is out of date. Refresh it before sending."
            )
        current = {
            int(material["id"]): material
            for material in job["materials"]
            if material.get("id") is not None
        }
        missing_ids = selected_ids - set(current)
        if missing_ids:
            missing_id = min(missing_ids)
            raise ValueError(
                f"Material {missing_id} is no longer part of this bid."
            )
        selected = [
            material
            for material in job["materials"]
            if int(material["id"]) in selected_ids
        ]
        for material in selected:
            if float(material.get("unit_price") or 0) > 0:
                raise ValueError(
                    f"{material.get('item_code') or material.get('description')} is already priced."
                )
        current_request_fingerprint = snapshot_hash(selected)
        if (
            not expected_request_fingerprint
            or not hmac.compare_digest(
                current_request_fingerprint,
                str(expected_request_fingerprint),
            )
        ):
            raise ValueError(
                "This vendor group changed. Refresh it before sending."
            )
        approved_materials = [
            {**material, "vendor": vendor_name}
            for material in selected
        ]
        material_snapshot = _material_snapshot(approved_materials)
        material_hash = snapshot_hash(approved_materials)
        resolved_subject = subject.strip() or _quote_subject(job)
        resolved_body = body.strip() or _quote_body(job, selected)
        body_hash = _sha256_bytes(resolved_body.encode("utf-8"))
        open_rows = conn.execute(
            f"""SELECT id, material_ids, material_snapshot_json
                FROM quote_requests
                WHERE job_id=?
                  AND status IN ({",".join("?" for _ in OPEN_REQUEST_STATUSES)})""",
            (
                job_id,
                *sorted(OPEN_REQUEST_STATUSES),
            ),
        ).fetchall()
        for open_row in open_rows:
            if selected_ids & _request_material_ids(dict(open_row)):
                raise ValueError(
                    f"Quote request {open_row['id']} already covers one of these materials."
                )
        cursor = conn.execute(
            """INSERT INTO quote_requests
               (job_id, vendor_id, vendor_name, vendor_email, status, material_ids,
                material_snapshot_json, material_snapshot_hash, source_fingerprint,
                body_hash, subject, request_text, approved_at, approved_by,
                mailbox_email, created_at)
               VALUES (?, ?, ?, ?, 'approved', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                job_id,
                vendor_id,
                vendor_name,
                normalize_email(vendor_email),
                json.dumps(
                    [
                        {"id": material["id"], "item_code": material.get("item_code") or ""}
                        for material in selected
                    ]
                ),
                json.dumps(material_snapshot),
                material_hash,
                current_source_fingerprint,
                body_hash,
                resolved_subject,
                resolved_body,
                now,
                reviewer_name.strip() or "Estimator",
                normalize_email(mailbox_email),
                now,
            ),
        )
        request_id = int(cursor.lastrowid)
        for material in material_snapshot:
            conn.execute(
                """INSERT INTO quote_request_materials
                   (quote_request_id, material_id, item_code, description, unit,
                    quantity, vendor_name, snapshot_hash, created_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    request_id,
                    material.get("id"),
                    material.get("item_code"),
                    material.get("description"),
                    material.get("unit"),
                    material.get("quantity"),
                    vendor_name,
                    material_hash,
                    now,
                ),
            )
        placeholders = ",".join("?" for _ in selected_ids)
        conn.execute(
            f"""UPDATE job_materials
                SET vendor=?
                WHERE job_id=? AND id IN ({placeholders})
                  AND COALESCE(unit_price, 0)<=0""",
            (vendor_name, job_id, *selected_ids),
        )
        conn.commit()
        row = conn.execute(
            "SELECT * FROM quote_requests WHERE id=?", (request_id,)
        ).fetchone()
        return dict(row)
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def mark_request_sent(
    request_id: int,
    *,
    graph_message_id: str = "",
    internet_message_id: str = "",
    conversation_id: str = "",
    sent_at: datetime | None = None,
) -> dict:
    sent = sent_at or utc_now()
    conn = _get_conn()
    try:
        conn.execute(
            """UPDATE quote_requests
               SET status=CASE
                       WHEN status IN ('complete','cancelled','stale',
                                       'needs_review','received_partial')
                       THEN status ELSE 'waiting' END,
                   sent_at=COALESCE(sent_at, ?),
                   outlook_message_id=CASE WHEN ? != '' THEN ? ELSE outlook_message_id END,
                   internet_message_id=CASE WHEN ? != '' THEN ? ELSE internet_message_id END,
                   conversation_id=CASE WHEN ? != '' THEN ? ELSE conversation_id END,
                   last_error=''
               WHERE id=?""",
            (
                sent.isoformat(),
                graph_message_id,
                graph_message_id,
                internet_message_id,
                internet_message_id,
                conversation_id,
                conversation_id,
                request_id,
            ),
        )
        conn.commit()
        row = conn.execute(
            "SELECT * FROM quote_requests WHERE id=?", (request_id,)
        ).fetchone()
    finally:
        conn.close()
    schedule_followups(request_id, sent)
    return dict(row) if row else {}


def mark_request_accepted(
    request_id: int,
    *,
    sent_at: datetime | None = None,
) -> dict:
    sent = sent_at or utc_now()
    conn = _get_conn()
    try:
        conn.execute(
            """UPDATE quote_requests
               SET status='sent', sent_at=COALESCE(sent_at, ?), last_error=''
               WHERE id=? AND status='sending'""",
            (sent.isoformat(), request_id),
        )
        conn.commit()
        row = conn.execute(
            "SELECT * FROM quote_requests WHERE id=?",
            (request_id,),
        ).fetchone()
    finally:
        conn.close()
    schedule_followups(request_id, sent)
    return dict(row) if row else {}


def claim_request_send(request_id: int) -> bool:
    conn = _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        request = conn.execute(
            "SELECT * FROM quote_requests WHERE id=?",
            (request_id,),
        ).fetchone()
        if not request:
            conn.rollback()
            return False
        job_row = conn.execute(
            "SELECT * FROM jobs WHERE id=?",
            (request["job_id"],),
        ).fetchone()
        materials = [
            dict(row)
            for row in conn.execute(
                "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
                (request["job_id"],),
            ).fetchall()
        ]
        request_state = {
            **dict(request),
            "material_snapshot": _json_loads(
                request["material_snapshot_json"],
                [],
            ),
        }
        job_state = {**dict(job_row), "materials": materials} if job_row else {}
        if (
            not job_row
            or _request_is_stale(request_state, job_state)
            or _request_has_priced_pending_materials(conn, request_id)
        ):
            conn.execute(
                """UPDATE quote_requests
                   SET status='stale',
                       last_error='Bid changed before the email was sent.'
                   WHERE id=?""",
                (request_id,),
            )
            conn.commit()
            return False
        cursor = conn.execute(
            """UPDATE quote_requests
               SET status='sending', last_error=''
               WHERE id=? AND status IN ('approved','send_failed')""",
            (request_id,),
        )
        conn.commit()
        return cursor.rowcount == 1
    finally:
        conn.close()


def record_request_send_token(request_id: int, send_token: str) -> None:
    token = str(send_token or "").strip()
    if not token:
        raise ValueError("Quote send token is required.")
    conn = _get_conn()
    try:
        cursor = conn.execute(
            """UPDATE quote_requests
               SET send_token=?
               WHERE id=? AND status='sending'""",
            (token, request_id),
        )
        conn.commit()
        if cursor.rowcount != 1:
            raise ValueError("Quote request is no longer waiting to send.")
    finally:
        conn.close()


@contextmanager
def request_send_guard(request_id: int):
    conn = _get_conn()
    sendable = False
    try:
        conn.execute("BEGIN IMMEDIATE")
        request = conn.execute(
            "SELECT * FROM quote_requests WHERE id=?",
            (request_id,),
        ).fetchone()
        job_row = (
            conn.execute(
                "SELECT * FROM jobs WHERE id=?",
                (request["job_id"],),
            ).fetchone()
            if request
            else None
        )
        materials = (
            [
                dict(row)
                for row in conn.execute(
                    "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
                    (request["job_id"],),
                ).fetchall()
            ]
            if request
            else []
        )
        if request and job_row and request["status"] == "sending":
            request_state = {
                **dict(request),
                "material_snapshot": _json_loads(
                    request["material_snapshot_json"],
                    [],
                ),
            }
            sendable = (
                not _request_is_stale(
                    request_state,
                    {**dict(job_row), "materials": materials},
                )
                and not _request_has_priced_pending_materials(conn, request_id)
            )
        if not sendable and request and request["status"] == "sending":
            conn.execute(
                """UPDATE quote_requests
                   SET status='stale',
                       last_error='Bid changed before the email was sent.'
                   WHERE id=?""",
                (request_id,),
            )
            conn.commit()
            yield False
            return
        yield sendable
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def mark_request_send_uncertain(request_id: int, error: str) -> None:
    conn = _get_conn()
    try:
        conn.execute(
            """UPDATE quote_requests
               SET status='send_uncertain', last_error=?
               WHERE id=? AND status='sending'""",
            (str(error)[:1000], request_id),
        )
        conn.commit()
    finally:
        conn.close()


def record_sent_request_evidence(
    request_id: int,
    *,
    sender_email: str,
    sent_at: str,
    raw_message: bytes | None = None,
) -> dict:
    request = get_request(request_id)
    if not request:
        raise ValueError("Quote request not found.")
    if not raw_message:
        return {"verified": False, "reason": "Microsoft MIME proof is pending."}
    message = EmailMessage()
    message["From"] = normalize_email(sender_email)
    message["To"] = normalize_email(request.get("vendor_email"))
    message["Subject"] = str(request.get("subject") or "")
    message["Date"] = str(sent_at or iso_now())
    message.set_content(str(request.get("request_text") or ""))
    raw = raw_message
    manifest = _write_artifact(
        raw,
        filename=f"quote-request-{request_id}.eml",
        job_id=int(request["job_id"]),
        artifact_kind="quote_email_sent",
    )
    conn = _get_conn()
    try:
        conn.execute(
            """UPDATE quote_requests
               SET body_hash=?, sent_artifact_path=?, sent_artifact_hash=?
               WHERE id=?""",
            (
                _sha256_bytes(str(request.get("request_text") or "").encode("utf-8")),
                manifest["artifact_path"],
                manifest["file_hash"],
                request_id,
            ),
        )
        conn.commit()
    finally:
        conn.close()
    return manifest


def record_followup_evidence(
    event_id: int,
    *,
    sender_email: str,
    recipient_email: str,
    subject: str,
    body: str,
    sent_at: str,
    raw_message: bytes | None = None,
) -> dict:
    conn = _get_conn()
    try:
        row = conn.execute(
            """SELECT qfe.quote_request_id, qr.job_id
               FROM quote_followup_events qfe
               JOIN quote_requests qr ON qr.id=qfe.quote_request_id
               WHERE qfe.id=?""",
            (event_id,),
        ).fetchone()
    finally:
        conn.close()
    if not row:
        raise ValueError("Follow-up event not found.")
    if not raw_message:
        return {"verified": False, "reason": "Microsoft MIME proof is pending."}
    message = EmailMessage()
    message["From"] = normalize_email(sender_email)
    message["To"] = normalize_email(recipient_email)
    message["Subject"] = str(subject or "")
    message["Date"] = str(sent_at or iso_now())
    message.set_content(str(body or ""))
    manifest = _write_artifact(
        raw_message,
        filename=f"quote-followup-{event_id}.eml",
        job_id=int(row["job_id"]),
        artifact_kind="quote_email_sent",
    )
    conn = _get_conn()
    try:
        conn.execute(
            """UPDATE quote_followup_events
               SET body_hash=?, artifact_path=?, artifact_hash=?
               WHERE id=?""",
            (
                _sha256_bytes(str(body or "").encode("utf-8")),
                manifest["artifact_path"],
                manifest["file_hash"],
                event_id,
            ),
        )
        conn.commit()
    finally:
        conn.close()
    return manifest


def mark_request_send_failed(request_id: int, error: str) -> None:
    conn = _get_conn()
    try:
        conn.execute(
            """UPDATE quote_requests
               SET status='send_failed', last_error=?
               WHERE id=? AND status IN ('approved','sending','send_failed')""",
            (str(error)[:1000], request_id),
        )
        conn.commit()
    finally:
        conn.close()


def get_request(request_id: int) -> dict | None:
    conn = _get_conn()
    try:
        row = conn.execute(
            "SELECT * FROM quote_requests WHERE id=?", (request_id,)
        ).fetchone()
        if not row:
            return None
        result = dict(row)
        if str(result.get("status") or "").lower() == "received":
            result["status"] = "complete"
        result["material_snapshot"] = _json_loads(
            result.get("material_snapshot_json"), []
        )
        result["material_ids"] = _json_loads(result.get("material_ids"), [])
        return result
    finally:
        conn.close()


def _request_is_stale(request: dict, job: dict) -> bool:
    source_fingerprint = str(request.get("source_fingerprint") or "")
    if source_fingerprint and source_fingerprint != _job_source_fingerprint(job):
        return True
    current = {
        int(material["id"]): material
        for material in job.get("materials") or []
        if material.get("id") is not None
    }
    snapshot = request.get("material_snapshot") or _json_loads(
        request.get("material_snapshot_json"), []
    )
    selected = []
    for item in snapshot:
        try:
            selected.append(current[int(item.get("id"))])
        except (KeyError, TypeError, ValueError):
            return True
    return snapshot_hash(selected) != str(request.get("material_snapshot_hash") or "")


def _request_has_priced_pending_materials(conn, request_id: int) -> bool:
    return bool(
        conn.execute(
            """SELECT 1
               FROM quote_request_materials qrm
               JOIN job_materials jm ON jm.id=qrm.material_id
               WHERE qrm.quote_request_id=?
                 AND qrm.status='requested'
                 AND COALESCE(jm.unit_price, 0)>0
               LIMIT 1""",
            (request_id,),
        ).fetchone()
    )


def _request_status(request: dict, job: dict) -> str:
    stored = str(request.get("status") or "draft").lower()
    if stored == "received":
        return "complete"
    if stored in {
        "cancelled",
        "complete",
        "send_failed",
        "needs_review",
        "stale",
    }:
        return stored
    if _request_is_stale(request, job):
        return "stale"
    if stored == "sent":
        return "sent"
    if stored in {"waiting", "approved", "overdue", "received_partial"}:
        sent = _parse_iso(request.get("sent_at"))
        if sent and utc_now() >= add_business_days(sent, 3):
            return "overdue"
        return "waiting" if sent else stored
    return stored


def _request_workflow_state(request: dict, job: dict) -> dict:
    stored = str(request.get("status") or "draft").lower()
    normalized = "complete" if stored == "received" else stored
    normalized_request = {**request, "status": normalized}
    if normalized in {"complete", "cancelled"}:
        stale = False
    elif normalized == "stale":
        stale = True
    else:
        stale = _request_is_stale(normalized_request, job)
    return {
        "status": "stale" if stale else _request_status(normalized_request, job),
        "stale": stale,
    }


def mark_request_stale(request_id: int) -> None:
    conn = _get_conn()
    try:
        conn.execute(
            """UPDATE quote_requests
               SET status='stale'
               WHERE id=? AND status NOT IN ('complete','cancelled','stale')""",
            (request_id,),
        )
        conn.execute(
            """UPDATE quote_followup_events
               SET status='cancelled', error='Bid materials changed after approval.'
               WHERE quote_request_id=? AND status='scheduled'""",
            (request_id,),
        )
        conn.commit()
    finally:
        conn.close()


def _followups_for_request(request_id: int) -> list[dict]:
    conn = _get_conn()
    try:
        return [
            dict(row)
            for row in conn.execute(
                """SELECT * FROM quote_followup_events
                   WHERE quote_request_id=? ORDER BY followup_number""",
                (request_id,),
            ).fetchall()
        ]
    finally:
        conn.close()


def _request_materials(request_id: int) -> list[dict]:
    conn = _get_conn()
    try:
        return [
            dict(row)
            for row in conn.execute(
                """SELECT * FROM quote_request_materials
                   WHERE quote_request_id=? ORDER BY id""",
                (request_id,),
            ).fetchall()
        ]
    finally:
        conn.close()


def _messages_for_request(request_id: int) -> list[dict]:
    conn = _get_conn()
    try:
        messages = [
            dict(row)
            for row in conn.execute(
                """SELECT id, direction, sender_email, subject, received_at,
                          sent_at, raw_artifact_path, raw_hash,
                          attachment_manifest_json, match_status, match_method
                   FROM quote_email_messages
                   WHERE matched_request_id=?
                   ORDER BY COALESCE(received_at, sent_at, created_at), id""",
                (request_id,),
            ).fetchall()
        ]
        for message in messages:
            message["attachments"] = _json_loads(
                message.pop("attachment_manifest_json", "[]"),
                [],
            )
        return messages
    finally:
        conn.close()


def quote_review(
    job_id: int | None = None,
    mailbox_email: str | None = None,
) -> dict:
    conn = _get_conn()
    try:
        message_sql = """
            SELECT qem.*, j.project_name AS matched_job_name,
                   j.slug AS matched_job_slug
            FROM quote_email_messages qem
            LEFT JOIN jobs j ON j.id=qem.matched_job_id
            WHERE qem.match_status IN ('needs_review','assigning')
        """
        message_params: list[Any] = []
        if mailbox_email:
            message_sql += " AND lower(qem.mailbox_email)=?"
            message_params.append(normalize_email(mailbox_email))
        if job_id is not None:
            message_sql += """ AND (
                qem.matched_job_id=?
                OR EXISTS (
                    SELECT 1 FROM quote_match_candidates qmc
                    WHERE qmc.message_id=qem.id AND qmc.job_id=?
                      AND qmc.status='candidate'
                )
            )"""
            message_params.extend((job_id, job_id))
        message_sql += " ORDER BY qem.received_at DESC, qem.id DESC"
        messages = [
            dict(row) for row in conn.execute(message_sql, tuple(message_params))
        ]
        for message in messages:
            assigning = str(message.get("match_status") or "") == "assigning"
            started_at = _parse_iso(message.get("assignment_started_at"))
            retry_at = (
                started_at + timedelta(seconds=ASSIGNMENT_LEASE_SECONDS)
                if assigning and started_at
                else None
            )
            message["assignment_in_progress"] = assigning
            message["assignment_retry_at"] = (
                retry_at.isoformat() if retry_at else None
            )
            message["assignment_retry_available"] = bool(
                assigning and (retry_at is None or retry_at <= utc_now())
            )
            message.pop("assignment_token", None)
            message["evidence"] = _json_loads(message.get("evidence_json"), [])
            message["attachments"] = _json_loads(
                message.get("attachment_manifest_json"), []
            )
            message["candidates"] = [
                {
                    **dict(candidate),
                    "evidence": _json_loads(candidate["evidence_json"], []),
                }
                for candidate in conn.execute(
                    """SELECT qmc.*, j.project_name, j.slug, qr.vendor_name
                       FROM quote_match_candidates qmc
                       JOIN jobs j ON j.id=qmc.job_id
                       LEFT JOIN quote_requests qr ON qr.id=qmc.quote_request_id
                       WHERE qmc.message_id=? AND qmc.status='candidate'
                       ORDER BY j.project_name""",
                    (message["id"],),
                ).fetchall()
            ]

        price_sql = """
            SELECT qpm.*, j.project_name, j.slug,
                   jm.description AS material_description,
                   COALESCE(jm.unit_price, 0) AS current_price,
                   COALESCE(jm.price_source, '') AS current_price_source
            FROM quote_price_matches qpm
            JOIN jobs j ON j.id=qpm.job_id
            JOIN quote_requests qr ON qr.id=qpm.quote_request_id
            LEFT JOIN quote_email_messages qem ON qem.id=qpm.message_id
            LEFT JOIN job_materials jm ON jm.id=qpm.material_id
            WHERE qpm.status='needs_review'
              AND (
                  qpm.message_id IS NULL
                  OR (
                      qem.match_status='matched'
                      AND qem.matched_job_id=qpm.job_id
                      AND qem.matched_request_id=qpm.quote_request_id
                  )
              )
              AND lower(qr.status) NOT IN ('cancelled','complete','received','stale')
        """
        price_params: list[Any] = []
        if mailbox_email:
            price_sql += """ AND EXISTS (
                SELECT 1 FROM quote_email_messages qem
                WHERE qem.id=qpm.message_id AND lower(qem.mailbox_email)=?
            )"""
            price_params.append(normalize_email(mailbox_email))
        if job_id is not None:
            price_sql += " AND qpm.job_id=?"
            price_params.append(job_id)
        price_sql += " ORDER BY ABS((qpm.quote_price-qpm.accepted_price_before)*qpm.quantity) DESC, qpm.id DESC"
        price_matches = [
            dict(row) for row in conn.execute(price_sql, tuple(price_params))
        ]
        for match in price_matches:
            match["dollar_impact"] = round(
                (float(match.get("quote_price") or 0) - float(match.get("accepted_price_before") or 0))
                * float(match.get("quantity") or 0),
                2,
            )
        return {"messages": messages, "price_matches": price_matches}
    finally:
        conn.close()


def quote_workflow(
    job_id: int,
    outlook_status: dict | None = None,
    mailbox_email: str | None = None,
) -> dict:
    job = load_job(job_id)
    if not job:
        raise ValueError("Job not found")
    plan = build_quote_plan(job_id)
    requests = []
    counts: dict[str, int] = {}
    for request in list_quote_requests(job_id):
        request_mailbox = normalize_email(request.get("mailbox_email"))
        if (
            mailbox_email
            and request_mailbox
            and request_mailbox != normalize_email(mailbox_email)
        ):
            continue
        request["material_snapshot"] = request.get("material_snapshot") or _json_loads(
            request.get("material_snapshot_json"), []
        )
        request["materials"] = _request_materials(request["id"])
        request["followups"] = _followups_for_request(request["id"])
        request["messages"] = _messages_for_request(request["id"])
        workflow_state = _request_workflow_state(request, job)
        request["stale"] = workflow_state["stale"]
        if request["stale"] and request.get("status") not in {
            "complete",
            "cancelled",
            "stale",
            "received",
        }:
            mark_request_stale(int(request["id"]))
        request["status"] = workflow_state["status"]
        counts[request["status"]] = counts.get(request["status"], 0) + 1
        request.pop("send_token", None)
        requests.append(request)
    review = quote_review(job_id, mailbox_email)
    return {
        "job_id": job_id,
        "source_fingerprint": plan["source_fingerprint"],
        "summary": {
            "unpriced_count": plan["unpriced_count"],
            "ready_group_count": plan["ready_group_count"],
            "request_count": len(requests),
            "needs_matching_count": len(review["messages"]),
            "price_review_count": len(review["price_matches"]),
            "statuses": counts,
        },
        "groups": plan["groups"],
        "price_history": plan["price_history"],
        "requests": requests,
        "review": review,
        "has_quote_activity": bool(
            any(
                str(request.get("status") or "").lower()
                not in {"complete", "cancelled"}
                for request in requests
            )
            or review["messages"]
            or review["price_matches"]
        ),
        "outlook": outlook_status or {"configured": False, "connected": False},
        "matching_engine": plan["matching_engine"],
    }


def _message_headers(message: dict) -> dict[str, str]:
    raw = message.get("headers") or {}
    if isinstance(raw, list):
        return {
            str(item.get("name") or "").lower(): str(item.get("value") or "")
            for item in raw
            if isinstance(item, dict)
        }
    return {str(key).lower(): str(value) for key, value in raw.items()}


def _open_request_candidates() -> list[dict]:
    conn = _get_conn()
    try:
        rows = conn.execute(
            """SELECT qr.*, j.project_name, j.slug, j.gc_name, j.address,
                      j.city, j.state, j.zip, v.contact_email
               FROM quote_requests qr
               JOIN jobs j ON j.id=qr.job_id
               LEFT JOIN vendors v ON v.id=qr.vendor_id
               WHERE qr.status IN ('approved','sending','send_uncertain',
                                   'sent','waiting','overdue',
                                   'received_partial','needs_review','send_failed',
                                   'stale')
               ORDER BY qr.created_at DESC"""
        ).fetchall()
        result = []
        for row in rows:
            request = dict(row)
            request["material_snapshot"] = _json_loads(
                request.get("material_snapshot_json"), []
            )
            result.append(request)
        return result
    finally:
        conn.close()


def _exact_job_evidence(request: dict, searchable: str) -> list[dict]:
    evidence = []
    fields = (
        ("project_name", request.get("project_name")),
        ("address", request.get("address")),
        ("gc_name", request.get("gc_name")),
        ("job_slug", request.get("slug")),
    )
    for kind, value in fields:
        normalized = normalize_text(value)
        if len(normalized) >= 5 and contains_exact_fact(searchable, value):
            evidence.append({"type": kind, "value": str(value)})
    for material in request.get("material_snapshot") or []:
        code = normalize_item_code(material.get("item_code"))
        if len(code) >= 3 and contains_exact_fact(
            searchable,
            material.get("item_code"),
        ):
            evidence.append(
                {"type": "material_code", "value": material.get("item_code") or ""}
            )
    unique = []
    seen = set()
    for item in evidence:
        key = (item["type"], normalize_text(item["value"]))
        if key not in seen:
            seen.add(key)
            unique.append(item)
    return unique


def _sender_matches_request(sender: str, request: dict) -> bool:
    address = normalize_email(sender)
    if not address:
        return False
    approved = {
        normalize_email(request.get("vendor_email")),
        normalize_email(request.get("contact_email")),
    } - {""}
    if address in approved:
        return True
    domain = email_domain(address)
    approved_domains = {email_domain(value) for value in approved} - {""}
    return bool(
        domain
        and domain not in GENERIC_EMAIL_DOMAINS
        and domain in approved_domains
    )


def match_message_to_requests(message: dict, requests: list[dict]) -> dict:
    mailbox_email = normalize_email(message.get("mailbox_email"))
    if mailbox_email:
        requests = [
            request
            for request in requests
            if normalize_email(request.get("mailbox_email")) == mailbox_email
        ]
    headers = _message_headers(message)
    sender = normalize_email(message.get("sender_email"))
    references = " ".join(
        value
        for key, value in headers.items()
        if key in {"in-reply-to", "references"}
    )
    reference_ids = set(re.findall(r"<[^>]+>", references))
    if not reference_ids:
        reference_ids = {value for value in references.split() if value}
    conversation_id = str(message.get("conversation_id") or "")
    direct = []
    for request in requests:
        internet_id = str(request.get("internet_message_id") or "")
        if internet_id and internet_id in reference_ids:
            direct.append(
                (
                    request,
                    [{"type": "reply_header", "value": internet_id}],
                    "direct_reply",
                )
            )
    if len(direct) == 1:
        request, evidence, method = direct[0]
        sender_verified = _sender_matches_request(sender, request)
        evidence.append(
            {
                "type": "vendor_sender" if sender_verified else "unverified_sender",
                "value": sender,
            }
        )
        return {
            "status": "matched",
            "job_id": request["job_id"],
            "request_id": request["id"],
            "method": method if sender_verified else "direct_reply_unverified_sender",
            "evidence": evidence,
            "candidates": [],
        }
    conversation = [
        request
        for request in requests
        if conversation_id
        and str(request.get("conversation_id") or "") == conversation_id
    ]
    if len(conversation) == 1:
        request = conversation[0]
        sender_verified = _sender_matches_request(sender, request)
        return {
            "status": "matched",
            "job_id": request["job_id"],
            "request_id": request["id"],
            "method": (
                "conversation"
                if sender_verified
                else "conversation_unverified_sender"
            ),
            "evidence": [
                {"type": "conversation", "value": conversation_id},
                {
                    "type": (
                        "vendor_sender" if sender_verified else "unverified_sender"
                    ),
                    "value": sender,
                },
            ],
            "candidates": [],
        }

    sender_domain = email_domain(sender)
    attachment_names = " ".join(
        str(item.get("name") or item.get("file_name") or "")
        for item in message.get("attachments") or []
        if isinstance(item, dict)
    )
    searchable = " ".join(
        [
            str(message.get("subject") or ""),
            str(message.get("body_text") or ""),
            attachment_names,
            str(message.get("attachment_text") or ""),
        ]
    )
    candidates = []
    for request in requests:
        addresses = {
            normalize_email(request.get("vendor_email")),
            normalize_email(request.get("contact_email")),
        } - {""}
        domains = {email_domain(address) for address in addresses} - {""}
        sender_evidence = []
        if sender in addresses:
            sender_evidence.append({"type": "vendor_email", "value": sender})
        elif (
            sender_domain
            and sender_domain not in GENERIC_EMAIL_DOMAINS
            and sender_domain in domains
        ):
            sender_evidence.append(
                {"type": "vendor_domain", "value": sender_domain}
            )
        if not sender_evidence:
            continue
        job_evidence = _exact_job_evidence(request, searchable)
        candidates.append(
            {
                "job_id": request["job_id"],
                "request_id": request["id"],
                "evidence": sender_evidence + job_evidence,
                "has_job_fact": bool(job_evidence),
            }
        )

    qualified = [candidate for candidate in candidates if candidate["has_job_fact"]]
    unique_pairs = {(item["job_id"], item["request_id"]) for item in qualified}
    if len(unique_pairs) == 1:
        candidate = qualified[0]
        return {
            "status": "matched",
            "job_id": candidate["job_id"],
            "request_id": candidate["request_id"],
            "method": "standalone_exact",
            "evidence": candidate["evidence"],
            "candidates": [],
        }
    if qualified:
        return {
            "status": "needs_review",
            "job_id": None,
            "request_id": None,
            "method": "ambiguous",
            "evidence": [{"type": "sender", "value": sender}],
            "candidates": qualified,
        }
    return {
        "status": "ignored",
        "job_id": None,
        "request_id": None,
        "method": "no_exact_job_fact" if candidates else "no_open_request",
        "evidence": [],
        "candidates": [],
    }


def _stored_message_match(conn, row) -> dict:
    return {
        "status": str(row["match_status"] or ""),
        "job_id": row["matched_job_id"],
        "request_id": row["matched_request_id"],
        "method": str(row["match_method"] or ""),
        "evidence": _json_loads(row["evidence_json"], []),
        "candidates": [
            {
                "job_id": candidate["job_id"],
                "request_id": candidate["quote_request_id"],
                "evidence": _json_loads(candidate["evidence_json"], []),
            }
            for candidate in conn.execute(
                """SELECT job_id, quote_request_id, evidence_json
                   FROM quote_match_candidates
                   WHERE message_id=? AND status='candidate'
                   ORDER BY id""",
                (row["id"],),
            ).fetchall()
        ],
    }


def _persist_message(
    message: dict,
    match: dict,
) -> tuple[int, bool, dict]:
    graph_id = str(message.get("graph_message_id") or "").strip()
    if not graph_id:
        raise ValueError("graph_message_id is required")
    raw_bytes = message.get("raw_bytes")
    raw_artifact = {"artifact_path": "", "file_hash": ""}
    if isinstance(raw_bytes, (bytes, bytearray)) and raw_bytes:
        raw_artifact = _write_artifact(
            bytes(raw_bytes),
            filename=f"{graph_id}.eml",
            artifact_kind="quote_email",
        )
    attachment_manifest = []
    for attachment in message.get("attachments") or []:
        data = attachment.get("data") if isinstance(attachment, dict) else None
        if not isinstance(data, (bytes, bytearray)):
            if isinstance(attachment, dict):
                attachment_manifest.append(
                    {
                        "file_name": str(attachment.get("name") or "attachment"),
                        "file_hash": "",
                        "file_size": int(attachment.get("size") or 0),
                        "artifact_path": "",
                        "artifact_kind": "quote_email_attachment",
                        "error": str(attachment.get("error") or "Attachment was not saved."),
                    }
                )
            continue
        attachment_manifest.append(
            _write_artifact(
                bytes(data),
                filename=str(attachment.get("name") or "attachment"),
                artifact_kind="quote_email_attachment",
            )
        )
    conn = _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        existing = conn.execute(
            "SELECT * FROM quote_email_messages WHERE graph_message_id=?",
            (graph_id,),
        ).fetchone()
        if existing:
            existing_status = str(existing["match_status"] or "")
            existing_method = str(existing["match_method"] or "")
            can_upgrade_match = (
                match.get("status") == "matched"
                and existing_status in {"ignored", "needs_review"}
                and existing_method not in {"manual", "manual_ignore"}
            )
            if can_upgrade_match:
                now = iso_now()
                evidence = list(match.get("evidence") or [])
                evidence.append(
                    {
                        "type": "automatic_rematch",
                        "value": "Sent email proof became available.",
                    }
                )
                match = {**match, "evidence": evidence}
                conn.execute(
                    """UPDATE quote_email_messages
                       SET match_status='matched', match_method=?,
                           matched_job_id=?, matched_request_id=?,
                           evidence_json=?, processed_at=NULL
                       WHERE id=?""",
                    (
                        match.get("method") or "",
                        match.get("job_id"),
                        match.get("request_id"),
                        json.dumps(evidence),
                        existing["id"],
                    ),
                )
                conn.execute(
                    """UPDATE quote_match_candidates
                       SET status='resolved',
                           decision='automatic_receipt_match',
                           reviewer_name='System', resolved_at=?
                       WHERE message_id=? AND status='candidate'""",
                    (now, existing["id"]),
                )
                conn.commit()
                return int(existing["id"]), False, match
            result = (
                int(existing["id"]),
                bool(existing["processed_at"]),
                _stored_message_match(conn, existing),
            )
            conn.commit()
            return result
        now = iso_now()
        cursor = conn.execute(
            """INSERT INTO quote_email_messages
               (graph_message_id, internet_message_id, conversation_id,
                mailbox_email, direction,
                sender_email, recipients_json, subject, body_text, received_at,
                sent_at, raw_artifact_path, raw_hash, attachment_manifest_json,
                match_status, match_method, matched_job_id, matched_request_id,
                evidence_json, created_at)
               VALUES (?, ?, ?, ?, 'incoming', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                graph_id,
                str(message.get("internet_message_id") or ""),
                str(message.get("conversation_id") or ""),
                normalize_email(message.get("mailbox_email")),
                normalize_email(message.get("sender_email")),
                json.dumps(message.get("recipients") or []),
                str(message.get("subject") or ""),
                str(message.get("body_text") or ""),
                message.get("received_at") or now,
                message.get("sent_at"),
                raw_artifact["artifact_path"],
                raw_artifact["file_hash"],
                json.dumps(
                    [
                        {
                            key: value
                            for key, value in item.items()
                            if key != "absolute_path"
                        }
                        for item in attachment_manifest
                    ]
                ),
                match["status"],
                match["method"],
                match.get("job_id"),
                match.get("request_id"),
                json.dumps(match.get("evidence") or []),
                now,
            ),
        )
        message_id = int(cursor.lastrowid)
        for candidate in match.get("candidates") or []:
            conn.execute(
                """INSERT OR IGNORE INTO quote_match_candidates
                   (message_id, job_id, quote_request_id, match_method,
                    evidence_json, created_at)
                   VALUES (?, ?, ?, 'standalone_candidate', ?, ?)""",
                (
                    message_id,
                    candidate["job_id"],
                    candidate["request_id"],
                    json.dumps(candidate.get("evidence") or []),
                    now,
                ),
            )
        conn.commit()
        return message_id, False, match
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _header_key(value: Any) -> str:
    text = str(value or "").lower().replace("#", " number ")
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", text)).strip("_")


def _positive_price(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return round(number, 4) if math.isfinite(number) and number > 0 else None
    text = str(value or "").strip()
    match = re.fullmatch(
        r"(?i)(?:USD\s*)?\$?\s*"
        r"(?P<number>(?:(?:\d{1,3}(?:,\d{3})+|\d+)"
        r"(?:\.\d{1,4})?|\.\d{1,4}))",
        text,
    )
    if not match:
        return None
    try:
        number = float(match.group("number").replace(",", ""))
    except (TypeError, ValueError):
        return None
    return round(number, 4) if math.isfinite(number) and number > 0 else None


def _product_name_with_code(code: str, description: str) -> str:
    if code and description and not contains_exact_fact(description, code):
        return f"{code} - {description}"
    return description or code


def _products_from_rows(rows: list[list[Any]]) -> list[dict]:
    header_index = None
    mapping: dict[int, str] = {}
    for index, row in enumerate(rows[:30]):
        candidate = {
            column: HEADER_ALIASES.get(_header_key(value))
            for column, value in enumerate(row)
        }
        canonical = [value for value in candidate.values() if value]
        values = set(canonical)
        possible_unit_prices = sum(
            value in {"unit_price", "ambiguous_price"}
            for value in canonical
        )
        if possible_unit_prices > 1:
            raise ValueError(
                "Quote has more than one possible unit-price column."
            )
        if "unit_price" in values and ({"item_code", "description"} & values):
            header_index = index
            mapping = {column: value for column, value in candidate.items() if value}
            break
    if header_index is None:
        return []
    products = []
    for row in rows[header_index + 1 :]:
        mapped = {
            key: (row[column] if column < len(row) else "")
            for column, key in mapping.items()
        }
        price = _positive_price(mapped.get("unit_price"))
        code = str(mapped.get("item_code") or "").strip()
        description = str(mapped.get("description") or "").strip()
        if price is None or not (code or description):
            continue
        products.append(
            {
                "item_code": code,
                "product_name": _product_name_with_code(code, description),
                "description": description,
                "vendor": str(mapped.get("vendor") or "").strip(),
                "unit_price": price,
                "unit": normalize_quote_unit(mapped.get("unit")),
                "source_unit": _exact_unit_text(mapped.get("unit")),
                "freight": mapped.get("freight") or "",
                "lead_time": str(mapped.get("lead_time") or "").strip(),
                "notes": str(mapped.get("notes") or "").strip(),
            }
        )
    return products


def _products_from_text(text: str) -> list[dict]:
    products = []
    line_pattern = re.compile(
        r"(?P<code>[A-Za-z][A-Za-z0-9._/-]{2,})"
        r"(?P<description>.*?)\$(?P<price>\d{1,6}(?:,\d{3})*(?:\.\d{1,4})?)"
        r"\s*(?:/|per\s+)?(?P<unit>SF|SY|LF|EA|PC|ROLL|BOX|CTN)\b",
        re.IGNORECASE,
    )
    for line in text.splitlines():
        compact = " ".join(line.split())
        match = line_pattern.search(compact)
        if not match:
            continue
        code = match.group("code")
        if not (re.search(r"\d", code) or "-" in code):
            continue
        price = _positive_price(match.group("price"))
        if price is None:
            continue
        products.append(
            {
                "item_code": code,
                "product_name": _product_name_with_code(
                    code,
                    match.group("description").strip(" -,:"),
                ),
                "description": match.group("description").strip(" -,:"),
                "vendor": "",
                "unit_price": price,
                "unit": normalize_quote_unit(match.group("unit")),
                "source_unit": _exact_unit_text(match.group("unit")),
                "freight": "",
                "lead_time": "",
                "notes": "",
            }
        )

    item_label_pattern = re.compile(
        r"(?im)^\s*(?:ITEM\s*(?:CODE|NUMBER|#)|SKU|PRODUCT\s+CODE)"
        r"\s*[:#-]\s*(?P<code>[A-Za-z][A-Za-z0-9._/-]{2,})\s*$"
    )
    item_labels = list(item_label_pattern.finditer(text))
    vendor_match = re.search(
        r"(?im)^\s*VENDOR\s*[:#-]\s*(?P<vendor>[^\r\n]{2,120})\s*$",
        text,
    )
    for index, item_match in enumerate(item_labels):
        block_end = (
            item_labels[index + 1].start()
            if index + 1 < len(item_labels)
            else len(text)
        )
        block = text[item_match.start():block_end]
        price_match = re.search(
            r"(?im)^\s*(?:UNIT\s*(?:PRICE|COST)|PRICE\s+PER\s+UNIT)"
            r"\s*[:#-]?\s*\$?\s*"
            r"(?P<price>\d{1,6}(?:,\d{3})*(?:\.\d{1,4})?)"
            r"\s*(?:(?:/|PER)\s*(?P<unit>SF|SY|LF|EA|PC|ROLL|BOX|CTN))?\s*$",
            block,
        )
        if not price_match:
            continue
        source_unit = price_match.group("unit")
        if not source_unit:
            unit_match = re.search(
                r"(?im)^\s*(?:UOM|UNIT(?:\s+OF\s+MEASURE)?)"
                r"\s*[:#-]\s*(?P<unit>SF|SY|LF|EA|PC|ROLL|BOX|CTN)\s*$",
                block,
            )
            source_unit = unit_match.group("unit") if unit_match else ""
        price = _positive_price(price_match.group("price"))
        code = item_match.group("code")
        if (
            price is None
            or not source_unit
            or not (re.search(r"\d", code) or "-" in code)
        ):
            continue
        description_match = re.search(
            r"(?im)^\s*(?:PRODUCT|ITEM\s+DESCRIPTION|DESCRIPTION)"
            r"\s*[:#-]\s*(?P<description>[^\r\n]{2,240})\s*$",
            block,
        )
        description = (
            description_match.group("description").strip()
            if description_match
            else code
        )
        products.append(
            {
                "item_code": code,
                "product_name": _product_name_with_code(code, description),
                "description": description,
                "vendor": (
                    vendor_match.group("vendor").strip()
                    if vendor_match
                    else ""
                ),
                "unit_price": price,
                "unit": normalize_quote_unit(source_unit),
                "source_unit": _exact_unit_text(source_unit),
                "freight": "",
                "lead_time": "",
                "notes": "",
            }
        )

    unique = {}
    for product in products:
        key = (
            normalize_item_code(product["item_code"]),
            product["unit"],
            _exact_unit_key(product.get("source_unit")),
            product["unit_price"],
        )
        unique[key] = product
    return list(unique.values())


def parse_quote_file_deterministic(path: str, *, depth: int = 0) -> list[dict]:
    if depth > 2:
        raise ValueError("Nested email depth exceeds the supported limit.")
    file_path = Path(path)
    if not file_path.is_file():
        raise ValueError("Quote file is missing.")
    if file_path.stat().st_size > MAX_ATTACHMENT_BYTES:
        raise ValueError("Quote file is larger than 25 MB.")
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        with file_path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            return _products_from_rows([list(row) for row in csv.reader(handle)])
    if suffix == ".xlsx":
        import openpyxl

        workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        try:
            products = []
            for sheet in workbook.worksheets:
                rows = [
                    list(row)
                    for row in itertools.islice(
                        sheet.iter_rows(values_only=True),
                        MAX_SPREADSHEET_ROWS,
                    )
                ]
                products.extend(_products_from_rows(rows))
            return products
        finally:
            workbook.close()
    if suffix == ".txt":
        return _products_from_text(
            file_path.read_text(encoding="utf-8", errors="replace")
        )
    if suffix == ".pdf":
        text_parts = []
        with pdfplumber.open(str(file_path)) as pdf:
            for page in pdf.pages[:50]:
                page_text = page.extract_text() or ""
                if page_text:
                    text_parts.append(page_text)
            if not text_parts and pdf.pages:
                try:
                    import pytesseract
                    from pdf2image import convert_from_path
                except ImportError:
                    return []
                for image in convert_from_path(
                    str(file_path),
                    dpi=180,
                    first_page=1,
                    last_page=min(12, len(pdf.pages)),
                ):
                    try:
                        text_parts.append(pytesseract.image_to_string(image))
                    finally:
                        image.close()
        return _products_from_text("\n".join(text_parts))
    if suffix == ".eml":
        with file_path.open("rb") as handle:
            message = email.message_from_binary_file(handle, policy=policy.default)
        products = []
        body = []
        for part in message.walk():
            if part.get_content_disposition() == "attachment" and part.get_filename():
                data = part.get_payload(decode=True)
                nested_suffix = Path(part.get_filename()).suffix.lower()
                if data and nested_suffix in {".csv", ".xlsx", ".pdf", ".txt", ".eml", ".msg"}:
                    temp = ARTIFACT_ROOT / "shared" / "tmp" / f"{_sha256_bytes(data)[:12]}{nested_suffix}"
                    temp.parent.mkdir(parents=True, exist_ok=True)
                    temp.write_bytes(data)
                    nested_products = parse_quote_file_deterministic(
                        str(temp),
                        depth=depth + 1,
                    )
                    if not nested_products:
                        raise ValueError(
                            f"Attachment {part.get_filename()} has no exact quote rows."
                        )
                    products.extend(nested_products)
            elif part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True)
                if payload:
                    body.append(payload.decode("utf-8", errors="replace"))
        products.extend(_products_from_text("\n".join(body)))
        return products
    if suffix == ".msg":
        import extract_msg

        message = extract_msg.openMsg(str(file_path))
        try:
            products = _products_from_text(str(getattr(message, "body", "") or ""))
            for attachment in getattr(message, "attachments", []) or []:
                name = (
                    getattr(attachment, "longFilename", None)
                    or getattr(attachment, "shortFilename", None)
                    or "attachment"
                )
                data = getattr(attachment, "data", None)
                nested_suffix = Path(str(name)).suffix.lower()
                if isinstance(data, (bytes, bytearray)) and nested_suffix in {
                    ".csv",
                    ".xlsx",
                    ".pdf",
                    ".txt",
                    ".eml",
                    ".msg",
                }:
                    temp = ARTIFACT_ROOT / "shared" / "tmp" / f"{_sha256_bytes(bytes(data))[:12]}{nested_suffix}"
                    temp.parent.mkdir(parents=True, exist_ok=True)
                    temp.write_bytes(bytes(data))
                    nested_products = parse_quote_file_deterministic(
                        str(temp),
                        depth=depth + 1,
                    )
                    if not nested_products:
                        raise ValueError(
                            f"Attachment {name} has no exact quote rows."
                        )
                    products.extend(nested_products)
            return products
        finally:
            message.close()
    raise ValueError(f"Unsupported quote file type: {suffix}")


def repair_existing_price_evidence(
    job_id: int,
    files: list[tuple[str, bytes]],
    *,
    reviewer_name: str,
) -> dict:
    job = load_job(job_id)
    if not job:
        raise ValueError("Job not found.")
    parsed_sources = []
    issues = []
    for filename, data in files:
        manifest = _write_artifact(
            data,
            filename=filename,
            job_id=job_id,
            artifact_kind="vendor_quote",
        )
        try:
            products = parse_quote_file_deterministic(manifest["absolute_path"])
        except Exception as exc:
            issues.append(
                {
                    "file_name": filename,
                    "message": f"Could not read exact prices: {str(exc)[:240]}",
                }
            )
            continue
        if not products:
            issues.append(
                {
                    "file_name": filename,
                    "message": "No explicit item code, unit, and unit price rows were found.",
                }
            )
            continue
        record_imported_file(
            job_id,
            manifest["file_name"],
            manifest["file_hash"],
            manifest["file_size"],
            source="manual_exact_repair",
            artifact_path=manifest["artifact_path"],
            artifact_kind="vendor_quote",
        )
        parsed_sources.append((manifest, products))

    repairs = []
    conflicts = []
    for material in job.get("materials") or []:
        accepted_price = float(material.get("unit_price") or 0)
        material_code = normalize_item_code(material.get("item_code"))
        material_unit = _exact_unit_text(material.get("unit"))
        if (
            accepted_price <= 0
            or not material_code
            or not material_unit
            or str(material.get("price_source") or "").lower() != "vendor_quote"
        ):
            continue
        candidates = []
        changed_prices = []
        for manifest, products in parsed_sources:
            for product in products:
                if (
                    normalize_item_code(product.get("item_code")) != material_code
                    or _exact_unit_key(
                        product.get("source_unit") or product.get("unit")
                    )
                    != _exact_unit_key(material_unit)
                ):
                    continue
                quoted_price = float(product.get("unit_price") or 0)
                if abs(quoted_price - accepted_price) <= 0.0001:
                    candidates.append((manifest, product))
                else:
                    changed_prices.append(
                        {
                            "file_name": manifest["file_name"],
                            "quoted_price": quoted_price,
                        }
                    )
        unique_candidates = {
            manifest["file_hash"]: (manifest, product)
            for manifest, product in candidates
        }
        if len(unique_candidates) == 1:
            manifest, _ = next(iter(unique_candidates.values()))
            repairs.append(
                {
                    "material_id": int(material["id"]),
                    "item_code": material.get("item_code") or "",
                    "accepted_price": accepted_price,
                    "unit": material_unit,
                    "file_name": manifest["file_name"],
                    "file_hash": manifest["file_hash"],
                }
            )
        elif len(unique_candidates) > 1 or changed_prices:
            conflicts.append(
                {
                    "material_id": int(material["id"]),
                    "item_code": material.get("item_code") or "",
                    "accepted_price": accepted_price,
                    "unit": material_unit,
                    "reason": (
                        "More than one exact proof matched."
                        if len(unique_candidates) > 1
                        else "The uploaded quote has a different price."
                    ),
                    "quoted": changed_prices,
                }
            )

    conn = _get_conn()
    applied = []
    try:
        conn.execute("BEGIN IMMEDIATE")
        for repair in repairs:
            cursor = conn.execute(
                """UPDATE job_materials
                   SET quote_source_hash=?, quote_file_name=?,
                       quote_status='quoted'
                   WHERE id=? AND job_id=?
                     AND COALESCE(unit_price, 0)=?
                     AND lower(COALESCE(price_source, ''))='vendor_quote'""",
                (
                    repair["file_hash"],
                    repair["file_name"],
                    repair["material_id"],
                    job_id,
                    repair["accepted_price"],
                ),
            )
            if cursor.rowcount == 1:
                applied.append(repair)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    if applied or conflicts or issues:
        log_activity(
            job_id,
            "vendor_quote_evidence_repair",
            (
                f"Exact-only receipt repair linked {len(applied)} material(s); "
                f"{len(conflicts) + len(issues)} need review"
            ),
            {
                "reviewer": reviewer_name or "Estimator",
                "applied": applied,
                "conflicts": conflicts,
                "issues": issues,
                "ai_calls": 0,
            },
        )
    return {
        "status": "repaired" if applied else "needs_review",
        "repaired_count": len(applied),
        "repaired": applied,
        "conflicts": conflicts,
        "issues": issues,
        "prices_changed": 0,
        "ai_calls": 0,
    }


def attachment_match_text(attachments: list[dict]) -> str:
    """Extract exact searchable facts without making any bid decision."""
    searchable = []
    supported = {".csv", ".xlsx", ".pdf", ".txt", ".eml", ".msg"}
    for attachment in attachments[:20]:
        if not isinstance(attachment, dict):
            continue
        data = attachment.get("data")
        name = str(attachment.get("name") or "attachment")
        suffix = Path(name).suffix.lower()
        if (
            suffix not in supported
            or not isinstance(data, (bytes, bytearray))
            or len(data) > MAX_ATTACHMENT_BYTES
        ):
            continue
        raw = bytes(data)
        temp = (
            ARTIFACT_ROOT
            / "shared"
            / "tmp"
            / f"{_sha256_bytes(raw)[:12]}{suffix}"
        )
        temp.parent.mkdir(parents=True, exist_ok=True)
        if not temp.exists():
            temp.write_bytes(raw)
        if suffix in {".txt", ".csv"}:
            searchable.append(raw[:1_000_000].decode("utf-8", errors="replace"))
        try:
            products = parse_quote_file_deterministic(str(temp))
        except Exception:
            products = []
        searchable.extend(
            " ".join(
                str(value or "")
                for value in (
                    product.get("item_code"),
                    product.get("product_name"),
                    product.get("description"),
                    product.get("vendor"),
                )
            )
            for product in products
        )
    return "\n".join(searchable)[:2_000_000]


def _extract_product_code(product: dict) -> str:
    explicit = normalize_item_code(product.get("item_code"))
    if explicit:
        return explicit
    name = str(product.get("product_name") or "")
    candidates = re.findall(r"\b[A-Za-z]{1,8}[-./]?\d{2,}[A-Za-z0-9.-]*\b", name)
    return normalize_item_code(candidates[0]) if len(candidates) == 1 else ""


def evaluate_price_candidate(
    product: dict,
    *,
    requested_materials: list[dict],
    current_materials: list[dict],
    match_method: str,
    request_stale: bool,
) -> dict:
    code = _extract_product_code(product)
    requested = [
        row
        for row in requested_materials
        if code and normalize_item_code(row.get("item_code")) == code
    ]
    current_by_id = {
        int(material["id"]): material
        for material in current_materials
        if material.get("id") is not None
    }
    requested_row = requested[0] if len(requested) == 1 else None
    material = (
        current_by_id.get(int(requested_row["material_id"]))
        if requested_row and requested_row.get("material_id") is not None
        else None
    )
    quote_unit = _exact_unit_text(
        product.get("source_unit") or product.get("unit")
    )
    material_unit = _exact_unit_text(material.get("unit")) if material else ""
    direct = match_method in {"direct_reply", "conversation"}
    current_price = float(material.get("unit_price") or 0) if material else 0

    if not code or not requested_row or not material:
        status = "needs_review"
        reason = "No single requested material has this exact vendor item code."
    elif request_stale:
        status = "needs_review"
        reason = "The bid changed after this quote request was approved."
    elif (
        not _exact_unit_key(quote_unit)
        or _exact_unit_key(quote_unit) != _exact_unit_key(material_unit)
    ):
        status = "needs_review"
        reason = (
            f"Unit mismatch: quote uses {quote_unit or 'unknown'}, "
            f"bid uses {material_unit or 'unknown'}."
        )
    elif not direct:
        status = "needs_review"
        reason = (
            "This was not a verified direct vendor reply and needs estimator approval."
        )
    elif current_price > 0:
        status = "needs_review"
        reason = "This material already has a price. It will not be replaced automatically."
    elif str(requested_row.get("status") or "requested") != "requested":
        status = "needs_review"
        reason = "This requested material was already resolved."
    else:
        status = "ready_to_apply"
        reason = "Verified direct reply with exact requested item code and unit."
    return {
        "code": code,
        "requested_material": requested_row,
        "material": material,
        "status": status,
        "reason": reason,
        "quote_unit": quote_unit,
        "material_unit": material_unit,
    }


def _require_assignment_owner(
    conn,
    *,
    message_id: int,
    assignment_token: str,
) -> None:
    row = conn.execute(
        """SELECT 1 FROM quote_email_messages
           WHERE id=? AND match_status='assigning' AND assignment_token=?""",
        (message_id, assignment_token),
    ).fetchone()
    if not row:
        raise AssignmentLeaseLostError(
            "A newer worker owns this email assignment. Refresh and try again."
        )


def _record_price_match(
    *,
    message_id: int,
    request: dict,
    material: dict | None,
    product: dict,
    source: dict,
    match_method: str,
    status: str,
    reason: str,
    assignment_token: str | None = None,
    conn=None,
) -> int:
    now = iso_now()
    material = material or {}
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        if assignment_token:
            _require_assignment_owner(
                conn,
                message_id=message_id,
                assignment_token=assignment_token,
            )
        existing = conn.execute(
            """SELECT id FROM quote_price_matches
               WHERE message_id=? AND quote_request_id=?
                 AND material_id IS ? AND source_hash=?
                 AND quote_price=? AND quote_unit=?""",
            (
                message_id,
                request["id"],
                material.get("id"),
                source["file_hash"],
                float(product.get("unit_price") or 0),
                _exact_unit_text(
                    product.get("source_unit") or product.get("unit")
                ),
            ),
        ).fetchone()
        if existing:
            conn.commit()
            return int(existing["id"])
        cursor = conn.execute(
            """INSERT OR IGNORE INTO quote_price_matches
               (message_id, job_id, quote_request_id, material_id, item_code,
                description, quote_price, quote_unit, material_unit,
                accepted_price_before, quantity, source_hash, source_file,
                match_method, status, reason, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                message_id,
                request["job_id"],
                request["id"],
                material.get("id"),
                product.get("item_code") or material.get("item_code") or "",
                product.get("product_name") or product.get("description") or "",
                float(product.get("unit_price") or 0),
                _exact_unit_text(product.get("source_unit") or product.get("unit")),
                _exact_unit_text(material.get("unit")),
                float(material.get("unit_price") or 0),
                float(material.get("order_qty") or material.get("installed_qty") or 0),
                source["file_hash"],
                source["file_name"],
                match_method,
                status,
                reason,
                now,
            ),
        )
        if cursor.rowcount != 1:
            raise ValueError(
                "This email price is already linked to another quote request."
            )
        conn.commit()
        return int(cursor.lastrowid or 0)
    except Exception:
        conn.rollback()
        raise
    finally:
        if owns_connection:
            conn.close()


def _apply_exact_price(
    match_id: int,
    *,
    reviewer_name: str,
    reason: str,
    automatic: bool = False,
    expected_material_id: int | None = None,
    expected_current_price: float | None = None,
    conn=None,
) -> bool:
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        match = conn.execute(
            "SELECT * FROM quote_price_matches WHERE id=?", (match_id,)
        ).fetchone()
        if not match or not match["material_id"]:
            conn.rollback()
            return False
        request = conn.execute(
            "SELECT * FROM quote_requests WHERE id=? AND job_id=?",
            (match["quote_request_id"], match["job_id"]),
        ).fetchone()
        requested_material = conn.execute(
            """SELECT * FROM quote_request_materials
               WHERE quote_request_id=? AND material_id=?""",
            (match["quote_request_id"], match["material_id"]),
        ).fetchone()
        material = conn.execute(
            "SELECT * FROM job_materials WHERE id=? AND job_id=?",
            (match["material_id"], match["job_id"]),
        ).fetchone()
        job_row = conn.execute(
            "SELECT * FROM jobs WHERE id=?", (match["job_id"],)
        ).fetchone()
        if not request or not requested_material or not material or not job_row:
            conn.rollback()
            return False
        if str(request["status"] or "").lower() in {
            "complete",
            "cancelled",
            "received",
            "stale",
        }:
            conn.rollback()
            return False
        current_materials = [
            dict(row)
            for row in conn.execute(
                "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
                (match["job_id"],),
            ).fetchall()
        ]
        job_state = {**dict(job_row), "materials": current_materials}
        request_dict = {
            **dict(request),
            "material_snapshot": _json_loads(
                request["material_snapshot_json"], []
            ),
        }
        if _request_is_stale(request_dict, job_state):
            conn.rollback()
            mark_request_stale(int(request["id"]))
            return False
        current_price = float(material["unit_price"] or 0)
        quote_price = float(match["quote_price"])
        quantity = float(material["order_qty"] or material["installed_qty"] or 0)
        expected_status = "ready_to_apply" if automatic else "needs_review"
        if str(match["status"]) != expected_status:
            conn.rollback()
            return False
        if (
            not automatic
            and expected_material_id is not None
            and int(match["material_id"] or 0) != int(expected_material_id)
        ):
            conn.rollback()
            return False
        if not automatic:
            if expected_current_price is None or not math.isclose(
                current_price,
                float(expected_current_price),
                abs_tol=0.0001,
            ):
                conn.rollback()
                return False
        if automatic and (
            current_price > 0
            or str(requested_material["status"]) != "requested"
            or normalize_item_code(match["item_code"])
            != normalize_item_code(requested_material["item_code"])
            or _exact_unit_key(match["quote_unit"])
            != _exact_unit_key(requested_material["unit"])
        ):
            conn.rollback()
            return False
        material_update = conn.execute(
            f"""UPDATE job_materials
                SET unit_price=?, extended_cost=?,
                    vendor=COALESCE(NULLIF(vendor,''), ?),
                    quote_status='quoted', price_source='vendor_quote',
                    quote_source_hash=?, quote_file_name=?
                WHERE id=? AND job_id=?
                {"AND COALESCE(unit_price, 0)<=0" if automatic else ""}""",
            (
                quote_price,
                round(quantity * quote_price, 2),
                request["vendor_name"] or "",
                match["source_hash"],
                match["source_file"],
                material["id"],
                match["job_id"],
            ),
        )
        if material_update.rowcount != 1:
            conn.rollback()
            return False
        conn.execute(
            """UPDATE quote_price_matches
               SET status='applied', decision='use_quote', reviewer_name=?,
                   reason=?, resolved_at=?
               WHERE id=?""",
            (reviewer_name, reason, iso_now(), match_id),
        )
        conn.execute(
            """UPDATE quote_request_materials
               SET status='quoted', quoted_price=?, source_message_id=?, resolved_at=?
               WHERE quote_request_id=? AND material_id=?""",
            (
                quote_price,
                match["message_id"],
                iso_now(),
                match["quote_request_id"],
                material["id"],
            ),
        )
        record_material_price_decision(
            job_id=int(match["job_id"]),
            material_id=int(material["id"]),
            item_code=str(material["item_code"] or ""),
            decision="use_quote",
            accepted_price_before=current_price,
            resolved_price=quote_price,
            material_unit=str(material["unit"] or ""),
            quote_price=quote_price,
            quote_unit=str(match["quote_unit"] or ""),
            source_hash=str(match["source_hash"]),
            source_file=str(match["source_file"] or ""),
            reason=reason,
            reviewer_name=reviewer_name,
            conn=conn,
        )
        match_data = dict(match)
        material_data = dict(material)
        request_data = dict(request)
        product = {
            "item_code": material_data.get("item_code")
            or match_data.get("item_code")
            or "",
            "product_name": (
                match_data.get("description")
                or material_data.get("description")
                or material_data.get("item_code")
                or "Vendor quote"
            ),
            "description": match_data.get("description")
            or material_data.get("description")
            or "",
            "vendor": request_data.get("vendor_name") or "",
            "unit_price": quote_price,
            "unit": match_data.get("quote_unit") or material_data.get("unit") or "",
            "quantity": quantity,
            "file_name": match_data.get("source_file") or "",
            "_source_hash": match_data.get("source_hash") or "",
            "verified": True,
            "verified_at": iso_now(),
            "verified_by": reviewer_name,
            "quote_request_id": match_data.get("quote_request_id"),
            "price_match_id": match_id,
        }
        save_quotes(int(match_data["job_id"]), [product], conn=conn)
        save_vendor_prices_from_quotes(
            int(match_data["job_id"]),
            [product],
            conn=conn,
        )
        _refresh_request_completion(
            int(match_data["quote_request_id"]),
            conn=conn,
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        if owns_connection:
            conn.close()
    return True


def _refresh_request_completion(
    request_id: int,
    conn=None,
) -> None:
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        rows = conn.execute(
            "SELECT status FROM quote_request_materials WHERE quote_request_id=?",
            (request_id,),
        ).fetchall()
        statuses = {str(row["status"]) for row in rows}
        unresolved_reviews = int(
            conn.execute(
                """SELECT COUNT(*) AS count
                   FROM quote_price_matches
                   WHERE quote_request_id=? AND status='needs_review'""",
                (request_id,),
            ).fetchone()["count"]
            or 0
        )
        if unresolved_reviews:
            conn.execute(
                """UPDATE quote_requests
                   SET status='needs_review',
                       received_at=COALESCE(received_at, ?)
                   WHERE id=? AND status NOT IN ('cancelled','complete','stale')""",
                (iso_now(), request_id),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='cancelled',
                       error='Estimator review is required before follow-up.'
                   WHERE quote_request_id=? AND status='scheduled'""",
                (request_id,),
            )
        elif rows and statuses <= {"quoted", "ignored"}:
            conn.execute(
                """UPDATE quote_requests
                   SET status='complete', completed_at=?
                   WHERE id=? AND status!='cancelled'""",
                (iso_now(), request_id),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='cancelled'
                   WHERE quote_request_id=? AND status='scheduled'""",
                (request_id,),
            )
        elif "needs_review" in statuses:
            conn.execute(
                """UPDATE quote_requests
                   SET status='needs_review',
                       received_at=COALESCE(received_at, ?)
                   WHERE id=? AND status NOT IN ('cancelled','complete','stale')""",
                (iso_now(), request_id),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='cancelled',
                       error='Estimator review is required before follow-up.'
                   WHERE quote_request_id=? AND status='scheduled'""",
                (request_id,),
            )
        elif "quoted" in statuses:
            conn.execute(
                """UPDATE quote_requests
                   SET status='received_partial', received_at=COALESCE(received_at, ?)
                   WHERE id=? AND status!='cancelled'""",
                (iso_now(), request_id),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='scheduled', error=''
                   WHERE quote_request_id=? AND status='cancelled'
                     AND error='Estimator review is required before follow-up.'""",
                (request_id,),
            )
        elif statuses == {"requested"}:
            conn.execute(
                """UPDATE quote_requests
                   SET status=CASE WHEN received_at IS NULL THEN 'waiting'
                                   ELSE 'received_partial' END
                   WHERE id=? AND status NOT IN ('cancelled','complete','stale')""",
                (request_id,),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='scheduled', error=''
                   WHERE quote_request_id=? AND status='cancelled'
                     AND error='Estimator review is required before follow-up.'""",
                (request_id,),
            )
        if owns_connection:
            conn.commit()
    finally:
        if owns_connection:
            conn.close()


def _message_price_decision(
    evaluation: dict,
    parse_errors: list[dict],
) -> tuple[str, str]:
    if parse_errors:
        return (
            "needs_review",
            (
                "One or more files in this email could not be read. "
                "No price from the email can be automatic."
            ),
        )
    return str(evaluation["status"]), str(evaluation["reason"])


def _store_no_price_result(
    conn,
    *,
    message_id: int,
    request_id: int,
    evidence_json: str,
    defer_request_state: bool,
    assignment_token: str | None = None,
) -> None:
    if defer_request_state:
        if not assignment_token:
            raise ValueError("Deferred pricing requires an assignment token.")
        _require_assignment_owner(
            conn,
            message_id=message_id,
            assignment_token=assignment_token,
        )
        updated = conn.execute(
            """UPDATE quote_email_messages
               SET evidence_json=?
               WHERE id=? AND match_status='assigning'
                 AND assignment_token=?""",
            (evidence_json, message_id, assignment_token),
        )
        if updated.rowcount != 1:
            raise AssignmentLeaseLostError(
                "A newer worker owns this email assignment. Refresh and try again."
            )
        return
    conn.execute(
        """UPDATE quote_email_messages
           SET match_status='needs_review', evidence_json=?
           WHERE id=?""",
        (evidence_json, message_id),
    )
    conn.execute(
        """UPDATE quote_requests SET status='needs_review'
           WHERE id=? AND status NOT IN
               ('complete','cancelled','received','stale')""",
        (request_id,),
    )


def _mark_material_for_price_review(
    *,
    message_id: int,
    request_id: int,
    material_id: int,
    assignment_token: str | None,
    conn=None,
) -> None:
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        if assignment_token:
            conn.execute("BEGIN IMMEDIATE")
            _require_assignment_owner(
                conn,
                message_id=message_id,
                assignment_token=assignment_token,
            )
        conn.execute(
            """UPDATE quote_request_materials
               SET status='needs_review', source_message_id=?
               WHERE quote_request_id=? AND material_id=?
                 AND status='requested'""",
            (message_id, request_id, material_id),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        if owns_connection:
            conn.close()


def _price_products_for_message(
    message_id: int,
    request: dict,
    match_method: str,
    *,
    defer_request_state: bool = False,
    assignment_token: str | None = None,
) -> dict:
    if defer_request_state and not assignment_token:
        raise ValueError("Deferred pricing requires an assignment token.")
    if defer_request_state and match_method != "manual":
        raise ValueError("Deferred pricing is allowed only for manual bid assignment.")
    conn = _get_conn()
    try:
        row = conn.execute(
            "SELECT * FROM quote_email_messages WHERE id=?", (message_id,)
        ).fetchone()
        if not row:
            return {"products": 0, "applied": 0, "needs_review": 0}
        message = dict(row)
    finally:
        conn.close()
    manifests = _json_loads(message.get("attachment_manifest_json"), [])
    raw_manifest = {
        "file_name": f"{message.get('graph_message_id')}.eml",
        "file_hash": message.get("raw_hash"),
        "artifact_path": message.get("raw_artifact_path"),
    }
    parse_errors = [
        {
            "type": "attachment_failure",
            "value": (
                f"{manifest.get('file_name') or 'attachment'}: "
                f"{manifest.get('error') or 'Attachment was not saved.'}"
            ),
        }
        for manifest in manifests
        if isinstance(manifest, dict) and manifest.get("error")
    ]
    sources = []
    for manifest in [raw_manifest, *manifests]:
        source = (
            _manifest_artifact_source(manifest)
            if defer_request_state
            else _copy_artifact_to_job(manifest, request["job_id"])
        )
        if source:
            sources.append(source)
    products_with_source = []
    raw_products_found = False
    parsed_sources: dict[str, dict] = {}
    for source_index, source in enumerate(sources):
        if source_index > 0 and raw_products_found:
            break
        try:
            products = parse_quote_file_deterministic(source["absolute_path"])
        except Exception as exc:
            products = []
            parse_errors.append(
                {
                    "type": "parse_failure",
                    "value": (
                        f"{source['file_name']}: "
                        f"{type(exc).__name__}: {str(exc)[:240]}"
                    ),
                }
            )
        if source_index == 0 and products:
            # A preserved .eml includes its attachments. Do not parse those
            # attachments a second time and create duplicate price candidates.
            raw_products_found = True
        if products and not defer_request_state:
            record_imported_file(
                request["job_id"],
                source["file_name"],
                source["file_hash"],
                source["file_size"],
                source="outlook",
                artifact_path=source["artifact_path"],
                artifact_kind="vendor_quote",
            )
        if products:
            parsed_sources[source["file_hash"]] = {
                key: source[key]
                for key in (
                    "file_name",
                    "file_hash",
                    "file_size",
                    "artifact_path",
                    "artifact_kind",
                )
            }
        for product in products:
            product["vendor"] = product.get("vendor") or request.get("vendor_name") or ""
            product["file_name"] = source["file_name"]
            product["_source_hash"] = source["file_hash"]
            products_with_source.append((product, source))

    job = load_job(request["job_id"]) or {}
    materials = job.get("materials") or []
    requested_materials = _request_materials(int(request["id"]))
    request_stale = _request_is_stale(request, job)
    applied = 0
    needs_review = 0
    for product, source in products_with_source:
        evaluation = evaluate_price_candidate(
            product,
            requested_materials=requested_materials,
            current_materials=materials,
            match_method=match_method,
            request_stale=request_stale,
        )
        material = evaluation["material"]
        status, reason = _message_price_decision(evaluation, parse_errors)
        match_id = _record_price_match(
            message_id=message_id,
            request=request,
            material=material,
            product=product,
            source=source,
            match_method=(
                "exact_requested_code_unit"
                if material
                and evaluation["quote_unit"]
                and evaluation["quote_unit"] == evaluation["material_unit"]
                else "review"
            ),
            status=status,
            reason=reason,
            assignment_token=assignment_token,
        )
        if status == "ready_to_apply" and match_id:
            applied_ok = _apply_exact_price(
                match_id,
                reviewer_name="Deterministic quote automation",
                reason=reason,
                automatic=True,
            )
            if applied_ok:
                applied += 1
            else:
                needs_review += 1
                conn = _get_conn()
                try:
                    conn.execute(
                        """UPDATE quote_price_matches
                           SET status='needs_review',
                               reason='Safety re-check stopped automatic pricing.'
                           WHERE id=? AND status='ready_to_apply'""",
                        (match_id,),
                    )
                    conn.commit()
                finally:
                    conn.close()
                if material:
                    _mark_material_for_price_review(
                        message_id=message_id,
                        request_id=int(request["id"]),
                        material_id=int(material["id"]),
                        assignment_token=assignment_token,
                    )
        else:
            needs_review += 1
            if material:
                _mark_material_for_price_review(
                    message_id=message_id,
                    request_id=int(request["id"]),
                    material_id=int(material["id"]),
                    assignment_token=assignment_token,
                )
    if not products_with_source:
        no_price_evidence = [
            *parse_errors,
            {
                "type": "parse_failure",
                "value": (
                    "No explicit price rows were found. "
                    "Free-text prices must include a dollar sign."
                ),
            },
        ]
        evidence = [
            *_json_loads(message.get("evidence_json"), []),
            *no_price_evidence,
        ]
        conn = _get_conn()
        try:
            if defer_request_state:
                conn.execute("BEGIN IMMEDIATE")
            _store_no_price_result(
                conn,
                message_id=message_id,
                request_id=int(request["id"]),
                evidence_json=json.dumps(evidence),
                defer_request_state=defer_request_state,
                assignment_token=assignment_token,
            )
            conn.commit()
        finally:
            conn.close()
    elif needs_review and not defer_request_state:
        _refresh_request_completion(int(request["id"]))
        conn = _get_conn()
        try:
            conn.execute(
                """UPDATE quote_requests
                   SET status='needs_review',
                       received_at=COALESCE(received_at, ?)
                   WHERE id=? AND status NOT IN ('complete','cancelled','stale')""",
                (iso_now(), request["id"]),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='cancelled',
                       error='Estimator review is required before follow-up.'
                   WHERE quote_request_id=? AND status='scheduled'""",
                (request["id"],),
            )
            conn.commit()
        finally:
            conn.close()
    return {
        "products": len(products_with_source),
        "applied": applied,
        "needs_review": needs_review,
        "parse_errors": parse_errors,
        "parsed_sources": list(parsed_sources.values()),
    }


def process_incoming_message(message: dict) -> dict:
    match = match_message_to_requests(message, _open_request_candidates())
    message_id, already_processed, match = _persist_message(message, match)
    if already_processed:
        return {"status": "duplicate", "message_id": message_id}
    result = {"status": match["status"], "message_id": message_id, "match": match}
    if match["status"] == "matched":
        request = get_request(int(match["request_id"]))
        if request:
            result["pricing"] = _price_products_for_message(
                message_id, request, match["method"]
            )
            conn = _get_conn()
            try:
                conn.execute(
                    """UPDATE quote_requests
                       SET received_at=COALESCE(received_at, ?),
                           status=CASE
                               WHEN status IN ('complete','needs_review','stale','cancelled')
                               THEN status ELSE 'received_partial' END
                       WHERE id=?""",
                    (message.get("received_at") or iso_now(), request["id"]),
                )
                conn.execute(
                    "UPDATE quote_email_messages SET processed_at=? WHERE id=?",
                    (iso_now(), message_id),
                )
                conn.commit()
            finally:
                conn.close()
            log_activity(
                request["job_id"],
                "quote_email_matched",
                f"Matched vendor reply from {message.get('sender_email') or 'vendor'} without AI",
                {
                    "message_id": message_id,
                    "request_id": request["id"],
                    "method": match["method"],
                    "evidence": match["evidence"],
                },
            )
    else:
        conn = _get_conn()
        try:
            conn.execute(
                "UPDATE quote_email_messages SET processed_at=? WHERE id=?",
                (iso_now(), message_id),
            )
            conn.commit()
        finally:
            conn.close()
    return result


def preview_incoming_match(message: dict) -> dict:
    """Classify a message without persisting it or changing any bid data."""
    return match_message_to_requests(message, _open_request_candidates())


def _assignment_request_is_current(conn, request: dict) -> tuple[bool, str]:
    if str(request.get("status") or "").lower() not in OPEN_REQUEST_STATUSES:
        return False, "That quote request is no longer open."
    job_row = conn.execute(
        "SELECT * FROM jobs WHERE id=?",
        (request["job_id"],),
    ).fetchone()
    if not job_row:
        return False, "That bid no longer exists."
    materials = [
        dict(row)
        for row in conn.execute(
            "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
            (request["job_id"],),
        ).fetchall()
    ]
    request_state = {
        **request,
        "material_snapshot": _json_loads(
            request.get("material_snapshot_json"),
            [],
        ),
    }
    if not _request_is_stale(
        request_state,
        {**dict(job_row), "materials": materials},
    ):
        return True, ""
    conn.execute(
        """UPDATE quote_requests
           SET status='stale'
           WHERE id=? AND status NOT IN ('complete','cancelled','received','stale')""",
        (request["id"],),
    )
    conn.execute(
        """UPDATE quote_followup_events
           SET status='cancelled',
               error='Bid materials changed after approval.'
           WHERE quote_request_id=? AND status='scheduled'""",
        (request["id"],),
    )
    return False, "That quote request is stale because the bid changed."


def _cleanup_assignment_side_effects(
    conn,
    *,
    message_id: int,
    request_id: int,
) -> None:
    conn.execute(
        """DELETE FROM quote_price_matches
           WHERE message_id=? AND quote_request_id=?
             AND status!='applied'""",
        (message_id, request_id),
    )
    conn.execute(
        """UPDATE quote_request_materials
           SET status='requested', quoted_price=NULL,
               source_message_id=NULL, resolved_at=NULL
           WHERE quote_request_id=? AND source_message_id=?
             AND status='needs_review'""",
        (request_id, message_id),
    )


def _assignment_evidence(
    current: Any,
    *,
    evidence_type: str,
    value: str,
) -> str:
    evidence = [
        item
        for item in _json_loads(current, [])
        if not (
            isinstance(item, dict)
            and item.get("type") == evidence_type
        )
    ]
    evidence.append({"type": evidence_type, "value": value})
    return json.dumps(evidence)


def _reject_assignment_target(
    conn,
    message_id: int,
    *,
    job_id: int,
    request_id: int,
    reviewer_name: str,
    decision: str,
    reason: str,
    assignment_token: str | None,
) -> bool:
    message = conn.execute(
        """SELECT match_status, assignment_token, evidence_json
           FROM quote_email_messages WHERE id=?""",
        (message_id,),
    ).fetchone()
    if not message:
        return False
    if assignment_token is not None and (
        str(message["match_status"] or "") != "assigning"
        or not hmac.compare_digest(
            str(message["assignment_token"] or ""),
            assignment_token,
        )
    ):
        return False
    if assignment_token is None and str(message["match_status"] or "") != "needs_review":
        return False

    _cleanup_assignment_side_effects(
        conn,
        message_id=message_id,
        request_id=request_id,
    )
    now = iso_now()
    conn.execute(
        """UPDATE quote_match_candidates
           SET status='rejected', decision=?, reviewer_name=?, resolved_at=?
           WHERE message_id=? AND job_id=? AND quote_request_id=?
             AND status='candidate'""",
        (
            decision,
            reviewer_name,
            now,
            message_id,
            job_id,
            request_id,
        ),
    )
    remaining = int(
        conn.execute(
            """SELECT COUNT(*) FROM quote_match_candidates
               WHERE message_id=? AND status='candidate'""",
            (message_id,),
        ).fetchone()[0]
    )
    where = (
        "id=? AND match_status='assigning' AND assignment_token=?"
        if assignment_token is not None
        else "id=? AND match_status='needs_review'"
    )
    params: tuple[Any, ...] = (
        (message_id, assignment_token)
        if assignment_token is not None
        else (message_id,)
    )
    updated = conn.execute(
        f"""UPDATE quote_email_messages
            SET match_status=?, match_method=?,
                matched_job_id=NULL, matched_request_id=NULL,
                assignment_token='', assignment_started_at=NULL,
                evidence_json=?
            WHERE {where}""",
        (
            "needs_review" if remaining else "ignored",
            "candidate_review" if remaining else decision,
            _assignment_evidence(
                message["evidence_json"],
                evidence_type="assignment_rejected",
                value=reason,
            ),
            *params,
        ),
    )
    return updated.rowcount == 1


def _abort_message_assignment(
    conn,
    message_id: int,
    *,
    request_id: int,
    assignment_token: str,
) -> bool:
    message = conn.execute(
        """SELECT evidence_json FROM quote_email_messages
           WHERE id=? AND match_status='assigning' AND assignment_token=?""",
        (message_id, assignment_token),
    ).fetchone()
    if not message:
        return False
    _cleanup_assignment_side_effects(
        conn,
        message_id=message_id,
        request_id=request_id,
    )
    updated = conn.execute(
        """UPDATE quote_email_messages
           SET match_status='needs_review', match_method='candidate_review',
               matched_job_id=NULL, matched_request_id=NULL,
               assignment_token='', assignment_started_at=NULL,
               evidence_json=?
           WHERE id=? AND match_status='assigning' AND assignment_token=?""",
        (
            _assignment_evidence(
                message["evidence_json"],
                evidence_type="assignment_retry",
                value="The assignment stopped safely. No price was kept.",
            ),
            message_id,
            assignment_token,
        ),
    )
    return updated.rowcount == 1


def _reserve_message_assignment(
    conn,
    message_id: int,
    *,
    job_id: int,
    request_id: int,
    mailbox_email: str,
    assignment_token: str,
) -> dict:
    if not assignment_token:
        raise ValueError("A secure assignment token is required.")
    row = conn.execute(
        "SELECT * FROM quote_email_messages WHERE id=?",
        (message_id,),
    ).fetchone()
    if not row:
        raise ValueError("Message not found.")
    if normalize_email(row["mailbox_email"]) != normalize_email(mailbox_email):
        raise ValueError("That email belongs to another mailbox.")
    current_status = str(row["match_status"] or "")
    reserved_job_id = int(row["matched_job_id"] or 0)
    reserved_request_id = int(row["matched_request_id"] or 0)
    if current_status == "assigning":
        if (
            reserved_job_id != int(job_id)
            or reserved_request_id != int(request_id)
        ):
            raise ValueError("That email is already being assigned to another bid.")
        request_row = conn.execute(
            "SELECT * FROM quote_requests WHERE id=? AND job_id=?",
            (request_id, job_id),
        ).fetchone()
        request = dict(request_row) if request_row else {}
        request_allowed = bool(
            request_row
            and normalize_email(request.get("mailbox_email"))
            == normalize_email(mailbox_email)
        )
        current, current_error = (
            _assignment_request_is_current(conn, request)
            if request_allowed
            else (False, "That quote request is no longer available.")
        )
        candidate = conn.execute(
            """SELECT 1 FROM quote_match_candidates
               WHERE message_id=? AND job_id=? AND quote_request_id=?
                 AND status='candidate'""",
            (message_id, job_id, request_id),
        ).fetchone()
        if not current or not candidate:
            _reject_assignment_target(
                conn,
                message_id,
                job_id=job_id,
                request_id=request_id,
                reviewer_name=mailbox_email,
                decision="request_unavailable_during_assignment",
                reason=current_error or "That bid is no longer an available match.",
                assignment_token=str(row["assignment_token"] or ""),
            )
            raise AssignmentRejectedError(
                current_error or "That bid is no longer an available match."
            )
        started_at = _parse_iso(row["assignment_started_at"])
        lease_expired = (
            started_at is None
            or utc_now() - started_at
            >= timedelta(seconds=ASSIGNMENT_LEASE_SECONDS)
        )
        if not lease_expired:
            raise ValueError(
                "That email assignment is already running. Try again in two minutes."
            )
        _cleanup_assignment_side_effects(
            conn,
            message_id=message_id,
            request_id=request_id,
        )
        updated = conn.execute(
            """UPDATE quote_email_messages
               SET assignment_token=?, assignment_started_at=?
               WHERE id=? AND match_status='assigning'
                 AND assignment_token=?""",
            (
                assignment_token,
                iso_now(),
                message_id,
                str(row["assignment_token"] or ""),
            ),
        )
        if updated.rowcount != 1:
            raise ValueError("That email assignment changed. Refresh and try again.")
        return request
    if current_status != "needs_review":
        raise ValueError("That email has already been handled.")
    if (
        (reserved_job_id or reserved_request_id)
        and (
            reserved_job_id != int(job_id)
            or reserved_request_id != int(request_id)
        )
    ):
        raise ValueError("That email is already being assigned to another bid.")

    candidate = conn.execute(
        """SELECT 1 FROM quote_match_candidates
           WHERE message_id=? AND job_id=? AND quote_request_id=?
             AND status='candidate'""",
        (message_id, job_id, request_id),
    ).fetchone()
    if not candidate:
        raise ValueError("That bid is not an available match for this email.")

    request_row = conn.execute(
        "SELECT * FROM quote_requests WHERE id=?",
        (request_id,),
    ).fetchone()
    if not request_row or int(request_row["job_id"]) != int(job_id):
        _reject_assignment_target(
            conn,
            message_id,
            job_id=job_id,
            request_id=request_id,
            reviewer_name=mailbox_email,
            decision="request_unavailable",
            reason="That quote request does not belong to this bid.",
            assignment_token=None,
        )
        raise AssignmentRejectedError(
            "That quote request does not belong to this bid."
        )
    request = dict(request_row)
    if normalize_email(request.get("mailbox_email")) != normalize_email(
        mailbox_email
    ):
        raise ValueError("That quote request belongs to another mailbox.")
    current, current_error = _assignment_request_is_current(conn, request)
    if not current:
        _reject_assignment_target(
            conn,
            message_id,
            job_id=job_id,
            request_id=request_id,
            reviewer_name=mailbox_email,
            decision="request_unavailable",
            reason=current_error,
            assignment_token=None,
        )
        raise AssignmentRejectedError(current_error)

    updated = conn.execute(
        """UPDATE quote_email_messages
           SET match_status='assigning',
               matched_job_id=?, matched_request_id=?,
               assignment_token=?, assignment_started_at=?
           WHERE id=? AND match_status='needs_review'""",
        (
            job_id,
            request_id,
            assignment_token,
            iso_now(),
            message_id,
        ),
    )
    if updated.rowcount != 1:
        raise ValueError("That email has already been handled.")
    return request


def _finalize_message_assignment(
    conn,
    message_id: int,
    *,
    job_id: int,
    request_id: int,
    reviewer_name: str,
    assignment_token: str,
    pricing: dict,
    staged_sources: Iterable[dict] = (),
) -> bool:
    message = conn.execute(
        """SELECT match_status, matched_job_id, matched_request_id,
                  evidence_json, assignment_token
           FROM quote_email_messages WHERE id=?""",
        (message_id,),
    ).fetchone()
    if (
        not message
        or int(message["matched_job_id"] or 0) != int(job_id)
        or int(message["matched_request_id"] or 0) != int(request_id)
        or str(message["match_status"] or "") != "assigning"
        or not hmac.compare_digest(
            str(message["assignment_token"] or ""),
            assignment_token,
        )
    ):
        raise ValueError("That email has already been handled.")

    request = conn.execute(
        "SELECT * FROM quote_requests WHERE id=? AND job_id=?",
        (request_id, job_id),
    ).fetchone()
    request_data = dict(request) if request else {}
    current, current_error = (
        _assignment_request_is_current(conn, request_data)
        if request
        else (False, "That quote request no longer exists.")
    )
    if not current:
        _reject_assignment_target(
            conn,
            message_id,
            job_id=job_id,
            request_id=request_id,
            reviewer_name=reviewer_name,
            decision="request_closed_during_assignment",
            reason=current_error,
            assignment_token=assignment_token,
        )
        return False

    for copied in staged_sources:
        record_job_artifact(
            job_id,
            "vendor_quote",
            copied["artifact_path"],
            copied["file_hash"],
            copied["file_size"],
            conn=conn,
        )
        record_imported_file(
            job_id,
            copied["file_name"],
            copied["file_hash"],
            copied["file_size"],
            source="outlook",
            artifact_path=copied["artifact_path"],
            artifact_kind="vendor_quote",
            conn=conn,
        )

    now = iso_now()
    evidence = [
        item
        for item in _json_loads(message["evidence_json"], [])
        if not (
            isinstance(item, dict)
            and item.get("type") == "manual_assignment"
        )
    ]
    evidence.append(
        {"type": "manual_assignment", "value": reviewer_name or "Estimator"}
    )
    updated = conn.execute(
        """UPDATE quote_email_messages
           SET match_status='matched', match_method='manual',
               evidence_json=?, assignment_token='',
               assignment_started_at=NULL
           WHERE id=? AND matched_job_id=? AND matched_request_id=?
             AND match_status='assigning' AND assignment_token=?""",
        (
            json.dumps(evidence),
            message_id,
            job_id,
            request_id,
            assignment_token,
        ),
    )
    if updated.rowcount != 1:
        raise ValueError("That email has already been handled.")
    conn.execute(
        """UPDATE quote_match_candidates
           SET status=CASE WHEN job_id=? AND quote_request_id=?
                           THEN 'selected' ELSE 'rejected' END,
               decision='manual_assignment', reviewer_name=?, resolved_at=?
           WHERE message_id=? AND status='candidate'""",
        (job_id, request_id, reviewer_name, now, message_id),
    )
    needs_review = bool(
        int(pricing.get("needs_review") or 0)
        or int(pricing.get("products") or 0) == 0
        or pricing.get("parse_errors")
    )
    if needs_review:
        conn.execute(
            """UPDATE quote_requests
               SET status='needs_review',
                   received_at=COALESCE(received_at, ?)
               WHERE id=? AND status IN
                   ('approved','sending','send_uncertain','sent','waiting',
                    'overdue','received_partial','needs_review','send_failed')""",
            (now, request_id),
        )
        conn.execute(
            """UPDATE quote_followup_events
               SET status='cancelled',
                   error='Estimator review is required before follow-up.'
               WHERE quote_request_id=? AND status='scheduled'""",
            (request_id,),
        )
    else:
        conn.execute(
            """UPDATE quote_requests
               SET status='received_partial',
                   received_at=COALESCE(received_at, ?)
               WHERE id=? AND status IN
                   ('approved','sending','send_uncertain','sent','waiting',
                    'overdue','received_partial','send_failed')""",
            (now, request_id),
        )
    return True


def assign_message(
    message_id: int,
    *,
    job_id: int,
    request_id: int,
    reviewer_name: str,
    mailbox_email: str,
) -> dict:
    assignment_token = uuid.uuid4().hex
    conn = _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        request = _reserve_message_assignment(
            conn,
            message_id,
            job_id=job_id,
            request_id=request_id,
            mailbox_email=mailbox_email,
            assignment_token=assignment_token,
        )
        conn.commit()
    except AssignmentRejectedError:
        conn.commit()
        raise
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    try:
        pricing = _price_products_for_message(
            message_id,
            request,
            "manual",
            defer_request_state=True,
            assignment_token=assignment_token,
        )
    except Exception:
        conn = _get_conn()
        try:
            conn.execute("BEGIN IMMEDIATE")
            _abort_message_assignment(
                conn,
                message_id,
                request_id=request_id,
                assignment_token=assignment_token,
            )
            conn.commit()
        except Exception:
            conn.rollback()
        finally:
            conn.close()
        raise
    try:
        staged_sources, created_artifact_paths = _stage_assignment_artifacts(
            pricing.get("parsed_sources") or [],
            job_id=job_id,
            assignment_token=assignment_token,
        )
    except Exception:
        conn = _get_conn()
        try:
            conn.execute("BEGIN IMMEDIATE")
            _abort_message_assignment(
                conn,
                message_id,
                request_id=request_id,
                assignment_token=assignment_token,
            )
            conn.commit()
        except Exception:
            conn.rollback()
        finally:
            conn.close()
        raise
    conn = _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        finalized = _finalize_message_assignment(
            conn,
            message_id,
            job_id=job_id,
            request_id=request_id,
            reviewer_name=reviewer_name,
            assignment_token=assignment_token,
            pricing=pricing,
            staged_sources=staged_sources,
        )
        conn.commit()
    except Exception:
        conn.rollback()
        _remove_uncommitted_artifact_files(created_artifact_paths)
        try:
            conn.execute("BEGIN IMMEDIATE")
            _abort_message_assignment(
                conn,
                message_id,
                request_id=request_id,
                assignment_token=assignment_token,
            )
            conn.commit()
        except Exception:
            conn.rollback()
        raise
    finally:
        conn.close()
    if not finalized:
        _remove_uncommitted_artifact_files(created_artifact_paths)
        raise ValueError(
            "That quote request changed while the email was being assigned."
        )
    return {"status": "matched", "message_id": message_id, "pricing": pricing}


def ignore_message(
    message_id: int,
    reviewer_name: str = "Estimator",
    *,
    job_id: int | None = None,
    mailbox_email: str,
    conn=None,
) -> bool:
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        message = conn.execute(
            "SELECT * FROM quote_email_messages WHERE id=?",
            (message_id,),
        ).fetchone()
        if (
            not message
            or normalize_email(message["mailbox_email"])
            != normalize_email(mailbox_email)
            or message["match_status"] != "needs_review"
        ):
            conn.rollback()
            return False
        now = iso_now()
        if job_id is not None:
            rejected = conn.execute(
                """UPDATE quote_match_candidates
                   SET status='rejected', decision='ignore_for_bid',
                       reviewer_name=?, resolved_at=?
                   WHERE message_id=? AND job_id=? AND status='candidate'""",
                (reviewer_name, now, message_id, int(job_id)),
            )
            if (
                rejected.rowcount < 1
                and int(message["matched_job_id"] or 0) != int(job_id)
            ):
                conn.rollback()
                return False
            if rejected.rowcount:
                remaining = int(
                    conn.execute(
                        """SELECT COUNT(*) FROM quote_match_candidates
                           WHERE message_id=? AND status='candidate'""",
                        (message_id,),
                    ).fetchone()[0]
                )
                if remaining:
                    conn.commit()
                    return True
        conn.execute(
            """UPDATE quote_email_messages
               SET match_status='ignored', match_method='manual_ignore',
                   evidence_json=?
               WHERE id=? AND match_status='needs_review'""",
            (
                json.dumps([{"type": "ignored_by", "value": reviewer_name}]),
                message_id,
            ),
        )
        conn.execute(
            """UPDATE quote_match_candidates
               SET status='rejected', decision='ignore', reviewer_name=?, resolved_at=?
               WHERE message_id=? AND status='candidate'""",
            (reviewer_name, now, message_id),
        )
        conn.commit()
        return True
    except Exception:
        conn.rollback()
        raise
    finally:
        if owns_connection:
            conn.close()


def decide_price_match(
    match_id: int,
    *,
    decision: str,
    reviewer_name: str,
    reason: str,
    material_id: int | None = None,
    expected_material_id: int | None = None,
    expected_current_price: float | None = None,
) -> dict:
    allowed = {"use_quote", "keep_current", "match_elsewhere", "ignore"}
    if decision not in allowed:
        raise ValueError("Unknown price decision.")
    conn = _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        row = conn.execute(
            "SELECT * FROM quote_price_matches WHERE id=?", (match_id,)
        ).fetchone()
        if not row:
            raise ValueError("Price match not found.")
        match = dict(row)
        if match["status"] != "needs_review":
            raise ValueError("This price match has already been resolved.")
        if match.get("message_id"):
            message = conn.execute(
                """SELECT match_status, matched_job_id, matched_request_id
                   FROM quote_email_messages WHERE id=?""",
                (match["message_id"],),
            ).fetchone()
            if (
                not message
                or str(message["match_status"] or "") != "matched"
                or int(message["matched_job_id"] or 0) != int(match["job_id"])
                or int(message["matched_request_id"] or 0)
                != int(match["quote_request_id"] or 0)
            ):
                raise ValueError(
                    "This email is still being matched to a bid. Refresh and try again."
                )
        if (
            expected_material_id is not None
            and int(match.get("material_id") or 0) != int(expected_material_id)
        ):
            raise ValueError("This price match changed. Refresh before deciding.")
        current_material = (
            conn.execute(
                """SELECT * FROM job_materials
                   WHERE id=? AND job_id=?""",
                (match["material_id"], match["job_id"]),
            ).fetchone()
            if match.get("material_id")
            else None
        )
        if expected_current_price is not None and (
            not current_material
            or not math.isclose(
                float(current_material["unit_price"] or 0),
                float(expected_current_price),
                abs_tol=0.0001,
            )
        ):
            raise ValueError("This material price changed. Refresh before deciding.")
        if decision == "match_elsewhere":
            if not material_id:
                raise ValueError("Choose the material that this price belongs to.")
            material = conn.execute(
                "SELECT * FROM job_materials WHERE id=? AND job_id=?",
                (material_id, match["job_id"]),
            ).fetchone()
            if not material:
                raise ValueError("Chosen material is not part of this bid.")
            requested_target = conn.execute(
                """SELECT 1 FROM quote_request_materials
                   WHERE quote_request_id=? AND material_id=?
                     AND status='requested'""",
                (match["quote_request_id"], material_id),
            ).fetchone()
            if not requested_target:
                raise ValueError(
                    "Choose a material from this vendor's locked quote request."
                )
            original_material_id = match.get("material_id")
            conn.execute(
                """UPDATE quote_price_matches
                   SET material_id=?, material_unit=?,
                       accepted_price_before=?, quantity=?,
                       reason=?, reviewer_name=?
                   WHERE id=?""",
                (
                    material_id,
                    _exact_unit_text(material["unit"]),
                    float(material["unit_price"] or 0),
                    float(material["order_qty"] or material["installed_qty"] or 0),
                    "Estimator selected the material. Confirm Use Quote separately.",
                    reviewer_name,
                    match_id,
                ),
            )
            if original_material_id:
                conn.execute(
                    """UPDATE quote_request_materials
                       SET status='requested'
                       WHERE quote_request_id=? AND material_id=?
                         AND status='needs_review'""",
                    (match["quote_request_id"], original_material_id),
                )
            conn.execute(
                """UPDATE quote_request_materials
                   SET status='needs_review', source_message_id=?
                   WHERE quote_request_id=? AND material_id=?
                     AND status='requested'""",
                (match["message_id"], match["quote_request_id"], material_id),
            )
            conn.commit()
            return {
                "status": "needs_review",
                "decision": "match_elsewhere",
                "material_id": int(material_id),
                "requires_use_quote_confirmation": True,
            }
        elif decision in {"ignore", "keep_current"}:
            material = (
                conn.execute(
                    "SELECT * FROM job_materials WHERE id=?", (match["material_id"],)
                ).fetchone()
                if match.get("material_id")
                else None
            )
            if decision == "keep_current" and (
                not material or float(material["unit_price"] or 0) <= 0
            ):
                raise ValueError("There is no current price to keep.")
            conn.execute(
                """UPDATE quote_price_matches
                   SET status='resolved', decision=?, reviewer_name=?,
                       reason=?, resolved_at=?
                   WHERE id=?""",
                (decision, reviewer_name, reason, iso_now(), match_id),
            )
            if material:
                conn.execute(
                    """UPDATE quote_request_materials
                       SET status=?, source_message_id=?, resolved_at=?
                       WHERE quote_request_id=? AND material_id=?""",
                    (
                        "quoted" if decision == "keep_current" else "requested",
                        match["message_id"],
                        iso_now() if decision == "keep_current" else None,
                        match["quote_request_id"],
                        material["id"],
                    ),
                )
            if decision == "keep_current" and material:
                record_material_price_decision(
                    job_id=int(match["job_id"]),
                    material_id=int(material["id"]),
                    item_code=str(material["item_code"] or ""),
                    decision="keep_accepted",
                    accepted_price_before=float(material["unit_price"] or 0),
                    resolved_price=float(material["unit_price"] or 0),
                    material_unit=str(material["unit"] or ""),
                    quote_price=float(match["quote_price"]),
                    quote_unit=str(match["quote_unit"] or ""),
                    source_hash=str(match["source_hash"]),
                    source_file=str(match["source_file"] or ""),
                    reason=reason,
                    reviewer_name=reviewer_name,
                    conn=conn,
                )
            conn.commit()
            _refresh_request_completion(int(match["quote_request_id"]))
            return {"status": "resolved", "decision": decision}
        conn.rollback()
    finally:
        conn.close()
    if not _apply_exact_price(
        match_id,
        reviewer_name=reviewer_name,
        reason=reason,
        expected_material_id=expected_material_id,
        expected_current_price=expected_current_price,
    ):
        raise ValueError("The selected quote price could not be applied.")
    return {"status": "applied", "decision": "use_quote"}


def due_followups(
    now: datetime | None = None,
    conn=None,
) -> list[dict]:
    current_dt = now or utc_now()
    current = current_dt.isoformat()
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        uncertain_before = (current_dt - timedelta(minutes=30)).isoformat()
        uncertain_requests = [
            int(row["quote_request_id"])
            for row in conn.execute(
                """SELECT DISTINCT quote_request_id
                   FROM quote_followup_events
                   WHERE status='sending' AND claimed_at<=?""",
                (uncertain_before,),
            ).fetchall()
        ]
        for request_id in uncertain_requests:
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='uncertain',
                       error='Microsoft may have accepted this follow-up; it was not sent again.'
                   WHERE quote_request_id=? AND status='sending'""",
                (request_id,),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='cancelled',
                       error='An earlier follow-up has uncertain delivery.'
                   WHERE quote_request_id=? AND status='scheduled'""",
                (request_id,),
            )
            conn.execute(
                """UPDATE quote_requests
                   SET status='needs_review',
                       last_error='A follow-up has uncertain delivery.'
                   WHERE id=? AND status NOT IN ('complete','cancelled','stale')""",
                (request_id,),
            )

        candidates = [
            dict(row)
            for row in conn.execute(
                """SELECT qfe.*, qr.job_id, qr.vendor_name, qr.vendor_email,
                          qr.subject, qr.request_text,
                          qr.mailbox_email,
                          qr.material_snapshot_json,
                          qr.material_snapshot_hash,
                          qr.source_fingerprint,
                          qr.status AS request_status
                   FROM quote_followup_events qfe
                   JOIN quote_requests qr ON qr.id=qfe.quote_request_id
                   WHERE qfe.status='scheduled' AND qfe.scheduled_for<=?
                     AND qr.status IN ('sent','waiting','overdue','received_partial')
                   ORDER BY qfe.scheduled_for, qfe.id
                   LIMIT 50""",
                (current,),
            ).fetchall()
        ]
        claimed = []
        for event in candidates:
            request_id = int(event["quote_request_id"])
            if conn.execute(
                """SELECT 1 FROM quote_followup_events
                   WHERE quote_request_id=? AND status='sending'""",
                (request_id,),
            ).fetchone():
                continue
            job_row = conn.execute(
                "SELECT * FROM jobs WHERE id=?", (event["job_id"],)
            ).fetchone()
            materials = [
                dict(row)
                for row in conn.execute(
                    "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
                    (event["job_id"],),
                ).fetchall()
            ]
            request = {
                **event,
                "material_snapshot": _json_loads(
                    event.get("material_snapshot_json"), []
                ),
            }
            job = {**dict(job_row), "materials": materials} if job_row else {}
            manually_priced = conn.execute(
                """SELECT 1
                   FROM quote_request_materials qrm
                   JOIN job_materials jm ON jm.id=qrm.material_id
                   WHERE qrm.quote_request_id=?
                     AND qrm.status='requested'
                     AND COALESCE(jm.unit_price, 0)>0
                   LIMIT 1""",
                (request_id,),
            ).fetchone()
            if not job or _request_is_stale(request, job) or manually_priced:
                conn.execute(
                    """UPDATE quote_requests SET status='stale'
                       WHERE id=? AND status NOT IN ('complete','cancelled')""",
                    (request_id,),
                )
                conn.execute(
                    """UPDATE quote_followup_events
                       SET status='cancelled',
                           error='Bid materials changed after approval.'
                       WHERE quote_request_id=? AND status='scheduled'""",
                    (request_id,),
                )
                continue
            missing = conn.execute(
                """SELECT COUNT(*) AS count
                   FROM quote_request_materials
                   WHERE quote_request_id=? AND status='requested'""",
                (request_id,),
            ).fetchone()["count"]
            if not missing:
                continue
            cursor = conn.execute(
                """UPDATE quote_followup_events
                   SET status='sending', claimed_at=?, error=''
                   WHERE id=? AND status='scheduled'""",
                (current, event["id"]),
            )
            if cursor.rowcount == 1:
                event["status"] = "sending"
                event["claimed_at"] = current
                claimed.append(event)
        conn.commit()
        return claimed
    except Exception:
        conn.rollback()
        raise
    finally:
        if owns_connection:
            conn.close()


def followup_body(request_id: int, followup_number: int) -> str:
    materials = [
        material
        for material in _request_materials(request_id)
        if material.get("status") == "requested"
    ]
    lines = [
        "Hello,",
        "",
        "Following up on the pricing request below. We still need pricing for:",
        "",
    ]
    for material in materials:
        label = " - ".join(
            value
            for value in (
                str(material.get("item_code") or "").strip(),
                str(material.get("description") or "").strip(),
            )
            if value
        )
        lines.append(
            f"- {label}: {float(material.get('quantity') or 0):g} {material.get('unit') or ''}".rstrip()
        )
    lines.extend(
        [
            "",
            "Please send unit pricing, freight, availability, and lead time.",
            "",
            "Thank you.",
        ]
    )
    return "\n".join(lines)


def _followup_send_state(conn, event_id: int) -> tuple[dict | None, bool, bool]:
    raw = conn.execute(
        """SELECT qfe.status AS event_status, qr.status AS request_status,
                  qr.id AS request_id, qr.job_id,
                  qr.material_snapshot_json, qr.material_snapshot_hash,
                  qr.source_fingerprint,
                  EXISTS (
                      SELECT 1 FROM quote_request_materials qrm
                      WHERE qrm.quote_request_id=qr.id
                        AND qrm.status='requested'
                  ) AS has_missing,
                  EXISTS (
                      SELECT 1
                      FROM quote_request_materials qrm
                      JOIN job_materials jm ON jm.id=qrm.material_id
                      WHERE qrm.quote_request_id=qr.id
                        AND qrm.status='requested'
                        AND COALESCE(jm.unit_price, 0)>0
                  ) AS has_priced_missing
           FROM quote_followup_events qfe
           JOIN quote_requests qr ON qr.id=qfe.quote_request_id
           WHERE qfe.id=?""",
        (event_id,),
    ).fetchone()
    if not raw:
        return None, False, True
    row = dict(raw)
    job_row = conn.execute(
        "SELECT * FROM jobs WHERE id=?",
        (row["job_id"],),
    ).fetchone()
    materials = [
        dict(material)
        for material in conn.execute(
            "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
            (row["job_id"],),
        ).fetchall()
    ]
    request_state = {
        **row,
        "material_snapshot": _json_loads(
            row.get("material_snapshot_json"),
            [],
        ),
    }
    job_state = {**dict(job_row), "materials": materials} if job_row else {}
    snapshot_stale = not job_state or _request_is_stale(request_state, job_state)
    sendable = bool(
        row["event_status"] == "sending"
        and row["request_status"]
        in {"sent", "waiting", "overdue", "received_partial"}
        and row["has_missing"]
        and not row["has_priced_missing"]
        and not snapshot_stale
    )
    return row, sendable, snapshot_stale


def _cancel_blocked_followup(
    conn,
    event_id: int,
    row: dict | None,
    snapshot_stale: bool,
) -> None:
    if not row or row["event_status"] != "sending":
        return
    conn.execute(
        """UPDATE quote_followup_events
           SET status='cancelled',
               error='Request completed, changed, or was cancelled before send.'
           WHERE id=? AND status='sending'""",
        (event_id,),
    )
    if snapshot_stale or row["has_priced_missing"]:
        reason = (
            "Bid materials changed before follow-up."
            if snapshot_stale
            else "A requested material was priced before follow-up."
        )
        conn.execute(
            """UPDATE quote_requests
               SET status='stale', last_error=?
               WHERE id=? AND status NOT IN ('complete','cancelled','stale')""",
            (reason, row["request_id"]),
        )


def confirm_followup_send(event_id: int, conn=None) -> bool:
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        row, sendable, snapshot_stale = _followup_send_state(conn, event_id)
        if not sendable:
            _cancel_blocked_followup(conn, event_id, row, snapshot_stale)
        conn.commit()
        return sendable
    except Exception:
        conn.rollback()
        raise
    finally:
        if owns_connection:
            conn.close()


@contextmanager
def followup_send_guard(event_id: int):
    conn = _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        row, sendable, snapshot_stale = _followup_send_state(conn, event_id)
        if not sendable:
            _cancel_blocked_followup(conn, event_id, row, snapshot_stale)
            conn.commit()
            yield False
            return
        yield sendable
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def mark_followup_sent(
    event_id: int,
    graph_message_id: str = "",
    conn=None,
) -> None:
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        row = conn.execute(
            """SELECT quote_request_id, followup_number
               FROM quote_followup_events WHERE id=?""",
            (event_id,),
        ).fetchone()
        conn.execute(
            """UPDATE quote_followup_events
               SET status='sent',
                   graph_message_id=CASE WHEN ? != '' THEN ? ELSE graph_message_id END,
                   sent_at=COALESCE(sent_at, ?), error=''
               WHERE id=? AND status IN ('sending','sent','uncertain')""",
            (graph_message_id, graph_message_id, iso_now(), event_id),
        )
        if row:
            conn.execute(
                """UPDATE quote_requests
                   SET status='waiting', last_error=''
                   WHERE id=? AND status='needs_review'
                     AND lower(last_error) LIKE '%follow-up%uncertain%'""",
                (row["quote_request_id"],),
            )
            if int(row["followup_number"] or 0) == 1:
                conn.execute(
                    """UPDATE quote_followup_events
                       SET status='scheduled', error=''
                       WHERE quote_request_id=? AND followup_number=2
                         AND status='cancelled'
                         AND error='An earlier follow-up has uncertain delivery.'""",
                    (row["quote_request_id"],),
                )
        if owns_connection:
            conn.commit()
    finally:
        if owns_connection:
            conn.close()


def record_followup_send_token(event_id: int, send_token: str) -> None:
    token = str(send_token or "").strip()
    if not token:
        raise ValueError("Follow-up send token is required.")
    conn = _get_conn()
    try:
        cursor = conn.execute(
            """UPDATE quote_followup_events
               SET send_token=?
               WHERE id=? AND status='sending'""",
            (token, event_id),
        )
        conn.commit()
        if cursor.rowcount != 1:
            raise ValueError("Follow-up is no longer waiting to send.")
    finally:
        conn.close()


def mark_followup_failed(event_id: int, error: str) -> None:
    conn = _get_conn()
    try:
        row = conn.execute(
            "SELECT quote_request_id FROM quote_followup_events WHERE id=?",
            (event_id,),
        ).fetchone()
        conn.execute(
            """UPDATE quote_followup_events
               SET status='failed', error=?
               WHERE id=? AND status='sending'""",
            (str(error)[:1000], event_id),
        )
        if row:
            conn.execute(
                """UPDATE quote_requests
                   SET status='needs_review', last_error=?
                   WHERE id=? AND status NOT IN ('complete','cancelled','stale')""",
                (f"Follow-up failed: {str(error)[:900]}", row["quote_request_id"]),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='cancelled',
                       error='An earlier follow-up failed.'
                   WHERE quote_request_id=? AND status='scheduled'""",
                (row["quote_request_id"],),
            )
        conn.commit()
    finally:
        conn.close()


def mark_followup_uncertain(event_id: int, error: str) -> None:
    conn = _get_conn()
    try:
        row = conn.execute(
            "SELECT quote_request_id FROM quote_followup_events WHERE id=?",
            (event_id,),
        ).fetchone()
        conn.execute(
            """UPDATE quote_followup_events
               SET status='uncertain', error=?
               WHERE id=? AND status='sending'""",
            (str(error)[:1000], event_id),
        )
        if row:
            conn.execute(
                """UPDATE quote_requests
                   SET status='needs_review', last_error=?
                   WHERE id=? AND status NOT IN ('complete','cancelled','stale')""",
                (
                    f"Follow-up delivery is uncertain: {str(error)[:850]}",
                    row["quote_request_id"],
                ),
            )
            conn.execute(
                """UPDATE quote_followup_events
                   SET status='cancelled',
                       error='An earlier follow-up has uncertain delivery.'
                   WHERE quote_request_id=? AND status='scheduled'""",
                (row["quote_request_id"],),
            )
        conn.commit()
    finally:
        conn.close()


def cancel_request(request_id: int, conn=None) -> bool:
    owns_connection = conn is None
    conn = conn or _get_conn()
    try:
        conn.execute("BEGIN IMMEDIATE")
        sending = conn.execute(
            """SELECT 1 FROM quote_followup_events
               WHERE quote_request_id=? AND status='sending'""",
            (request_id,),
        ).fetchone()
        if sending:
            conn.rollback()
            return False
        cursor = conn.execute(
            """UPDATE quote_requests
               SET status='cancelled', cancelled_at=?
               WHERE id=? AND status NOT IN ('complete','cancelled','received')""",
            (iso_now(), request_id),
        )
        if cursor.rowcount != 1:
            conn.rollback()
            return False
        conn.execute(
            """UPDATE quote_followup_events
               SET status='cancelled'
               WHERE quote_request_id=? AND status='scheduled'""",
            (request_id,),
        )
        conn.execute(
            """UPDATE quote_price_matches
               SET status='cancelled', decision='request_cancelled',
                   reason='Quote request was cancelled before a price decision.',
                   resolved_at=?
               WHERE quote_request_id=?
                 AND status IN ('needs_review','ready_to_apply')""",
            (iso_now(), request_id),
        )
        conn.commit()
        return True
    except Exception:
        conn.rollback()
        raise
    finally:
        if owns_connection:
            conn.close()


def _live_job_fingerprint(job: dict) -> str:
    payload = {
        key: job.get(key)
        for key in (
            "id",
            "project_name",
            "gc_name",
            "address",
            "city",
            "state",
            "zip",
            "proposal_data",
            "bid_data",
        )
    }
    for key in ("materials", "sundries", "labor", "bundles", "quotes"):
        payload[key] = job.get(key) or []
    job_id = int(job.get("id") or 0)
    if job_id:
        conn = _get_conn()
        try:
            queries = {
                "quote_requests": (
                    "SELECT * FROM quote_requests WHERE job_id=? ORDER BY id",
                    (job_id,),
                ),
                "quote_request_materials": (
                    """SELECT qrm.* FROM quote_request_materials qrm
                       JOIN quote_requests qr ON qr.id=qrm.quote_request_id
                       WHERE qr.job_id=? ORDER BY qrm.id""",
                    (job_id,),
                ),
                "quote_email_messages": (
                    """SELECT * FROM quote_email_messages
                       WHERE matched_job_id=?
                          OR id IN (
                              SELECT message_id FROM quote_match_candidates
                              WHERE job_id=?
                          )
                       ORDER BY id""",
                    (job_id, job_id),
                ),
                "quote_match_candidates": (
                    """SELECT * FROM quote_match_candidates
                       WHERE job_id=? ORDER BY id""",
                    (job_id,),
                ),
                "quote_price_matches": (
                    "SELECT * FROM quote_price_matches WHERE job_id=? ORDER BY id",
                    (job_id,),
                ),
                "quote_followup_events": (
                    """SELECT qfe.* FROM quote_followup_events qfe
                       JOIN quote_requests qr ON qr.id=qfe.quote_request_id
                       WHERE qr.job_id=? ORDER BY qfe.id""",
                    (job_id,),
                ),
                "job_quotes": (
                    "SELECT * FROM job_quotes WHERE job_id=? ORDER BY id",
                    (job_id,),
                ),
                "vendor_prices": (
                    "SELECT * FROM vendor_prices WHERE job_id=? ORDER BY id",
                    (job_id,),
                ),
                "imported_files": (
                    "SELECT * FROM imported_files WHERE job_id=? ORDER BY id",
                    (job_id,),
                ),
                "material_price_decisions": (
                    """SELECT * FROM material_price_decisions
                       WHERE job_id=? ORDER BY id""",
                    (job_id,),
                ),
                "job_artifacts": (
                    "SELECT * FROM job_artifacts WHERE job_id=? ORDER BY id",
                    (job_id,),
                ),
            }
            payload["real_quote_records"] = {
                name: [
                    dict(row)
                    for row in conn.execute(sql, params).fetchall()
                ]
                for name, (sql, params) in queries.items()
            }
        finally:
            conn.close()
    return _sha256_bytes(
        json.dumps(payload, sort_keys=True, default=str, separators=(",", ":")).encode(
            "utf-8"
        )
    )


def _simulation_result(
    scenario: str,
    label: str,
    expected: Any,
    actual: Any,
    detail: str,
) -> dict:
    passed = expected == actual
    return {
        "scenario": scenario,
        "label": label,
        "expected": expected,
        "actual": actual,
        "detail": detail,
        "passed": passed,
        "status": "pass" if passed else "fail",
    }


def _insert_shadow_row(conn, table: str, values: dict) -> None:
    columns = {
        str(row["name"])
        for row in conn.execute(f"PRAGMA table_info({table})").fetchall()
    }
    selected = {key: value for key, value in values.items() if key in columns}
    names = list(selected)
    placeholders = ",".join("?" for _ in names)
    conn.execute(
        f"INSERT INTO {table} ({','.join(names)}) VALUES ({placeholders})",
        [selected[name] for name in names],
    )


def _simulation_shadow(
    job_id: int,
    isolated_materials: list[dict],
    virtual_now: datetime,
) -> tuple[sqlite3.Connection, int, int]:
    source = _get_conn()
    shadow = sqlite3.connect(":memory:")
    shadow.row_factory = sqlite3.Row
    shadow.execute("PRAGMA foreign_keys = ON")
    tables = (
        "jobs",
        "job_materials",
        "quote_requests",
        "quote_request_materials",
        "quote_email_messages",
        "quote_match_candidates",
        "quote_price_matches",
        "quote_followup_events",
        "material_price_decisions",
        "job_quotes",
        "vendor_prices",
        "vendors",
        "imported_files",
        "job_artifacts",
    )
    try:
        for table in tables:
            row = source.execute(
                """SELECT sql FROM sqlite_master
                   WHERE type='table' AND name=?""",
                (table,),
            ).fetchone()
            if not row or not row["sql"]:
                raise ValueError(f"Simulation schema is missing {table}.")
            shadow.execute(str(row["sql"]))
        shadow.execute(
            """CREATE UNIQUE INDEX idx_sim_imported_files_dedup
               ON imported_files(job_id, file_hash)"""
        )
        job_row = source.execute(
            "SELECT * FROM jobs WHERE id=?",
            (job_id,),
        ).fetchone()
        material_rows = source.execute(
            "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
            (job_id,),
        ).fetchall()
        if not job_row or not material_rows:
            raise ValueError("The selected bid has no material rows to copy.")
        _insert_shadow_row(shadow, "jobs", dict(job_row))
        base_materials = [dict(row) for row in material_rows]
        for index, simulated in enumerate(isolated_materials):
            base = dict(base_materials[min(index, len(base_materials) - 1)])
            base.update(simulated)
            base["job_id"] = job_id
            base["unit_price"] = 0
            base["extended_cost"] = 0
            base["price_source"] = ""
            base["quote_source_hash"] = ""
            base["quote_file_name"] = ""
            _insert_shadow_row(shadow, "job_materials", base)
    finally:
        source.close()

    request_id = 9_100_001
    match_id = 9_200_001
    material_snapshot = _material_snapshot(isolated_materials)
    material_hash = snapshot_hash(isolated_materials)
    job_state = {
        **dict(
            shadow.execute("SELECT * FROM jobs WHERE id=?", (job_id,)).fetchone()
        ),
        "materials": [
            dict(row)
            for row in shadow.execute(
                "SELECT * FROM job_materials WHERE job_id=? ORDER BY id",
                (job_id,),
            ).fetchall()
        ],
    }
    shadow.execute(
        """INSERT INTO quote_requests
           (id, job_id, vendor_name, vendor_email, status, material_ids,
            request_text, sent_at, created_at, mailbox_email, subject,
            material_snapshot_json, material_snapshot_hash, source_fingerprint,
            approved_at, approved_by)
           VALUES (?, ?, ?, ?, 'waiting', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            request_id,
            job_id,
            "Simulation Vendor One",
            "vendor-one@example.test",
            json.dumps([{"id": item["id"]} for item in isolated_materials]),
            "Simulation request",
            virtual_now.isoformat(),
            virtual_now.isoformat(),
            "estimator@example.test",
            "Simulation quote request",
            json.dumps(material_snapshot),
            material_hash,
            _job_source_fingerprint(job_state),
            virtual_now.isoformat(),
            "Simulation",
        ),
    )
    for material in material_snapshot:
        shadow.execute(
            """INSERT INTO quote_request_materials
               (quote_request_id, material_id, item_code, description, unit,
                quantity, vendor_name, snapshot_hash, status, created_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'requested', ?)""",
            (
                request_id,
                material["id"],
                material["item_code"],
                material["description"],
                material["unit"],
                material["quantity"],
                "Simulation Vendor One",
                material_hash,
                virtual_now.isoformat(),
            ),
        )
    first = material_snapshot[0]
    shadow.execute(
        """INSERT INTO quote_price_matches
           (id, job_id, quote_request_id, material_id, item_code, description,
            quote_price, quote_unit, material_unit, accepted_price_before,
            quantity, source_hash, source_file, match_method, status, reason,
            created_at)
           VALUES (?, ?, ?, ?, ?, ?, 4.25, ?, ?, 0, ?, ?, ?,
                   'direct_reply', 'ready_to_apply', ?, ?)""",
        (
            match_id,
            job_id,
            request_id,
            first["id"],
            first["item_code"],
            first["description"],
            first["unit"],
            first["unit"],
            first["quantity"],
            _sha256_bytes(b"simulation-quote-proof"),
            "simulation-quote.csv",
            "Verified direct simulation reply.",
            virtual_now.isoformat(),
        ),
    )
    for number, days in ((1, 3), (2, 6)):
        shadow.execute(
            """INSERT INTO quote_followup_events
               (quote_request_id, followup_number, scheduled_for, status, created_at)
               VALUES (?, ?, ?, 'scheduled', ?)""",
            (
                request_id,
                number,
                add_business_days(virtual_now, days).isoformat(),
                virtual_now.isoformat(),
            ),
        )
    shadow.commit()
    return shadow, request_id, match_id


def _add_simulation_candidate_request(
    conn,
    *,
    source_job_id: int,
    source_request_id: int,
    suffix: int,
) -> tuple[int, int]:
    source_job = conn.execute(
        "SELECT * FROM jobs WHERE id=?",
        (source_job_id,),
    ).fetchone()
    source_request = conn.execute(
        "SELECT * FROM quote_requests WHERE id=?",
        (source_request_id,),
    ).fetchone()
    if not source_job or not source_request:
        raise ValueError("Simulation candidate source is missing.")

    other_job = dict(source_job)
    other_job_id = int(source_job_id) + 10_000_000 + int(suffix)
    other_job["id"] = other_job_id
    other_job["project_name"] = f"Other Simulation Bid {suffix}"
    other_job["slug"] = f"other-simulation-bid-{suffix}"
    _insert_shadow_row(conn, "jobs", other_job)

    other_request = dict(source_request)
    other_request_id = int(source_request_id) + int(suffix)
    other_request["id"] = other_request_id
    other_request["job_id"] = other_job_id
    other_request["send_token"] = f"simulation-other-send-{suffix}"
    other_request["outlook_message_id"] = ""
    other_request["internet_message_id"] = (
        f"<simulation-other-request-{suffix}@example.test>"
    )
    other_request["conversation_id"] = f"simulation-other-conversation-{suffix}"
    _insert_shadow_row(conn, "quote_requests", other_request)
    conn.commit()
    return other_job_id, other_request_id


def _simulation_manual_price_case(
    job_id: int,
    isolated_materials: list[dict],
    virtual_now: datetime,
    *,
    current_price: float,
    expected_current_price: float,
    request_status: str = "waiting",
) -> dict:
    shadow, request_id, match_id = _simulation_shadow(
        job_id,
        isolated_materials,
        virtual_now,
    )
    material_id = int(isolated_materials[0]["id"])
    try:
        shadow.execute(
            """UPDATE quote_price_matches
               SET status='needs_review' WHERE id=?""",
            (match_id,),
        )
        shadow.execute(
            """UPDATE job_materials
               SET unit_price=?, extended_cost=?
               WHERE id=? AND job_id=?""",
            (
                current_price,
                round(
                    current_price
                    * float(
                        isolated_materials[0].get("order_qty")
                        or isolated_materials[0].get("installed_qty")
                        or 0
                    ),
                    2,
                ),
                material_id,
                job_id,
            ),
        )
        shadow.execute(
            "UPDATE quote_requests SET status=? WHERE id=?",
            (request_status, request_id),
        )
        shadow.commit()
        applied = _apply_exact_price(
            match_id,
            reviewer_name="Simulation",
            reason="Explicit estimator decision in an isolated bid copy.",
            expected_material_id=material_id,
            expected_current_price=expected_current_price,
            conn=shadow,
        )
        material = shadow.execute(
            "SELECT unit_price FROM job_materials WHERE id=?",
            (material_id,),
        ).fetchone()
        match = shadow.execute(
            "SELECT status FROM quote_price_matches WHERE id=?",
            (match_id,),
        ).fetchone()
        return {
            "applied": applied,
            "material_price": float(material["unit_price"] or 0),
            "match_status": str(match["status"] or ""),
        }
    finally:
        shadow.close()


def create_simulation_run(job_id: int, scenario: str = "all") -> dict:
    scenario = str(scenario or "all").strip().lower()
    if scenario not in SIMULATION_SCENARIOS:
        raise ValueError(f"Unknown simulation scenario: {scenario}")
    job = load_job(job_id)
    if not job:
        raise ValueError("Job not found")
    fingerprint_before = _live_job_fingerprint(job)
    virtual_now = utc_now()
    source_materials = list(job.get("materials") or [])
    if not source_materials:
        raise ValueError("Add at least one material before running this bid test.")
    isolated_materials = []
    for index, material in enumerate(source_materials[:2], start=1):
        isolated_materials.append(
            {
                **material,
                "id": int(material.get("id") or 9_000_000 + index),
                "item_code": str(
                    material.get("item_code") or f"SIM-{index:03d}"
                ).strip(),
                "description": str(
                    material.get("description") or f"Copied material {index}"
                ).strip(),
                "unit": normalize_quote_unit(material.get("unit")) or "EA",
                "order_qty": float(
                    material.get("order_qty")
                    or material.get("installed_qty")
                    or 1
                ),
                "installed_qty": float(
                    material.get("installed_qty")
                    or material.get("order_qty")
                    or 1
                ),
                "vendor": "Simulation Vendor One",
                "unit_price": 0,
            }
        )
    if len(isolated_materials) == 1:
        first = isolated_materials[0]
        isolated_materials.append(
            {
                **first,
                "id": int(first["id"]) + 9_000_000,
                "item_code": f"{first['item_code']}-SIM2",
                "description": f"{first['description']} second test line",
            }
        )
    requested_materials = [
        {
            "material_id": material["id"],
            "item_code": material["item_code"],
            "description": material["description"],
            "unit": material["unit"],
            "quantity": material["order_qty"],
            "status": "requested",
        }
        for material in isolated_materials
    ]
    fake_request = {
        "id": 9_100_001,
        "job_id": job_id,
        "vendor_email": "vendor-one@example.test",
        "contact_email": "vendor-one@example.test",
        "internet_message_id": "<simulation-request-1@example.test>",
        "conversation_id": "simulation-conversation-1",
        "project_name": job.get("project_name") or "Simulation Project",
        "slug": job.get("slug") or "simulation-project",
        "gc_name": job.get("gc_name") or "Simulation GC",
        "address": job.get("address") or "123 Simulation Street",
        "material_snapshot": _material_snapshot(isolated_materials),
    }
    second_request = {
        **fake_request,
        "id": 9_100_002,
        "job_id": job_id + 1,
        "project_name": "Other Simulation Bid",
        "slug": "other-simulation-bid",
        "conversation_id": "simulation-conversation-2",
        "internet_message_id": "<simulation-request-2@example.test>",
    }
    other_vendor_request = {
        **fake_request,
        "id": 9_100_003,
        "vendor_email": "vendor-two@example.test",
        "contact_email": "vendor-two@example.test",
        "conversation_id": "simulation-conversation-3",
        "internet_message_id": "<simulation-request-3@example.test>",
    }

    csv_buffer = io.StringIO()
    csv_writer = csv.writer(csv_buffer)
    csv_writer.writerow(["item_code", "description", "unit_price", "unit"])
    csv_writer.writerow(
        [
            isolated_materials[0]["item_code"],
            isolated_materials[0]["description"],
            "4.25",
            isolated_materials[0]["unit"],
        ]
    )
    csv_bytes = csv_buffer.getvalue().encode("utf-8")
    csv_path = (
        ARTIFACT_ROOT
        / "shared"
        / "tmp"
        / f"simulation-{_sha256_bytes(csv_bytes)[:12]}.csv"
    )
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    if not csv_path.exists():
        csv_path.write_bytes(csv_bytes)
    parsed_products = parse_quote_file_deterministic(str(csv_path))
    product = parsed_products[0] if parsed_products else {}

    direct_message = {
        "sender_email": "vendor-one@example.test",
        "subject": "Completely changed subject",
        "headers": {"In-Reply-To": "<simulation-request-1@example.test>"},
    }
    direct_match = match_message_to_requests(direct_message, [fake_request])
    direct_evaluation = evaluate_price_candidate(
        product,
        requested_materials=requested_materials,
        current_materials=isolated_materials,
        match_method=direct_match.get("method") or "",
        request_stale=False,
    )
    changed_subject_match = match_message_to_requests(
        {
            "sender_email": "vendor-one@example.test",
            "subject": "No original words remain",
            "conversation_id": "simulation-conversation-1",
        },
        [fake_request],
    )
    standalone_match = match_message_to_requests(
        {
            "sender_email": "vendor-one@example.test",
            "subject": f"Pricing for {fake_request['project_name']}",
            "body_text": (
                f"{isolated_materials[0]['item_code']} "
                f"$4.25/{isolated_materials[0]['unit']}"
            ),
        },
        [fake_request],
    )
    standalone_evaluation = evaluate_price_candidate(
        product,
        requested_materials=requested_materials,
        current_materials=isolated_materials,
        match_method=standalone_match.get("method") or "",
        request_stale=False,
    )
    ambiguous_match = match_message_to_requests(
        {
            "sender_email": "vendor-one@example.test",
            "subject": "Pricing attached",
            "body_text": isolated_materials[0]["item_code"],
        },
        [fake_request, second_request],
    )
    unrelated_vendor_match = match_message_to_requests(
        {
            "sender_email": "vendor-one@example.test",
            "subject": "Lunch schedule",
            "body_text": "This message contains no project or material facts.",
        },
        [fake_request, second_request],
    )
    wrong_unit_evaluation = evaluate_price_candidate(
        {
            **product,
            "source_unit": (
                "EA"
                if _exact_unit_key(product.get("source_unit") or product.get("unit"))
                != "EA"
                else "SF"
            ),
        },
        requested_materials=requested_materials,
        current_materials=isolated_materials,
        match_method="direct_reply",
        request_stale=False,
    )
    packaging_material = {
        **isolated_materials[0],
        "id": int(isolated_materials[0]["id"]) + 8_000_000,
        "item_code": "PKG-EXACT-100",
        "unit": "CTN",
    }
    packaging_unit_evaluation = evaluate_price_candidate(
        {
            "item_code": packaging_material["item_code"],
            "product_name": "Packaging exact-unit check",
            "unit_price": 10,
            "unit": "CTN",
            "source_unit": "BOX",
        },
        requested_materials=[
            {
                "material_id": packaging_material["id"],
                "item_code": packaging_material["item_code"],
                "unit": "CTN",
                "status": "requested",
            }
        ],
        current_materials=[packaging_material],
        match_method="direct_reply",
        request_stale=False,
    )
    changed_price_materials = [
        {**isolated_materials[0], "unit_price": 3.95},
        *isolated_materials[1:],
    ]
    changed_price_evaluation = evaluate_price_candidate(
        product,
        requested_materials=requested_materials,
        current_materials=changed_price_materials,
        match_method="direct_reply",
        request_stale=False,
    )
    stale_evaluation = evaluate_price_candidate(
        product,
        requested_materials=requested_materials,
        current_materials=isolated_materials,
        match_method="direct_reply",
        request_stale=True,
    )
    quantity_text = (
        ARTIFACT_ROOT / "shared" / "tmp" / "simulation-quantity-not-price.txt"
    )
    quantity_text.write_text(
        (
            f"{isolated_materials[0]['item_code']} "
            f"{isolated_materials[0]['description']} "
            f"{isolated_materials[0]['order_qty']:g} "
            f"{isolated_materials[0]['unit']}"
        ),
        encoding="utf-8",
    )
    quantity_products = parse_quote_file_deterministic(str(quantity_text))
    broken_path = ARTIFACT_ROOT / "shared" / "tmp" / "simulation-broken.xlsx"
    broken_path.write_bytes(b"not a valid spreadsheet")
    try:
        parse_quote_file_deterministic(str(broken_path))
        broken_attachment_status = "parsed"
    except Exception:
        broken_attachment_status = "needs_review"
    conn = _get_conn()
    try:
        table_sql = str(
            (
                conn.execute(
                    """SELECT sql FROM sqlite_master
                       WHERE type='table' AND name='quote_email_messages'"""
                ).fetchone()
                or {"sql": ""}
            )["sql"]
            or ""
        ).lower()
    finally:
        conn.close()
    duplicate_guard = (
        "graph_message_id text not null unique" in " ".join(table_sql.split())
    )
    multi_vendor_match = match_message_to_requests(
        direct_message,
        [fake_request, other_vendor_request],
    )

    module_source = Path(__file__).read_text(encoding="utf-8")
    openai_imports = re.findall(
        r"^\s*(?:from\s+openai\b|import\s+openai\b)",
        module_source,
        flags=re.MULTILINE,
    )
    shadow, shadow_request_id, shadow_match_id = _simulation_shadow(
        job_id,
        isolated_materials,
        virtual_now,
    )
    try:
        shadow_applied = _apply_exact_price(
            shadow_match_id,
            reviewer_name="Simulation",
            reason="Isolated production transaction test.",
            automatic=True,
            conn=shadow,
        )
        shadow_material = shadow.execute(
            "SELECT unit_price, quote_source_hash FROM job_materials WHERE id=?",
            (isolated_materials[0]["id"],),
        ).fetchone()
        shadow_request = shadow.execute(
            "SELECT status FROM quote_requests WHERE id=?",
            (shadow_request_id,),
        ).fetchone()
        shadow_atomic_counts = {
            "decision": int(
                shadow.execute(
                    "SELECT COUNT(*) FROM material_price_decisions"
                ).fetchone()[0]
            ),
            "job_quote": int(
                shadow.execute("SELECT COUNT(*) FROM job_quotes").fetchone()[0]
            ),
            "verified_history": int(
                shadow.execute(
                    "SELECT COUNT(*) FROM vendor_prices WHERE verified=1"
                ).fetchone()[0]
            ),
        }
        first_due = due_followups(
            add_business_days(virtual_now, 3),
            conn=shadow,
        )
        duplicate_first_due = due_followups(
            add_business_days(virtual_now, 3),
            conn=shadow,
        )
        if first_due:
            mark_followup_sent(
                int(first_due[0]["id"]),
                "simulation-followup-1",
                conn=shadow,
            )
            shadow.commit()
        second_due = due_followups(
            add_business_days(virtual_now, 6),
            conn=shadow,
        )
        duplicate_second_due = due_followups(
            add_business_days(virtual_now, 6),
            conn=shadow,
        )
        if second_due:
            mark_followup_sent(
                int(second_due[0]["id"]),
                "simulation-followup-2",
                conn=shadow,
            )
            shadow.commit()
        shadow_price_transaction_ok = bool(
            shadow_applied
            and shadow_material
            and float(shadow_material["unit_price"] or 0) == 4.25
            and shadow_material["quote_source_hash"]
            and shadow_request
            and shadow_request["status"] == "received_partial"
            and all(value == 1 for value in shadow_atomic_counts.values())
        )
        shadow_followups_once = bool(
            len(first_due) == 1
            and not duplicate_first_due
            and len(second_due) == 1
            and not duplicate_second_due
        )
    finally:
        shadow.close()

    changed_price_case = _simulation_manual_price_case(
        job_id,
        isolated_materials,
        virtual_now,
        current_price=3.95,
        expected_current_price=3.95,
    )
    stale_screen_case = _simulation_manual_price_case(
        job_id,
        isolated_materials,
        virtual_now,
        current_price=3.95,
        expected_current_price=0,
    )

    cancel_shadow, cancel_request_id, cancel_match_id = _simulation_shadow(
        job_id,
        isolated_materials,
        virtual_now,
    )
    try:
        cancel_ok = cancel_request(cancel_request_id, conn=cancel_shadow)
        cancelled_request = cancel_shadow.execute(
            "SELECT status FROM quote_requests WHERE id=?",
            (cancel_request_id,),
        ).fetchone()
        cancelled_match = cancel_shadow.execute(
            "SELECT status FROM quote_price_matches WHERE id=?",
            (cancel_match_id,),
        ).fetchone()
        cancelled_followups = {
            str(row["status"])
            for row in cancel_shadow.execute(
                """SELECT status FROM quote_followup_events
                   WHERE quote_request_id=?""",
                (cancel_request_id,),
            ).fetchall()
        }
        cancelled_request_is_terminal = bool(
            cancel_ok
            and cancelled_request
            and cancelled_request["status"] == "cancelled"
            and cancelled_match
            and cancelled_match["status"] == "cancelled"
            and cancelled_followups == {"cancelled"}
        )
    finally:
        cancel_shadow.close()

    ignore_shadow, ignore_request_id, _ = _simulation_shadow(
        job_id,
        isolated_materials,
        virtual_now,
    )
    try:
        message_id = 9_300_001
        ignore_shadow.execute(
            """INSERT INTO quote_email_messages
               (id, graph_message_id, mailbox_email, direction, sender_email,
                recipients_json, subject, body_text, attachment_manifest_json,
                match_status, evidence_json, processed_at, created_at)
               VALUES (?, ?, ?, 'inbound', ?, '[]', ?, '', '[]',
                       'needs_review', '[]', ?, ?)""",
            (
                message_id,
                "simulation-ambiguous-message",
                "estimator@example.test",
                "vendor-one@example.test",
                "Ambiguous simulation quote",
                virtual_now.isoformat(),
                virtual_now.isoformat(),
            ),
        )
        ignore_other_job_id, ignore_other_request_id = (
            _add_simulation_candidate_request(
                ignore_shadow,
                source_job_id=job_id,
                source_request_id=ignore_request_id,
                suffix=1,
            )
        )
        for candidate_job_id, candidate_request_id in (
            (job_id, ignore_request_id),
            (ignore_other_job_id, ignore_other_request_id),
        ):
            ignore_shadow.execute(
                """INSERT INTO quote_match_candidates
                   (message_id, job_id, quote_request_id, status,
                    match_method, evidence_json, created_at)
                   VALUES (?, ?, ?, 'candidate', 'exact_clues', '[]', ?)""",
                (
                    message_id,
                    candidate_job_id,
                    candidate_request_id,
                    virtual_now.isoformat(),
                ),
            )
        ignore_shadow.commit()
        ignored_for_one_bid = ignore_message(
            message_id,
            "Simulation",
            job_id=job_id,
            mailbox_email="estimator@example.test",
            conn=ignore_shadow,
        )
        candidate_states = {
            int(row["job_id"]): str(row["status"])
            for row in ignore_shadow.execute(
                """SELECT job_id, status FROM quote_match_candidates
                   WHERE message_id=?""",
                (message_id,),
            ).fetchall()
        }
        ignored_message = ignore_shadow.execute(
            """SELECT match_status, processed_at
               FROM quote_email_messages WHERE id=?""",
            (message_id,),
        ).fetchone()
        per_bid_ignore_isolated = bool(
            ignored_for_one_bid
            and candidate_states.get(job_id) == "rejected"
            and candidate_states.get(ignore_other_job_id) == "candidate"
            and ignored_message
            and ignored_message["match_status"] == "needs_review"
            and bool(ignored_message["processed_at"])
        )
    finally:
        ignore_shadow.close()

    assign_shadow, assign_request_id, assign_match_id = _simulation_shadow(
        job_id,
        isolated_materials,
        virtual_now,
    )
    try:
        assign_message_id = 9_300_002
        assign_shadow.execute(
            """INSERT INTO quote_email_messages
               (id, graph_message_id, mailbox_email, direction, sender_email,
                recipients_json, subject, body_text, attachment_manifest_json,
                match_status, evidence_json, processed_at, created_at)
               VALUES (?, ?, ?, 'inbound', ?, '[]', ?, '', '[]',
                       'needs_review', '[]', ?, ?)""",
            (
                assign_message_id,
                "simulation-assign-ambiguous-message",
                "estimator@example.test",
                "vendor-one@example.test",
                "Assignable ambiguous simulation quote",
                virtual_now.isoformat(),
                virtual_now.isoformat(),
            ),
        )
        assign_other_job_id, assign_other_request_id = (
            _add_simulation_candidate_request(
                assign_shadow,
                source_job_id=job_id,
                source_request_id=assign_request_id,
                suffix=2,
            )
        )
        for candidate_job_id, candidate_request_id in (
            (job_id, assign_request_id),
            (assign_other_job_id, assign_other_request_id),
        ):
            assign_shadow.execute(
                """INSERT INTO quote_match_candidates
                   (message_id, job_id, quote_request_id, status,
                    match_method, evidence_json, created_at)
                   VALUES (?, ?, ?, 'candidate', 'exact_clues', '[]', ?)""",
                (
                    assign_message_id,
                    candidate_job_id,
                    candidate_request_id,
                    virtual_now.isoformat(),
                ),
            )
        assign_shadow.commit()
        first_assignment_token = uuid.uuid4().hex
        assign_shadow.execute("BEGIN IMMEDIATE")
        assigned_request = _reserve_message_assignment(
            assign_shadow,
            assign_message_id,
            job_id=job_id,
            request_id=assign_request_id,
            mailbox_email="estimator@example.test",
            assignment_token=first_assignment_token,
        )
        assign_shadow.commit()
        reserved_message = assign_shadow.execute(
            """SELECT match_status, matched_job_id, matched_request_id,
                      assignment_token, processed_at
               FROM quote_email_messages WHERE id=?""",
            (assign_message_id,),
        ).fetchone()
        assign_shadow.execute(
            """UPDATE quote_price_matches
               SET message_id=?, status='needs_review'
               WHERE id=?""",
            (assign_message_id, assign_match_id),
        )
        assign_shadow.execute(
            """UPDATE quote_request_materials
               SET status='needs_review', source_message_id=?
               WHERE quote_request_id=? AND material_id=?""",
            (
                assign_message_id,
                assign_request_id,
                isolated_materials[0]["id"],
            ),
        )
        assign_shadow.commit()
        active_assignment_blocked = False
        try:
            assign_shadow.execute("BEGIN IMMEDIATE")
            _reserve_message_assignment(
                assign_shadow,
                assign_message_id,
                job_id=job_id,
                request_id=assign_request_id,
                mailbox_email="estimator@example.test",
                assignment_token=uuid.uuid4().hex,
            )
            assign_shadow.commit()
        except ValueError:
            assign_shadow.rollback()
            active_assignment_blocked = True
        assign_shadow.execute(
            """UPDATE quote_email_messages SET assignment_started_at=?
               WHERE id=?""",
            (
                (
                    virtual_now
                    - timedelta(seconds=ASSIGNMENT_LEASE_SECONDS + 1)
                ).isoformat(),
                assign_message_id,
            ),
        )
        assign_shadow.commit()
        retry_assignment_token = uuid.uuid4().hex
        assign_shadow.execute("BEGIN IMMEDIATE")
        retried_request = _reserve_message_assignment(
            assign_shadow,
            assign_message_id,
            job_id=job_id,
            request_id=assign_request_id,
            mailbox_email="estimator@example.test",
            assignment_token=retry_assignment_token,
        )
        assign_shadow.commit()
        retried_message = assign_shadow.execute(
            """SELECT assignment_token, assignment_started_at
               FROM quote_email_messages WHERE id=?""",
            (assign_message_id,),
        ).fetchone()
        old_worker_fenced = False
        try:
            _record_price_match(
                message_id=assign_message_id,
                request=assigned_request,
                material=isolated_materials[0],
                product={
                    "item_code": isolated_materials[0]["item_code"],
                    "product_name": isolated_materials[0]["description"],
                    "unit_price": 4.25,
                    "unit": isolated_materials[0]["unit"],
                    "source_unit": isolated_materials[0]["unit"],
                },
                source={
                    "file_hash": _sha256_bytes(b"expired-worker"),
                    "file_name": "expired-worker.csv",
                },
                match_method="review",
                status="needs_review",
                reason="Expired worker write must be fenced.",
                assignment_token=first_assignment_token,
                conn=assign_shadow,
            )
        except AssignmentLeaseLostError:
            old_worker_fenced = True
        takeover_cleaned_partial_rows = bool(
            int(
                assign_shadow.execute(
                    """SELECT COUNT(*) FROM quote_price_matches
                       WHERE message_id=?""",
                    (assign_message_id,),
                ).fetchone()[0]
            )
            == 0
            and str(
                assign_shadow.execute(
                    """SELECT status FROM quote_request_materials
                       WHERE quote_request_id=? AND material_id=?""",
                    (
                        assign_request_id,
                        isolated_materials[0]["id"],
                    ),
                ).fetchone()["status"]
            )
            == "requested"
        )
        interrupted_assignment_resumable = bool(
            active_assignment_blocked
            and old_worker_fenced
            and takeover_cleaned_partial_rows
            and assigned_request
            and retried_request
            and reserved_message
            and reserved_message["match_status"] == "assigning"
            and int(reserved_message["matched_job_id"]) == job_id
            and int(reserved_message["matched_request_id"]) == assign_request_id
            and reserved_message["assignment_token"] == first_assignment_token
            and retried_message
            and retried_message["assignment_token"] == retry_assignment_token
        )
        assign_shadow.execute("BEGIN IMMEDIATE")
        assignment_finalized = _finalize_message_assignment(
            assign_shadow,
            assign_message_id,
            job_id=job_id,
            request_id=assign_request_id,
            reviewer_name="Simulation",
            assignment_token=retry_assignment_token,
            pricing={
                "products": 0,
                "applied": 0,
                "needs_review": 0,
                "parse_errors": [],
                "parsed_sources": [],
            },
        )
        assign_shadow.commit()
        assigned_message = assign_shadow.execute(
            """SELECT match_status, matched_job_id, matched_request_id,
                      assignment_token, assignment_started_at, processed_at
               FROM quote_email_messages WHERE id=?""",
            (assign_message_id,),
        ).fetchone()
        assignment_states = {
            int(row["job_id"]): str(row["status"])
            for row in assign_shadow.execute(
                """SELECT job_id, status FROM quote_match_candidates
                   WHERE message_id=?""",
                (assign_message_id,),
            ).fetchall()
        }
        ambiguous_assignment_resolvable = bool(
            assignment_finalized
            and assigned_request
            and int(assigned_request["id"]) == assign_request_id
            and assigned_message
            and assigned_message["match_status"] == "matched"
            and int(assigned_message["matched_job_id"]) == job_id
            and int(assigned_message["matched_request_id"]) == assign_request_id
            and not assigned_message["assignment_token"]
            and assigned_message["assignment_started_at"] is None
            and assigned_message["processed_at"] == reserved_message["processed_at"]
            and assignment_states.get(job_id) == "selected"
            and assignment_states.get(assign_other_job_id) == "rejected"
        )
    finally:
        assign_shadow.close()

    cancel_assign_shadow, cancel_assign_request_id, cancel_assign_match_id = (
        _simulation_shadow(
            job_id,
            isolated_materials,
            virtual_now,
        )
    )
    try:
        cancel_assign_message_id = 9_300_003
        cancel_assign_shadow.execute(
            """INSERT INTO quote_email_messages
               (id, graph_message_id, mailbox_email, direction, sender_email,
                recipients_json, subject, body_text, attachment_manifest_json,
                match_status, evidence_json, processed_at, created_at)
               VALUES (?, ?, ?, 'inbound', ?, '[]', ?, '', '[]',
                       'needs_review', '[]', ?, ?)""",
            (
                cancel_assign_message_id,
                "simulation-cancel-during-assignment",
                "estimator@example.test",
                "vendor-one@example.test",
                "Cancelled assignment simulation quote",
                virtual_now.isoformat(),
                virtual_now.isoformat(),
            ),
        )
        cancel_other_job_id, cancel_other_request_id = (
            _add_simulation_candidate_request(
                cancel_assign_shadow,
                source_job_id=job_id,
                source_request_id=cancel_assign_request_id,
                suffix=3,
            )
        )
        for candidate_job_id, candidate_request_id in (
            (job_id, cancel_assign_request_id),
            (cancel_other_job_id, cancel_other_request_id),
        ):
            cancel_assign_shadow.execute(
                """INSERT INTO quote_match_candidates
                   (message_id, job_id, quote_request_id, status,
                    match_method, evidence_json, created_at)
                   VALUES (?, ?, ?, 'candidate', 'exact_clues', '[]', ?)""",
                (
                    cancel_assign_message_id,
                    candidate_job_id,
                    candidate_request_id,
                    virtual_now.isoformat(),
                ),
            )
        cancel_assign_shadow.commit()
        cancel_assignment_token = uuid.uuid4().hex
        cancel_assign_shadow.execute("BEGIN IMMEDIATE")
        _reserve_message_assignment(
            cancel_assign_shadow,
            cancel_assign_message_id,
            job_id=job_id,
            request_id=cancel_assign_request_id,
            mailbox_email="estimator@example.test",
            assignment_token=cancel_assignment_token,
        )
        cancel_assign_shadow.commit()
        cancel_assign_shadow.execute(
            """UPDATE quote_price_matches
               SET message_id=?, status='needs_review'
               WHERE id=?""",
            (cancel_assign_message_id, cancel_assign_match_id),
        )
        cancel_assign_shadow.execute(
            """UPDATE quote_request_materials
               SET status='needs_review', source_message_id=?
               WHERE quote_request_id=? AND material_id=?""",
            (
                cancel_assign_message_id,
                cancel_assign_request_id,
                isolated_materials[0]["id"],
            ),
        )
        cancel_assign_shadow.commit()
        cancelled_during_assignment = cancel_request(
            cancel_assign_request_id,
            conn=cancel_assign_shadow,
        )
        cancel_assign_shadow.execute("BEGIN IMMEDIATE")
        _store_no_price_result(
            cancel_assign_shadow,
            message_id=cancel_assign_message_id,
            request_id=cancel_assign_request_id,
            evidence_json='[{"type":"parse_failure","value":"No price"}]',
            defer_request_state=True,
            assignment_token=cancel_assignment_token,
        )
        deferred_message = cancel_assign_shadow.execute(
            """SELECT match_status FROM quote_email_messages WHERE id=?""",
            (cancel_assign_message_id,),
        ).fetchone()
        deferred_request = cancel_assign_shadow.execute(
            "SELECT status FROM quote_requests WHERE id=?",
            (cancel_assign_request_id,),
        ).fetchone()
        no_price_deferred_safe = bool(
            cancelled_during_assignment
            and
            deferred_message
            and deferred_message["match_status"] == "assigning"
            and deferred_request
            and deferred_request["status"] == "cancelled"
        )
        cancel_assign_shadow.commit()
        cancel_assign_shadow.execute("BEGIN IMMEDIATE")
        cancelled_finalize = _finalize_message_assignment(
            cancel_assign_shadow,
            cancel_assign_message_id,
            job_id=job_id,
            request_id=cancel_assign_request_id,
            reviewer_name="Simulation",
            assignment_token=cancel_assignment_token,
            pricing={
                "products": 0,
                "applied": 0,
                "needs_review": 0,
                "parse_errors": [],
                "parsed_sources": [],
            },
        )
        cancel_assign_shadow.commit()
        cancel_message = cancel_assign_shadow.execute(
            """SELECT match_status, assignment_token
               FROM quote_email_messages WHERE id=?""",
            (cancel_assign_message_id,),
        ).fetchone()
        cancel_candidate_states = {
            int(row["job_id"]): str(row["status"])
            for row in cancel_assign_shadow.execute(
                """SELECT job_id, status FROM quote_match_candidates
                   WHERE message_id=?""",
                (cancel_assign_message_id,),
            ).fetchall()
        }
        cancelled_assignment_released = bool(
            not cancelled_finalize
            and cancel_message
            and cancel_message["match_status"] == "needs_review"
            and not cancel_message["assignment_token"]
            and cancel_candidate_states.get(job_id) == "rejected"
            and cancel_candidate_states.get(cancel_other_job_id) == "candidate"
            and int(
                cancel_assign_shadow.execute(
                    """SELECT COUNT(*) FROM quote_price_matches
                       WHERE message_id=? AND quote_request_id=?""",
                    (cancel_assign_message_id, cancel_assign_request_id),
                ).fetchone()[0]
            )
            == 0
            and str(
                cancel_assign_shadow.execute(
                    """SELECT status FROM quote_request_materials
                       WHERE quote_request_id=? AND material_id=?""",
                    (
                        cancel_assign_request_id,
                        isolated_materials[0]["id"],
                    ),
                ).fetchone()["status"]
            )
            == "requested"
        )
    finally:
        cancel_assign_shadow.close()

    stale_assign_shadow, stale_assign_request_id, _ = _simulation_shadow(
        job_id,
        isolated_materials,
        virtual_now,
    )
    try:
        stale_assign_message_id = 9_300_004
        stale_assign_shadow.execute(
            """INSERT INTO quote_email_messages
               (id, graph_message_id, mailbox_email, direction, sender_email,
                recipients_json, subject, body_text, attachment_manifest_json,
                match_status, evidence_json, processed_at, created_at)
               VALUES (?, ?, ?, 'inbound', ?, '[]', ?, '', '[]',
                       'needs_review', '[]', ?, ?)""",
            (
                stale_assign_message_id,
                "simulation-stale-during-assignment",
                "estimator@example.test",
                "vendor-one@example.test",
                "Stale assignment simulation quote",
                virtual_now.isoformat(),
                virtual_now.isoformat(),
            ),
        )
        stale_other_job_id, stale_other_request_id = (
            _add_simulation_candidate_request(
                stale_assign_shadow,
                source_job_id=job_id,
                source_request_id=stale_assign_request_id,
                suffix=4,
            )
        )
        for candidate_job_id, candidate_request_id in (
            (job_id, stale_assign_request_id),
            (stale_other_job_id, stale_other_request_id),
        ):
            stale_assign_shadow.execute(
                """INSERT INTO quote_match_candidates
                   (message_id, job_id, quote_request_id, status,
                    match_method, evidence_json, created_at)
                   VALUES (?, ?, ?, 'candidate', 'exact_clues', '[]', ?)""",
                (
                    stale_assign_message_id,
                    candidate_job_id,
                    candidate_request_id,
                    virtual_now.isoformat(),
                ),
            )
        stale_assign_shadow.commit()
        stale_assignment_token = uuid.uuid4().hex
        stale_assign_shadow.execute("BEGIN IMMEDIATE")
        _reserve_message_assignment(
            stale_assign_shadow,
            stale_assign_message_id,
            job_id=job_id,
            request_id=stale_assign_request_id,
            mailbox_email="estimator@example.test",
            assignment_token=stale_assignment_token,
        )
        stale_assign_shadow.commit()
        stale_assign_shadow.execute(
            """UPDATE job_materials
               SET order_qty=COALESCE(order_qty, installed_qty, 0)+1
               WHERE id=? AND job_id=?""",
            (isolated_materials[0]["id"], job_id),
        )
        stale_assign_shadow.commit()
        stale_assign_shadow.execute("BEGIN IMMEDIATE")
        stale_finalize = _finalize_message_assignment(
            stale_assign_shadow,
            stale_assign_message_id,
            job_id=job_id,
            request_id=stale_assign_request_id,
            reviewer_name="Simulation",
            assignment_token=stale_assignment_token,
            pricing={
                "products": 0,
                "applied": 0,
                "needs_review": 0,
                "parse_errors": [],
                "parsed_sources": [],
            },
        )
        stale_assign_shadow.commit()
        stale_message = stale_assign_shadow.execute(
            """SELECT match_status, assignment_token
               FROM quote_email_messages WHERE id=?""",
            (stale_assign_message_id,),
        ).fetchone()
        stale_request = stale_assign_shadow.execute(
            "SELECT status FROM quote_requests WHERE id=?",
            (stale_assign_request_id,),
        ).fetchone()
        stale_candidate_states = {
            int(row["job_id"]): str(row["status"])
            for row in stale_assign_shadow.execute(
                """SELECT job_id, status FROM quote_match_candidates
                   WHERE message_id=?""",
                (stale_assign_message_id,),
            ).fetchall()
        }
        stale_assignment_released = bool(
            not stale_finalize
            and stale_message
            and stale_message["match_status"] == "needs_review"
            and not stale_message["assignment_token"]
            and stale_request
            and stale_request["status"] == "stale"
            and stale_candidate_states.get(job_id) == "rejected"
            and stale_candidate_states.get(stale_other_job_id) == "candidate"
        )
    finally:
        stale_assign_shadow.close()

    abort_assign_shadow, abort_assign_request_id, abort_assign_match_id = (
        _simulation_shadow(
            job_id,
            isolated_materials,
            virtual_now,
        )
    )
    try:
        abort_assign_message_id = 9_300_005
        abort_assign_shadow.execute(
            """INSERT INTO quote_email_messages
               (id, graph_message_id, mailbox_email, direction, sender_email,
                recipients_json, subject, body_text, attachment_manifest_json,
                match_status, evidence_json, processed_at, created_at)
               VALUES (?, ?, ?, 'inbound', ?, '[]', ?, '', '[]',
                       'needs_review', '[]', ?, ?)""",
            (
                abort_assign_message_id,
                "simulation-abort-during-assignment",
                "estimator@example.test",
                "vendor-one@example.test",
                "Aborted assignment simulation quote",
                virtual_now.isoformat(),
                virtual_now.isoformat(),
            ),
        )
        abort_assign_shadow.execute(
            """INSERT INTO quote_match_candidates
               (message_id, job_id, quote_request_id, status,
                match_method, evidence_json, created_at)
               VALUES (?, ?, ?, 'candidate', 'exact_clues', '[]', ?)""",
            (
                abort_assign_message_id,
                job_id,
                abort_assign_request_id,
                virtual_now.isoformat(),
            ),
        )
        abort_assign_shadow.commit()
        abort_assignment_token = uuid.uuid4().hex
        abort_assign_shadow.execute("BEGIN IMMEDIATE")
        _reserve_message_assignment(
            abort_assign_shadow,
            abort_assign_message_id,
            job_id=job_id,
            request_id=abort_assign_request_id,
            mailbox_email="estimator@example.test",
            assignment_token=abort_assignment_token,
        )
        abort_assign_shadow.commit()
        abort_assign_shadow.execute(
            """UPDATE quote_price_matches
               SET message_id=?, status='needs_review'
               WHERE id=?""",
            (abort_assign_message_id, abort_assign_match_id),
        )
        abort_assign_shadow.execute(
            """UPDATE quote_request_materials
               SET status='needs_review', source_message_id=?
               WHERE quote_request_id=? AND material_id=?""",
            (
                abort_assign_message_id,
                abort_assign_request_id,
                isolated_materials[0]["id"],
            ),
        )
        abort_assign_shadow.commit()
        abort_assign_shadow.execute("BEGIN IMMEDIATE")
        aborted = _abort_message_assignment(
            abort_assign_shadow,
            abort_assign_message_id,
            request_id=abort_assign_request_id,
            assignment_token=abort_assignment_token,
        )
        abort_assign_shadow.commit()
        abort_message = abort_assign_shadow.execute(
            """SELECT match_status, assignment_token
               FROM quote_email_messages WHERE id=?""",
            (abort_assign_message_id,),
        ).fetchone()
        aborted_assignment_clean = bool(
            aborted
            and abort_message
            and abort_message["match_status"] == "needs_review"
            and not abort_message["assignment_token"]
            and int(
                abort_assign_shadow.execute(
                    """SELECT COUNT(*) FROM quote_price_matches
                       WHERE message_id=?""",
                    (abort_assign_message_id,),
                ).fetchone()[0]
            )
            == 0
            and str(
                abort_assign_shadow.execute(
                    """SELECT status FROM quote_request_materials
                       WHERE quote_request_id=? AND material_id=?""",
                    (
                        abort_assign_request_id,
                        isolated_materials[0]["id"],
                    ),
                ).fetchone()["status"]
            )
            == "requested"
            and str(
                abort_assign_shadow.execute(
                    """SELECT status FROM quote_match_candidates
                       WHERE message_id=? AND job_id=?""",
                    (abort_assign_message_id, job_id),
                ).fetchone()["status"]
            )
            == "candidate"
        )
        artifact_test_job = dict(
            abort_assign_shadow.execute(
                "SELECT * FROM jobs WHERE id=?",
                (job_id,),
            ).fetchone()
        )
        artifact_test_job_id = int(job_id) + 40_000_000
        artifact_test_job["id"] = artifact_test_job_id
        artifact_test_job["slug"] = (
            f"simulation-artifact-rollback-{artifact_test_job_id}"
        )
        artifact_test_job["project_name"] = "Simulation Artifact Rollback"
        _insert_shadow_row(
            abort_assign_shadow,
            "jobs",
            artifact_test_job,
        )
        abort_assign_shadow.commit()
        artifact_source = _write_artifact(
            f"uncommitted-{uuid.uuid4().hex}".encode("utf-8"),
            filename="simulation-uncommitted-quote.txt",
            artifact_kind="simulation",
        )
        uncommitted_paths: list[str] = []
        staged_artifacts: list[dict] = []
        try:
            staged_artifacts, uncommitted_paths = (
                _stage_assignment_artifacts(
                    [artifact_source],
                    job_id=artifact_test_job_id,
                    assignment_token=uuid.uuid4().hex,
                )
            )
            abort_assign_shadow.execute("BEGIN IMMEDIATE")
            for staged_artifact in staged_artifacts:
                record_job_artifact(
                    artifact_test_job_id,
                    "vendor_quote",
                    staged_artifact["artifact_path"],
                    staged_artifact["file_hash"],
                    staged_artifact["file_size"],
                    conn=abort_assign_shadow,
                )
                record_imported_file(
                    artifact_test_job_id,
                    staged_artifact["file_name"],
                    staged_artifact["file_hash"],
                    staged_artifact["file_size"],
                    source="outlook",
                    artifact_path=staged_artifact["artifact_path"],
                    artifact_kind="vendor_quote",
                    conn=abort_assign_shadow,
                )
            receipts_written_before_rollback = bool(
                abort_assign_shadow.execute(
                    """SELECT 1 FROM job_artifacts
                       WHERE job_id=?""",
                    (artifact_test_job_id,),
                ).fetchone()
                and abort_assign_shadow.execute(
                    """SELECT 1 FROM imported_files
                       WHERE job_id=?""",
                    (artifact_test_job_id,),
                ).fetchone()
            )
            abort_assign_shadow.rollback()
            _remove_uncommitted_artifact_files(uncommitted_paths)
            artifact_rollback_clean = bool(
                receipts_written_before_rollback
                and staged_artifacts
                and uncommitted_paths
                and not (
                    ARTIFACT_ROOT
                    / str(staged_artifacts[0]["artifact_path"])
                ).is_file()
                and not abort_assign_shadow.execute(
                    """SELECT 1 FROM job_artifacts
                       WHERE job_id=?""",
                    (artifact_test_job_id,),
                ).fetchone()
                and not abort_assign_shadow.execute(
                    """SELECT 1 FROM imported_files
                       WHERE job_id=?""",
                    (artifact_test_job_id,),
                ).fetchone()
            )
        finally:
            abort_assign_shadow.rollback()
            try:
                Path(artifact_source["absolute_path"]).unlink()
            except OSError:
                pass
    finally:
        abort_assign_shadow.close()

    legacy_received_state = _request_workflow_state(
        {
            "status": "received",
            "material_snapshot": [{"id": 999_999_999}],
            "material_snapshot_hash": "legacy-missing-hash",
        },
        {"materials": []},
    )

    followup_shadow, followup_request_id, _ = _simulation_shadow(
        job_id,
        isolated_materials,
        virtual_now,
    )
    try:
        claimed_followups = due_followups(
            add_business_days(virtual_now, 3),
            conn=followup_shadow,
        )
        final_snapshot_blocked = False
        if claimed_followups:
            followup_event_id = int(claimed_followups[0]["id"])
            followup_shadow.execute(
                """UPDATE job_materials
                   SET order_qty=COALESCE(order_qty, installed_qty, 0)+1
                   WHERE id=? AND job_id=?""",
                (isolated_materials[0]["id"], job_id),
            )
            followup_shadow.commit()
            sendable_after_change = confirm_followup_send(
                followup_event_id,
                conn=followup_shadow,
            )
            followup_request = followup_shadow.execute(
                "SELECT status FROM quote_requests WHERE id=?",
                (followup_request_id,),
            ).fetchone()
            followup_event = followup_shadow.execute(
                "SELECT status FROM quote_followup_events WHERE id=?",
                (followup_event_id,),
            ).fetchone()
            final_snapshot_blocked = bool(
                not sendable_after_change
                and followup_request
                and followup_request["status"] == "stale"
                and followup_event
                and followup_event["status"] == "cancelled"
            )
    finally:
        followup_shadow.close()

    all_results = [
        _simulation_result(
            "normal_reply",
            "Real price transaction saves the price and all proof together",
            True,
            shadow_price_transaction_ok,
            "Runs the production price-application transaction in an isolated copy of this bid.",
        ),
        _simulation_result(
            "changed_price",
            "Josh can explicitly choose the quote after the current price changed",
            True,
            bool(
                changed_price_case["applied"]
                and changed_price_case["material_price"] == 4.25
            ),
            "The current price shown to Josh is checked again inside the locked transaction.",
        ),
        _simulation_result(
            "changed_price",
            "An old screen cannot overwrite a newer material price",
            True,
            bool(
                not stale_screen_case["applied"]
                and stale_screen_case["material_price"] == 3.95
                and stale_screen_case["match_status"] == "needs_review"
            ),
            "A stale expected price is rejected without changing the copied bid.",
        ),
        _simulation_result(
            "changed_price",
            "Cancelling a request also cancels its pending prices and follow-ups",
            True,
            cancelled_request_is_terminal,
            "The real cancellation transaction runs against an isolated copy.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "Removing an email from one bid keeps its other possible bid",
            True,
            per_bid_ignore_isolated,
            "The real per-bid ignore transaction keeps the other candidate open.",
        ),
        _simulation_result(
            "normal_reply",
            "Old received requests are treated as complete",
            True,
            bool(
                legacy_received_state["status"] == "complete"
                and legacy_received_state["stale"] is False
            ),
            "The production workflow state helper terminalizes the old record before checking drift.",
        ),
        _simulation_result(
            "no_response",
            "Real follow-up claims happen once at day 3 and day 6",
            True,
            shadow_followups_once,
            "Runs the production due-follow-up claim code in the isolated copy.",
        ),
        _simulation_result(
            "no_response",
            "A last-second quantity or unit change blocks the follow-up",
            True,
            final_snapshot_blocked,
            "The final production send check replays the locked material snapshot.",
        ),
        _simulation_result(
            "normal_reply",
            "Direct reply can auto-price only the requested exact item",
            "ready_to_apply",
            direct_evaluation["status"],
            "Uses the real matcher, CSV parser, and price safety evaluator.",
        ),
        _simulation_result(
            "changed_subject",
            "Changed subject still follows the recorded conversation",
            "conversation",
            changed_subject_match.get("method"),
            "The subject text is not used as an identifier.",
        ),
        _simulation_result(
            "standalone_email",
            "Standalone email price waits for Josh",
            "needs_review",
            standalone_evaluation["status"],
            "Exact job facts identify the bid, but the price is not automatic.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "Email that could fit two bids waits for Josh",
            "needs_review",
            ambiguous_match.get("status"),
            "The matcher refuses to choose between multiple open requests.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "Ambiguous email can be assigned after inbox processing",
            True,
            ambiguous_assignment_resolvable,
            "The production assignment transaction resolves a processed review email to one exact request.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "Interrupted assignment can retry the same exact bid",
            True,
            interrupted_assignment_resumable,
            "A stopped assignment can resume only after its two-minute safety lease expires.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "A second live worker cannot price the same email",
            True,
            active_assignment_blocked,
            "The active assignment token blocks a second click or browser from running at the same time.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "An expired worker cannot write after a retry takes over",
            True,
            bool(old_worker_fenced and takeover_cleaned_partial_rows),
            "The replacement token removes partial rows and fences every later write from the old worker.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "Cancelling during assignment releases the email without keeping prices",
            True,
            cancelled_assignment_released,
            "The selected bid is rejected, partial price rows are removed, and the other possible bid stays available.",
        ),
        _simulation_result(
            "materials_changed",
            "Changing bid materials during assignment makes the request stale",
            True,
            stale_assignment_released,
            "The final locked check rejects the changed bid and keeps the email available for any other exact candidate.",
        ),
        _simulation_result(
            "broken_attachment",
            "A no-price email cannot reopen a cancelled request",
            True,
            no_price_deferred_safe,
            "Deferred parsing keeps the assignment lock and leaves terminal request states untouched.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "A stopped price read removes partial rows and keeps the candidate",
            True,
            aborted_assignment_clean,
            "An unexpected parser or storage failure returns the email to review with no hidden price changes.",
        ),
        _simulation_result(
            "broken_attachment",
            "A rolled-back assignment removes its uncommitted job file",
            True,
            artifact_rollback_clean,
            "The durable file cleanup runs after the database rollback so orphaned vendor files are not left on the bid.",
        ),
        _simulation_result(
            "ambiguous_bid",
            "Unrelated email from a known vendor is ignored",
            "ignored",
            unrelated_vendor_match.get("status"),
            "A vendor address alone is not enough to create review work.",
        ),
        _simulation_result(
            "partial_quote",
            "Partial quote leaves only missing materials open",
            1,
            len(requested_materials) - len(parsed_products),
            "The fake attachment contains one of two requested materials.",
        ),
        _simulation_result(
            "changed_price",
            "Existing price is never silently replaced",
            "needs_review",
            changed_price_evaluation["status"],
            "The real safety evaluator sees a current bid price.",
        ),
        _simulation_result(
            "wrong_unit",
            "Wrong unit waits for Josh",
            "needs_review",
            wrong_unit_evaluation["status"],
            "The quoted unit and requested unit must be exactly equal.",
        ),
        _simulation_result(
            "wrong_unit",
            "BOX cannot silently equal CTN",
            "needs_review",
            packaging_unit_evaluation["status"],
            "Packaging aliases are shown to Josh instead of being treated as exact.",
        ),
        _simulation_result(
            "duplicate_reply",
            "The same Outlook message is stored once",
            True,
            duplicate_guard,
            "The deployed SQLite table enforces a unique immutable message ID.",
        ),
        _simulation_result(
            "broken_attachment",
            "Unreadable attachment waits for Josh",
            "needs_review",
            broken_attachment_status,
            "The real XLSX parser was given a broken file.",
        ),
        _simulation_result(
            "no_response",
            "Exactly two business-day follow-ups are scheduled",
            [3, 6],
            [
                business_days_between(
                    virtual_now,
                    add_business_days(virtual_now, day),
                )
                for day in (3, 6)
            ],
            "Dates are calculated by the production business-day function.",
        ),
        _simulation_result(
            "materials_changed",
            "Changed bid cannot auto-price",
            "needs_review",
            stale_evaluation["status"],
            "A stale locked snapshot always requires review.",
        ),
        _simulation_result(
            "send_failure",
            "Uncertain Microsoft delivery is never sent again automatically",
            ["send_failed", "send_uncertain"],
            [
                delivery_failure_status(microsoft_may_have_accepted=False),
                delivery_failure_status(microsoft_may_have_accepted=True),
            ],
            "A definite rejection and an uncertain timeout have different states.",
        ),
        _simulation_result(
            "normal_reply",
            "Requested quantity is not mistaken for a price",
            0,
            len(quantity_products),
            "Free-text prices require an explicit dollar sign.",
        ),
        _simulation_result(
            "normal_reply",
            "A second fake vendor cannot take the first vendor reply",
            9_100_001,
            multi_vendor_match.get("request_id"),
            "Reply headers identify one recorded request even with several vendors.",
        ),
        _simulation_result(
            "normal_reply",
            "Deterministic quote engine imports no OpenAI client",
            [],
            openai_imports,
            "This checks the module that performs matching, parsing, pricing safety, and follow-up rules.",
        ),
    ]
    results = (
        all_results
        if scenario == "all"
        else [result for result in all_results if result["scenario"] == scenario]
    )
    passed = bool(results) and all(result["passed"] for result in results)
    fingerprint_after = _live_job_fingerprint(load_job(job_id) or {})
    followups = [
        {
            "followup_number": number,
            "scheduled_for": add_business_days(virtual_now, days).isoformat(),
            "status": "scheduled",
            "sent_count": 0,
        }
        for number, days in ((1, 3), (2, 6))
    ]
    snapshot = {
        "job_id": job_id,
        "live_job_fingerprint": fingerprint_before,
        "source_job_snapshot": {
            key: job.get(key)
            for key in (
                "id",
                "project_name",
                "gc_name",
                "address",
                "city",
                "state",
                "zip",
            )
        },
        "source_material_snapshot": _material_snapshot(job.get("materials") or []),
        "isolated_materials": isolated_materials,
        "fake_request": fake_request,
        "fake_vendors": [
            "vendor-one@example.test",
            "vendor-two@example.test",
        ],
    }
    result = {
        "status": "pass" if passed else "fail",
        "checks": results,
        "timeline": [
            {"at": virtual_now.isoformat(), "event": "Isolated simulation started"},
            {
                "at": virtual_now.isoformat(),
                "event": (
                    f"{sum(1 for item in results if item['passed'])}/"
                    f"{len(results)} production-rule checks passed"
                ),
            },
        ],
        "followups": followups,
        "ai_calls": 0,
        "live_job_fingerprint_before": fingerprint_before,
        "live_job_fingerprint_after": fingerprint_after,
        "live_job_mutated": fingerprint_before != fingerprint_after,
        "engine_functions": [
            "match_message_to_requests",
            "parse_quote_file_deterministic",
            "evaluate_price_candidate",
            "add_business_days",
            "delivery_failure_status",
        ],
    }
    if result["live_job_mutated"]:
        result["status"] = "fail"
    conn = _get_conn()
    try:
        now = iso_now()
        cursor = conn.execute(
            """INSERT INTO quote_simulation_runs
               (job_id, scenario, status, virtual_now, snapshot_json,
                result_json, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                job_id,
                scenario,
                result["status"],
                virtual_now.isoformat(),
                json.dumps(snapshot),
                json.dumps(result),
                now,
                now,
            ),
        )
        conn.commit()
        run_id = int(cursor.lastrowid)
    finally:
        conn.close()
    return get_simulation_run(job_id, run_id) or {}


def get_simulation_run(job_id: int, run_id: int) -> dict | None:
    conn = _get_conn()
    try:
        row = conn.execute(
            "SELECT * FROM quote_simulation_runs WHERE id=? AND job_id=?",
            (run_id, job_id),
        ).fetchone()
        if not row:
            return None
        result = dict(row)
        result["snapshot"] = _json_loads(result.get("snapshot_json"), {})
        result["result"] = _json_loads(result.get("result_json"), {})
        return result
    finally:
        conn.close()


def advance_simulation_run(job_id: int, run_id: int, business_days: int = 1) -> dict:
    run = get_simulation_run(job_id, run_id)
    if not run:
        raise ValueError("Simulation run not found.")
    current = _parse_iso(run.get("virtual_now")) or utc_now()
    advanced = add_business_days(current, max(1, min(int(business_days), 30)))
    result = run["result"]
    timeline = list(result.get("timeline") or [])
    timeline.append(
        {
            "at": advanced.isoformat(),
            "event": f"Advanced simulated business time by {business_days} day(s)",
        }
    )
    followups = list(result.get("followups") or [])
    for followup in followups:
        due_at = _parse_iso(followup.get("scheduled_for"))
        if (
            followup.get("status") == "scheduled"
            and due_at
            and due_at <= advanced
        ):
            followup["status"] = "sent"
            followup["sent_count"] = int(followup.get("sent_count") or 0) + 1
            followup["sent_at"] = advanced.isoformat()
            timeline.append(
                {
                    "at": advanced.isoformat(),
                    "event": (
                        f"Follow-up {followup['followup_number']} "
                        "was previewed as due"
                    ),
                    "status": "scheduled",
                }
            )
    result["followups"] = followups
    followup_counts = [
        int(followup.get("sent_count") or 0)
        for followup in followups
    ]
    result["followups_sent_once"] = all(count <= 1 for count in followup_counts)
    if not result["followups_sent_once"]:
        result["status"] = "fail"
    result["timeline"] = timeline
    conn = _get_conn()
    try:
        conn.execute(
            """UPDATE quote_simulation_runs
               SET status=?, virtual_now=?, result_json=?, updated_at=?
               WHERE id=? AND job_id=?""",
            (
                result.get("status") or "fail",
                advanced.isoformat(),
                json.dumps(result),
                iso_now(),
                run_id,
                job_id,
            ),
        )
        conn.commit()
    finally:
        conn.close()
    return get_simulation_run(job_id, run_id) or {}


def simulation_report_markdown(job_id: int, run_id: int) -> str:
    run = get_simulation_run(job_id, run_id)
    if not run:
        raise ValueError("Simulation run not found.")
    result = run.get("result") or {}
    lines = [
        "# Material Quote Simulation Report",
        "",
        f"- Bid ID: {job_id}",
        f"- Run ID: {run_id}",
        f"- Scenario: {run.get('scenario')}",
        f"- Status: {str(result.get('status') or run.get('status') or '').upper()}",
        f"- AI calls: {int(result.get('ai_calls') or 0)}",
        f"- Real bid changed: {'YES' if result.get('live_job_mutated') else 'NO'}",
        "",
        "## Safety Checks",
        "",
    ]
    for check in result.get("checks") or []:
        mark = "PASS" if check.get("passed") else "FAIL"
        lines.append(f"- {mark}: {check.get('label') or check.get('scenario')}")
        if check.get("detail"):
            lines.append(f"  {check['detail']}")
        lines.append(
            f"  Expected: {check.get('expected')} | Actual: {check.get('actual')}"
        )
    lines.extend(["", "## Timeline", ""])
    for event in result.get("timeline") or []:
        lines.append(f"- {event.get('at')}: {event.get('event')}")
    lines.extend(
        [
            "",
            "## Plain English",
            "",
            "This test copied the selected bid into memory and used fake vendors, a fake mailbox, and a fake clock.",
            "It called the same deterministic matching, parsing, price-safety, and business-day functions used by the live quote workflow.",
            "It did not contact Microsoft, send email, call OpenAI, or change the real bid.",
            "",
        ]
    )
    return "\n".join(lines)


def deterministic_contract() -> dict:
    request = {
        "id": 1,
        "job_id": 10,
        "vendor_email": "vendor@example.com",
        "contact_email": "vendor@example.com",
        "internet_message_id": "<request@example.com>",
        "conversation_id": "conversation",
        "project_name": "Sun Valley",
        "slug": "sun-valley",
        "gc_name": "Harness GC",
        "address": "123 Main Street",
        "material_snapshot": [
            {
                "id": 101,
                "item_code": "CPT-100",
                "description": "Harness Carpet",
                "unit": "SY",
                "quantity": 100,
            }
        ],
    }
    direct_reply = match_message_to_requests(
        {
            "sender_email": "vendor@example.com",
            "subject": "Changed",
            "headers": {"References": "<request@example.com>"},
        },
        [request],
    )
    unverified_reply = match_message_to_requests(
        {
            "sender_email": "stranger@example.net",
            "subject": "Changed",
            "headers": {"References": "<request@example.com>"},
        },
        [request],
    )
    unit_and_total = _products_from_rows(
        [
            ["Item Code", "Unit Price", "Cost", "Unit"],
            ["CPT-100", "4.25", "425.00", "SY"],
        ]
    )
    generic_price_rows = _products_from_rows(
        [
            ["Item Code", "Qty", "Price", "Unit"],
            ["CPT-100", "100", "425.00", "SY"],
        ]
    )
    try:
        _products_from_rows(
            [
                ["Item Code", "Unit Price", "Net Price", "Unit"],
                ["CPT-100", "4.25", "4.10", "SY"],
            ]
        )
        duplicate_unit_columns_rejected = False
    except ValueError:
        duplicate_unit_columns_rejected = True
    labeled_ocr_products = _products_from_text(
        "\n".join(
            [
                "VENDOR: Harness Supply",
                "ITEM CODE: OCR-100",
                "PRODUCT: Harness Scanned Tile",
                "UNIT PRICE: $7.25 PER SF",
            ]
        )
    )
    labeled_ocr_product = (
        labeled_ocr_products[0]
        if len(labeled_ocr_products) == 1
        else {}
    )
    module_source = Path(__file__).read_text(encoding="utf-8")
    ai_dependency_imports = re.findall(
        r"^\s*(?:from\s+(?:openai|ai_client)\b|import\s+(?:openai|ai_client)\b)",
        module_source,
        flags=re.MULTILINE,
    )
    mixed_attachment_email = EmailMessage()
    mixed_attachment_email["From"] = "vendor@example.com"
    mixed_attachment_email["To"] = "estimator@example.com"
    mixed_attachment_email["Subject"] = "Mixed attachment safety proof"
    mixed_attachment_email.set_content("Pricing is attached.")
    mixed_attachment_email.add_attachment(
        (
            "item_code,description,unit_price,unit\n"
            "CPT-100,Harness Carpet,4.25,SY\n"
        ).encode("utf-8"),
        maintype="text",
        subtype="csv",
        filename="valid.csv",
    )
    mixed_attachment_email.add_attachment(
        b"not a workbook",
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="broken.xlsx",
    )
    mixed_email_bytes = mixed_attachment_email.as_bytes()
    mixed_email_path = (
        ARTIFACT_ROOT
        / "shared"
        / "tmp"
        / f"contract-mixed-{_sha256_bytes(mixed_email_bytes)[:12]}.eml"
    )
    mixed_email_path.parent.mkdir(parents=True, exist_ok=True)
    mixed_email_path.write_bytes(mixed_email_bytes)
    try:
        parse_quote_file_deterministic(str(mixed_email_path))
        mixed_attachment_parser_rejected = False
    except Exception:
        mixed_attachment_parser_rejected = True
    checks = {
        "direct_reply": direct_reply["status"]
        == "matched",
        "direct_reply_verified_sender": direct_reply["method"] == "direct_reply",
        "unverified_reply_cannot_auto_price": (
            unverified_reply["method"] == "direct_reply_unverified_sender"
        ),
        "quantity_not_price": len(
            _products_from_text("CPT-100 Harness Carpet 100 SY")
        )
        == 0,
        "explicit_text_price": len(
            _products_from_text("CPT-100 Harness Carpet $4.25/SY")
        )
        == 1,
        "short_code_not_substring": not contains_exact_fact(
            "Carpet pricing attached",
            "CAR",
        ),
        "punctuation_normalized_exactly": contains_exact_fact(
            "Pricing for AE 100",
            "AE-100",
        ),
        "total_cost_not_unit_price": (
            len(unit_and_total) == 1
            and unit_and_total[0]["unit_price"] == 4.25
        ),
        "generic_price_requires_review": generic_price_rows == [],
        "duplicate_unit_columns_rejected": duplicate_unit_columns_rejected,
        "accounting_negative_rejected": _positive_price("(5.00)") is None,
        "decimal_comma_rejected": _positive_price("7,25") is None,
        "broken_sibling_blocks_automatic": _message_price_decision(
            {
                "status": "ready_to_apply",
                "reason": "Exact direct reply.",
            },
            [{"type": "parse_failure", "value": "broken.xlsx"}],
        )[0]
        == "needs_review",
        "broken_sibling_parser_rejected": mixed_attachment_parser_rejected,
        "labeled_ocr_price": (
            labeled_ocr_product.get("item_code") == "OCR-100"
            and labeled_ocr_product.get("product_name")
            == "OCR-100 - Harness Scanned Tile"
            and labeled_ocr_product.get("vendor") == "Harness Supply"
            and labeled_ocr_product.get("unit_price") == 7.25
            and labeled_ocr_product.get("unit") == "SF"
            and labeled_ocr_product.get("source_unit") == "SF"
        ),
        "ai_dependencies_absent": not ai_dependency_imports,
        "standalone_exact": match_message_to_requests(
            {
                "sender_email": "vendor@example.com",
                "subject": "Sun Valley quote",
                "body_text": "CPT-100",
            },
            [request],
        )["status"]
        == "matched",
        "sender_only_ignored": match_message_to_requests(
            {
                "sender_email": "vendor@example.com",
                "subject": "Quote attached",
                "body_text": "",
            },
            [request],
        )["status"]
        == "ignored",
        "unrelated_ignored": match_message_to_requests(
            {
                "sender_email": "newsletter@example.net",
                "subject": "News",
                "body_text": "",
            },
            [request],
        )["status"]
        == "ignored",
        "ai_calls": 0,
    }
    passed = all(
        value is True
        for name, value in checks.items()
        if name != "ai_calls"
    ) and checks["ai_calls"] == 0
    return {
        "status": "pass" if passed else "fail",
        "checks": checks,
        "matching_engine": "deterministic-v1",
    }
